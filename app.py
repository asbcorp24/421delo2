# -*- coding: utf-8 -*-
import json
import os
import io
import csv
import hmac
import secrets
import sqlite3
import datetime as dt
from pathlib import Path
from typing import Optional
from functools import wraps
from models import *
from models import ActivityTemplate, MemoTemplate
from flask import (
    Flask, request, jsonify, render_template, redirect,
    url_for, flash, send_from_directory, send_file, abort
)
from sqlalchemy.orm import joinedload
from flask_sqlalchemy import SQLAlchemy
from flask_login import (
    LoginManager, UserMixin, login_user, logout_user,
    current_user, login_required
)
from sqlalchemy import text
from werkzeug.utils import secure_filename
from datetime import datetime
from sqlalchemy.orm import joinedload
from sqlalchemy import and_, or_, func
import uuid
from werkzeug.utils import secure_filename
import pytesseract
from PIL import Image
import pdfplumber
import json
from docx import Document
import docx
import re
from flask import current_app
import time
import glob
import queue
import threading
import urllib.error
import urllib.request
from config import SUPERADMIN_LOGIN, SUPERADMIN_PASSWORD,SECRET_PASSPHRASE
from dotenv import load_dotenv
import os
from urllib.parse import quote
from backup_database import BACKUP_DIR, create_backup, list_backups
load_dotenv()
from sqlalchemy import text as sql_text
pytesseract.pytesseract.tesseract_cmd = r"C:\Program Files\Tesseract-OCR\tesseract.exe"
# =========================
# ????????????
# =========================
import sys
from pathlib import Path
if getattr(sys, 'frozen', False):
    BASE_DIR = Path(sys.executable).resolve().parent
else:
    BASE_DIR = Path(__file__).resolve().parent
UPLOAD_DIR = BASE_DIR / "uploads"
UPLOAD_DIR.mkdir(exist_ok=True)
# === Crypto Utils ===
import base64
import os
from cryptography.hazmat.primitives.kdf.pbkdf2 import PBKDF2HMAC
from cryptography.hazmat.primitives import hashes
from cryptography.hazmat.primitives.ciphers.aead import AESGCM


def get_department_prefix(department_id, field_name: str, fallback: str) -> str:
    if department_id:
        department = db.session.get(Department, department_id)
        if department:
            department_prefix = (getattr(department, field_name, None) or "").strip()
            if department_prefix:
                return department_prefix
    return fallback.strip()


def get_department_start_number(department_id, field_name: str, fallback: int) -> int:
    if department_id:
        department = db.session.get(Department, department_id)
        if department:
            value = getattr(department, field_name, None)
            if isinstance(value, int) and value > 0:
                return value
    return fallback


def generate_next_letter_reg_number(department_id=None) -> str:
    prefix = get_department_prefix(
        department_id,
        "letter_prefix",
        get_app_setting("letter_prefix", ""),
    )
    start_number = get_department_start_number(
        department_id,
        "letter_start_number",
        get_app_setting_int("letter_start_number", 1),
    )
    memo_source_id = ensure_memo_source()
    max_number = max(start_number - 1, 0)
    department = db.session.get(Department, department_id) if department_id else None
    series = (department.letter_number_series or "").strip() if department else ""
    letter_query = db.session.query(Letter.reg_number).filter(or_(Letter.source_id.is_(None), Letter.source_id != memo_source_id))
    if series:
        department_ids = [d.id for d in Department.query.filter_by(letter_number_series=series).all()]
        letter_query = letter_query.join(User, Letter.author_id == User.id).filter(User.department_id.in_(department_ids))
    letter_numbers = letter_query.all()
    for value in letter_numbers:
        reg_number = (value[0] or "").strip()
        if series:
            match = re.search(r"(?:-|^)(\d+)$", reg_number)
        elif prefix:
            match = re.search(rf"^{re.escape(prefix)}-(\d+)$", reg_number)
        else:
            match = re.search(r"^(\d+)$", reg_number)
        if match:
            max_number = max(max_number, int(match.group(1)))
    next_number = max_number + 1
    return f"{prefix}-{next_number}" if prefix else str(next_number)


def get_app_setting(key: str, default: str = "") -> str:
    row = db.session.execute(
        sql_text("SELECT value FROM app_settings WHERE key = :key"),
        {"key": key},
    ).fetchone()
    if not row or row[0] is None:
        return default
    return str(row[0])


def set_app_setting(key: str, value: str) -> None:
    db.session.execute(
        sql_text("""
            INSERT INTO app_settings (key, value)
            VALUES (:key, :value)
            ON CONFLICT(key) DO UPDATE SET value = excluded.value
        """),
        {"key": key, "value": value},
    )
    db.session.commit()


def get_app_setting_int(key: str, default: int = 1) -> int:
    raw = (get_app_setting(key, str(default)) or "").strip()
    try:
        value = int(raw)
    except (TypeError, ValueError):
        return default
    return value if value > 0 else default


def ensure_memo_source() -> int:
    source = LetterSource.query.filter_by(name="Служебная записка").first()
    if source:
        return source.id
    source = LetterSource(name="Служебная записка")
    db.session.add(source)
    db.session.commit()
    return source.id


def generate_next_memo_reg_number(department_id=None) -> str:
    prefix = get_department_prefix(
        department_id,
        "memo_prefix",
        get_app_setting("memo_prefix", "421"),
    ) or "421"
    start_number = get_department_start_number(
        department_id,
        "memo_start_number",
        get_app_setting_int("memo_start_number", 1),
    )
    memo_source_id = ensure_memo_source()
    max_number = max(start_number - 1, 0)
    department = db.session.get(Department, department_id) if department_id else None
    series = (department.memo_number_series or "").strip() if department else ""
    memo_query = db.session.query(Letter.reg_number).filter(Letter.source_id == memo_source_id)
    if series:
        department_ids = [d.id for d in Department.query.filter_by(memo_number_series=series).all()]
        memo_query = memo_query.join(User, Letter.author_id == User.id).filter(User.department_id.in_(department_ids))
    memo_numbers = memo_query.all()
    for value in memo_numbers:
        reg_number = (value[0] or "").strip()
        match = re.search(r"(?:-|^)(\d+)$", reg_number) if series else re.search(rf"^{re.escape(prefix)}-(\d+)$", reg_number)
        if match:
            max_number = max(max_number, int(match.group(1)))
    return f"{prefix}-{max_number + 1}"


def user_can_manage_reg_numbers() -> bool:
    return current_user.is_authenticated and getattr(current_user, "role", None) in ("admin", "superadmin")


def validate_reg_number_uniqueness(reg_number: str, *, is_memo: bool, exclude_letter_id: Optional[int] = None) -> bool:
    query = Letter.query.filter(func.lower(Letter.reg_number) == reg_number.lower())
    if exclude_letter_id is not None:
        query = query.filter(Letter.id != exclude_letter_id)
    memo_source_id = ensure_memo_source()
    if is_memo:
        query = query.filter(Letter.source_id == memo_source_id)
    else:
        query = query.filter(or_(Letter.source_id.is_(None), Letter.source_id != memo_source_id))
    return query.first() is None


def get_chat_sender_name(message: ChatMessage) -> str:
    if getattr(message, "sender", None) and getattr(message.sender, "full_name", None):
        return message.sender.full_name
    if getattr(message, "sender_id", None) == -1:
        return "Суперадмин"
    return "Пользователь"


def serialize_chat_message(message: ChatMessage) -> dict:
    return {
        "id": message.id,
        "sender_id": message.sender_id,
        "sender": get_chat_sender_name(message),
        "receiver_id": message.receiver_id,
        "kind": "general" if message.receiver_id is None else "direct",
        "text": message.text or "",
        "file": url_for("static", filename=message.file_path) if message.file_path else None,
        "time": message.created_at.strftime("%H:%M"),
    }


def _derive_key(passphrase: str, salt: bytes) -> bytes:
    kdf = PBKDF2HMAC(
        algorithm=hashes.SHA256(),
        length=32,
        salt=salt,
        iterations=200_000,
    )
    return kdf.derive(passphrase.encode("utf-8"))

def encrypt_str(plaintext: str, passphrase: str) -> str:
    salt = os.urandom(16)
    key = _derive_key(passphrase, salt)
    aes = AESGCM(key)
    nonce = os.urandom(12)
    ct = aes.encrypt(nonce, plaintext.encode("utf-8"), None)
    return base64.b64encode(salt + nonce + ct).decode("utf-8")

def decrypt_str(token_b64: str, passphrase: str) -> str:
    data = base64.b64decode(token_b64)
    salt, nonce, ct = data[:16], data[16:28], data[28:]
    key = _derive_key(passphrase, salt)
    aes = AESGCM(key)
    pt = aes.decrypt(nonce, ct, None)
    return pt.decode("utf-8")
app = Flask(__name__, template_folder=str(BASE_DIR / "templates"))
app.config.update(
    SECRET_KEY=os.getenv("SECRET_KEY", "dev-secret"),
    SQLALCHEMY_DATABASE_URI="sqlite:///app.db",
    SQLALCHEMY_TRACK_MODIFICATIONS=False,
    UPLOAD_FOLDER=str(UPLOAD_DIR),
)
app.config["SQLALCHEMY_ECHO"] = True
import unicodedata
def set_password(self, raw: str):
    self.password_enc = encrypt_str(raw, SECRET_PASSPHRASE)

def check_password(self, raw: str) -> bool:
    return raw == decrypt_str(self.password_enc, SECRET_PASSPHRASE)
def cleanup_old_chat_files(days=3):
    limit_time = dt.datetime.utcnow() - dt.timedelta(days=days)
    folder = os.path.join(app.root_path, "static", "chat_uploads")
    if os.path.exists(folder):
        now = time.time()
        for f in glob.glob(os.path.join(folder, "*")):
            try:
                if now - os.path.getmtime(f) > days * 86400:
                    os.remove(f)
            except Exception as e:
                print(f"cleanup error {f}: {e}")

with app.app_context():
    cleanup_old_chat_files()

from flask_login import UserMixin

class SuperAdmin(UserMixin):
    def __init__(self):
        self.id = -1
        self.username = SUPERADMIN_LOGIN
        self.full_name = "??????????????????"
        self.role = "superadmin"
        self.is_approved = True

    def get_id(self):
        return str(self.id)


def secure_filename_rus(filename):
    filename = filename.strip()
    filename = filename.replace(" ", "_")
    filename = re.sub(r"[^?-??-?A-Za-z0-9._-]", "", filename)
    return filename[:120]

KREUZBERG_URL = os.getenv("KREUZBERG_URL", "http://127.0.0.1:8000").rstrip("/")


def extract_text_with_kreuzberg(path):
    """Extract document text through the local Kreuzberg Docker API."""
    suffix = Path(path).suffix.lower()
    supported_suffixes = {
        ".pdf", ".jpg", ".jpeg", ".png", ".bmp", ".tif", ".tiff", ".webp",
        ".txt", ".docx", ".doc", ".rtf", ".odt", ".xlsx", ".xls", ".pptx",
    }
    if suffix not in supported_suffixes:
        return None

    boundary = "----otd421{}".format(uuid.uuid4().hex)
    filename = os.path.basename(path)
    mime_type = mimetypes.guess_type(filename)[0] or "application/octet-stream"
    config = json.dumps({"force_ocr": True, "ocr": {"language": "rus"}})
    try:
        with open(path, "rb") as source_file:
            content = source_file.read()
        body = b"".join((
            "--{}\r\n".format(boundary).encode(),
            'Content-Disposition: form-data; name="files"; filename="{}"\r\n'.format(filename).encode("utf-8"),
            "Content-Type: {}\r\n\r\n".format(mime_type).encode(),
            content,
            "\r\n--{}\r\n".format(boundary).encode(),
            'Content-Disposition: form-data; name="config"\r\n\r\n'.encode(),
            config.encode("utf-8"),
            "\r\n--{}\r\n".format(boundary).encode(),
            'Content-Disposition: form-data; name="output_format"\r\n\r\nplain'.encode(),
            "\r\n--{}--\r\n".format(boundary).encode(),
        ))
        request_to_extract = urllib.request.Request(
            "{}/extract".format(KREUZBERG_URL),
            data=body,
            headers={"Content-Type": "multipart/form-data; boundary={}".format(boundary)},
            method="POST",
        )
        with urllib.request.urlopen(request_to_extract, timeout=180) as response:
            payload = json.loads(response.read().decode("utf-8"))
        results = payload if isinstance(payload, list) else [payload]
        recognized_text = "\n\n".join(
            result.get("content", "").strip()
            for result in results
            if isinstance(result, dict) and result.get("content")
        )
        return recognized_text.strip() or None
    except (OSError, ValueError, urllib.error.URLError, urllib.error.HTTPError, TimeoutError) as exc:
        print("Kreuzberg OCR unavailable: {}".format(exc))
        return None


def extract_text_from_file(path):
    ext = os.path.splitext(path)[-1].lower().strip(".")
    text = ""
    try:
        kreuzberg_text = extract_text_with_kreuzberg(path)
        if kreuzberg_text:
            return kreuzberg_text
        if ext in ("jpg", "jpeg", "png", "bmp", "tif", "tiff", "webp"):
            text = pytesseract.image_to_string(Image.open(path), lang="rus+eng")
        elif ext == "pdf":
            text_parts = []
            with pdfplumber.open(path) as pdf:
                for page in pdf.pages:
                    page_text = page.extract_text()
                    if page_text and page_text.strip():
                        text_parts.append(page_text)
                    else:
                        im = page.to_image(resolution=300).original
                        ocr_text = pytesseract.image_to_string(im, lang="rus+eng")
                        text_parts.append(ocr_text)
            text = "\n".join(text_parts)
        elif ext == "txt":
            with open(path, "r", encoding="utf-8", errors="ignore") as f:
                text = f.read()
        elif ext == "docx":
            doc = docx.Document(path)
            paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            text = "\n".join(paragraphs)
    except Exception as e:
        print(f"text extraction error from {path}: {e}")
        text = ""
    return text.strip()


# OCR is deliberately processed by one background worker.  This keeps the
# Kreuzberg container from receiving several large documents at once.
OCR_QUEUE = queue.Queue()
OCR_QUEUE_LOCK = threading.Lock()
OCR_QUEUED_KEYS = set()
OCR_WORKER_STARTED = False


def process_document_ocr(document_kind, document_id):
    model = ActivityDocument if document_kind == "activity" else LetterDocument
    document = db.session.get(model, document_id)
    if not document or document.ocr_status == "completed":
        return

    document.ocr_status = "processing"
    document.ocr_error = None
    document.ocr_started_at = dt.datetime.utcnow()
    db.session.commit()
    try:
        filepath = os.path.join(app.root_path, "static", document.filepath)
        if not os.path.isfile(filepath):
            raise FileNotFoundError("Загруженный файл не найден на сервере.")
        document.doc_rec = extract_text_from_file(filepath)
        document.ocr_status = "completed"
        if not document.doc_rec:
            document.ocr_error = "Текст в документе не обнаружен."
    except Exception as exc:
        document.ocr_status = "failed"
        document.ocr_error = str(exc)[:1000]
    finally:
        document.ocr_completed_at = dt.datetime.utcnow()
        db.session.commit()


def ocr_worker():
    while True:
        document_kind, document_id = OCR_QUEUE.get()
        try:
            with OCR_QUEUE_LOCK:
                OCR_QUEUED_KEYS.discard((document_kind, document_id))
            with app.app_context():
                process_document_ocr(document_kind, document_id)
        except Exception as exc:
            print("OCR worker error: {}".format(exc))
        finally:
            OCR_QUEUE.task_done()


def start_ocr_worker():
    global OCR_WORKER_STARTED
    with OCR_QUEUE_LOCK:
        if OCR_WORKER_STARTED:
            return
        OCR_WORKER_STARTED = True
        worker = threading.Thread(target=ocr_worker, name="document-ocr-worker", daemon=True)
        worker.start()

    # Resume documents left in the queue by an application restart.
    with app.app_context():
        ActivityDocument.query.filter(ActivityDocument.ocr_status == "processing").update(
            {"ocr_status": "queued"}, synchronize_session=False
        )
        LetterDocument.query.filter(LetterDocument.ocr_status == "processing").update(
            {"ocr_status": "queued"}, synchronize_session=False
        )
        db.session.commit()
        pending_jobs = [
            ("activity", document.id)
            for document in ActivityDocument.query.filter_by(ocr_status="queued").all()
        ] + [
            ("letter", document.id)
            for document in LetterDocument.query.filter_by(ocr_status="queued").all()
        ]
    for document_kind, document_id in pending_jobs:
        enqueue_document_ocr(document_kind, document_id)


def enqueue_document_ocr(document_kind, document_id):
    start_ocr_worker()
    key = (document_kind, document_id)
    with OCR_QUEUE_LOCK:
        if key in OCR_QUEUED_KEYS:
            return
        OCR_QUEUED_KEYS.add(key)
        OCR_QUEUE.put(key)

def safe_from_json(value):
    if not value:
        return []
    if isinstance(value, list):
        return value
    if isinstance(value, str):
        try:
            return json.loads(value)
        except Exception:
            return []
    return []

app.jinja_env.filters['from_json'] = safe_from_json

DB_PATH = os.path.join(BASE_DIR, "app.db")
app.config["SQLALCHEMY_DATABASE_URI"] = f"sqlite:///{DB_PATH}"
db.init_app(app)

ACTIVITY_STATUS_META = {
    "in_progress": {"label": "В работе", "class": "bg-warning text-dark", "style": ""},
    "done": {"label": "Вып.", "class": "bg-success", "style": ""},
    "not_done": {"label": "Не вып.", "class": "bg-danger", "style": ""},
    "postponed": {"label": "Перенос", "class": "text-dark", "style": "background-color:#fd7e14;"},
    "cancelled": {"label": "Отмена", "class": "", "style": "background-color:#6f4e37; color:#fff;"},
    "approved": {"label": "Вып.", "class": "bg-success", "style": ""},
    "rejected": {"label": "Отмена", "class": "", "style": "background-color:#6f4e37; color:#fff;"},
    "open": {"label": "В работе", "class": "bg-warning text-dark", "style": ""},
}

ACTIVITY_STATUS_FILTERS = (
    ("active_and_not_done", "В работе и не выполнено"),
    ("in_progress", "В работе"),
    ("not_done", "Не выполнено"),
    ("postponed", "Перенос"),
    ("done", "Выполнено"),
    ("cancelled", "Отмена"),
    ("all", "Все статусы"),
)

ACTIVITY_PRIORITY_META = {
    1: {"label": "1 - низкий", "class": "bg-secondary"},
    2: {"label": "2 - обычный", "class": "bg-info text-dark"},
    3: {"label": "3 - средний", "class": "bg-primary"},
    4: {"label": "4 - высокий", "class": "bg-warning text-dark"},
    5: {"label": "5 - критический", "class": "bg-danger"},
}

ACTIVITY_COMPLEXITY_META = {
    1: {"label": "1 - простая", "class": "bg-secondary", "default_hours": 1},
    2: {"label": "2 - базовая", "class": "bg-info text-dark", "default_hours": 2},
    3: {"label": "3 - средняя", "class": "bg-primary", "default_hours": 4},
    4: {"label": "4 - высокая", "class": "bg-warning text-dark", "default_hours": 6},
    5: {"label": "5 - критическая", "class": "bg-danger", "default_hours": 8},
}


def get_activity_priority_meta(priority):
    try:
        priority = int(priority)
    except (TypeError, ValueError):
        priority = 3
    return ACTIVITY_PRIORITY_META.get(priority, ACTIVITY_PRIORITY_META[3])


def parse_activity_priority(raw_priority):
    try:
        priority = int(raw_priority)
    except (TypeError, ValueError):
        return None
    return priority if priority in ACTIVITY_PRIORITY_META else None


def parse_activity_complexity(raw_complexity):
    try:
        complexity = int(raw_complexity)
    except (TypeError, ValueError):
        return None
    return complexity if complexity in ACTIVITY_COMPLEXITY_META else None


def get_activity_complexity_meta(complexity_level):
    complexity = parse_activity_complexity(complexity_level) or 3
    base = ACTIVITY_COMPLEXITY_META[complexity]
    return {
        **base,
        "level": complexity,
        "daily_hours": get_app_setting_int(
            "activity_complexity_hours_{}".format(complexity),
            base["default_hours"],
        ),
    }


def calculate_activity_effort_hours(activity):
    """Estimate workload by inclusive weekdays between start and due dates."""
    if not activity or not activity.start_date or not activity.end_date:
        return None
    start_date = activity.start_date.date()
    end_date = activity.end_date.date()
    if end_date < start_date:
        return 0
    working_days = sum(
        1 for offset in range((end_date - start_date).days + 1)
        if (start_date + dt.timedelta(days=offset)).weekday() < 5
    )
    return working_days * get_activity_complexity_meta(activity.complexity_level)["daily_hours"]


def apply_activity_status_filter(query, selected_status):
    """Apply UI status groups while keeping legacy status values visible."""
    if selected_status == "active_and_not_done":
        return query.filter(Activity.status.in_(("open", "in_progress", "not_done")))
    if selected_status == "in_progress":
        return query.filter(Activity.status.in_(("open", "in_progress")))
    if selected_status == "done":
        return query.filter(Activity.status.in_(("done", "approved")))
    if selected_status == "cancelled":
        return query.filter(Activity.status.in_(("cancelled", "rejected")))
    if selected_status in ("not_done", "postponed"):
        return query.filter(Activity.status == selected_status)
    return query


def activity_work_is_locked(activity):
    return activity.status in ("done", "not_done")


def can_modify_activity_work(activity):
    return (
        activity.owner_id == current_user.id
        or current_user.role in ("manager", "admin", "superadmin", "deputy")
        or (
            activity.assigned_department_id is not None
            and activity.assigned_department_id == getattr(current_user, "department_id", None)
        )
    )


def add_activity_history(activity, action, details):
    db.session.add(ActivityHistory(
        activity_id=activity.id,
        author_id=current_user.id if getattr(current_user, "id", -1) != -1 else None,
        action=action,
        details=details,
    ))

def get_activity_status_meta(status: str) -> dict:
    return ACTIVITY_STATUS_META.get(status, {"label": status or "-", "class": "bg-secondary", "style": ""})


@app.context_processor
def inject_template_helpers():
    deps = []
    try:
        deps = Department.query.order_by(Department.name).all()
    except Exception:
        deps = []
    return {
        "get_status_meta": get_activity_status_meta,
        "get_priority_meta": get_activity_priority_meta,
        "get_complexity_meta": get_activity_complexity_meta,
        "calculate_effort_hours": calculate_activity_effort_hours,
        "deps": deps,
    }


app.jinja_env.globals["get_status_meta"] = get_activity_status_meta
app.jinja_env.globals["get_priority_meta"] = get_activity_priority_meta
app.jinja_env.globals["get_complexity_meta"] = get_activity_complexity_meta
app.jinja_env.globals["calculate_effort_hours"] = calculate_activity_effort_hours


@app.get("/manifest.webmanifest")
def pwa_manifest():
    return send_from_directory(app.static_folder, "manifest.webmanifest", mimetype="application/manifest+json")


@app.get("/sw.js")
def service_worker():
    response = send_from_directory(app.static_folder, "sw.js", mimetype="application/javascript")
    response.headers["Service-Worker-Allowed"] = "/"
    response.headers["Cache-Control"] = "no-cache"
    return response

def normalize_activity_statuses():
    # Date fields are stored at midnight, but a deadline remains valid until
    # the end of that calendar day.
    today_start = dt.datetime.combine(dt.date.today(), dt.time.min)
    changed = (
        Activity.query
        .filter(Activity.end_date.isnot(None))
        .filter(Activity.end_date < today_start)
        .filter(Activity.status.in_(["open", "in_progress", "postponed"]))
        .update({"status": "not_done"}, synchronize_session=False)
    )
    if changed:
        db.session.commit()

def ensure_runtime_schema():
    with db.engine.begin() as conn:
        user_cols = {row[1] for row in conn.execute(sql_text("PRAGMA table_info(users)")).fetchall()}
        if user_cols and "is_enabled" not in user_cols:
            conn.execute(sql_text("ALTER TABLE users ADD COLUMN is_enabled BOOLEAN NOT NULL DEFAULT 1"))
        conn.execute(sql_text("""
            CREATE TABLE IF NOT EXISTS user_notifications (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id INTEGER NOT NULL REFERENCES users(id),
                sender_id INTEGER REFERENCES users(id),
                title VARCHAR(180) NOT NULL,
                message TEXT NOT NULL,
                created_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP,
                scheduled_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP,
                background_color VARCHAR(16) NOT NULL DEFAULT '#0d6efd',
                read_at DATETIME
            )
        """))
        notification_cols = {row[1] for row in conn.execute(sql_text("PRAGMA table_info(user_notifications)")).fetchall()}
        if "scheduled_at" not in notification_cols:
            conn.execute(sql_text("ALTER TABLE user_notifications ADD COLUMN scheduled_at DATETIME"))
            conn.execute(sql_text("UPDATE user_notifications SET scheduled_at = created_at WHERE scheduled_at IS NULL"))
        if "background_color" not in notification_cols:
            conn.execute(sql_text("ALTER TABLE user_notifications ADD COLUMN background_color VARCHAR(16) NOT NULL DEFAULT '#0d6efd'"))
        conn.execute(sql_text("""
            CREATE INDEX IF NOT EXISTS ix_user_notifications_user_read
            ON user_notifications (user_id, read_at, created_at)
        """))
        conn.execute(sql_text("""
            CREATE TABLE IF NOT EXISTS app_settings (
                key VARCHAR(100) PRIMARY KEY,
                value TEXT
            )
        """))
        conn.execute(
            sql_text("INSERT OR IGNORE INTO app_settings (key, value) VALUES (:key, :value)"),
            {"key": "memo_prefix", "value": "421"},
        )
        conn.execute(
            sql_text("INSERT OR IGNORE INTO app_settings (key, value) VALUES (:key, :value)"),
            {"key": "letter_prefix", "value": ""},
        )
        conn.execute(
            sql_text("INSERT OR IGNORE INTO app_settings (key, value) VALUES (:key, :value)"),
            {"key": "letter_start_number", "value": "1"},
        )
        conn.execute(
            sql_text("INSERT OR IGNORE INTO app_settings (key, value) VALUES (:key, :value)"),
            {"key": "memo_start_number", "value": "1"},
        )
        for complexity_level, daily_hours in ((1, 1), (2, 2), (3, 4), (4, 6), (5, 8)):
            conn.execute(
                sql_text("INSERT OR IGNORE INTO app_settings (key, value) VALUES (:key, :value)"),
                {
                    "key": "activity_complexity_hours_{}".format(complexity_level),
                    "value": str(daily_hours),
                },
            )
        conn.execute(sql_text("""
            CREATE TABLE IF NOT EXISTS letter_sources (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name VARCHAR(200) NOT NULL UNIQUE,
                bin_code VARCHAR(50),
                address VARCHAR(255),
                phone VARCHAR(100),
                email VARCHAR(150),
                contact_person VARCHAR(150),
                notes TEXT
            )
        """))
        for source_name in ("Входящее", "Исходящее", "Электронная почта"):
            conn.execute(
                sql_text("INSERT OR IGNORE INTO letter_sources (name) VALUES (:name)"),
                {"name": source_name},
            )
        letter_source_cols = {row[1] for row in conn.execute(sql_text("PRAGMA table_info(letter_sources)")).fetchall()}
        if "bin_code" not in letter_source_cols:
            conn.execute(sql_text("ALTER TABLE letter_sources ADD COLUMN bin_code VARCHAR(50)"))
        if "address" not in letter_source_cols:
            conn.execute(sql_text("ALTER TABLE letter_sources ADD COLUMN address VARCHAR(255)"))
        if "phone" not in letter_source_cols:
            conn.execute(sql_text("ALTER TABLE letter_sources ADD COLUMN phone VARCHAR(100)"))
        if "email" not in letter_source_cols:
            conn.execute(sql_text("ALTER TABLE letter_sources ADD COLUMN email VARCHAR(150)"))
        if "contact_person" not in letter_source_cols:
            conn.execute(sql_text("ALTER TABLE letter_sources ADD COLUMN contact_person VARCHAR(150)"))
        if "notes" not in letter_source_cols:
            conn.execute(sql_text("ALTER TABLE letter_sources ADD COLUMN notes TEXT"))

        department_cols = {row[1] for row in conn.execute(sql_text("PRAGMA table_info(departments)")).fetchall()}
        conn.execute(sql_text("""
            CREATE TABLE IF NOT EXISTS workshops (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name VARCHAR(200) NOT NULL UNIQUE,
                manager_id INTEGER REFERENCES users(id),
                created_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP
            )
        """))
        workshop_cols = {row[1] for row in conn.execute(sql_text("PRAGMA table_info(workshops)")).fetchall()}
        if "manager_id" not in workshop_cols:
            conn.execute(sql_text("ALTER TABLE workshops ADD COLUMN manager_id INTEGER"))
        if department_cols:
            if "workshop_id" not in department_cols:
                conn.execute(sql_text("ALTER TABLE departments ADD COLUMN workshop_id INTEGER"))
            if "letter_prefix" not in department_cols:
                conn.execute(sql_text("ALTER TABLE departments ADD COLUMN letter_prefix VARCHAR(100)"))
            if "memo_prefix" not in department_cols:
                conn.execute(sql_text("ALTER TABLE departments ADD COLUMN memo_prefix VARCHAR(100)"))
            if "letter_start_number" not in department_cols:
                conn.execute(sql_text("ALTER TABLE departments ADD COLUMN letter_start_number INTEGER"))
            if "memo_start_number" not in department_cols:
                conn.execute(sql_text("ALTER TABLE departments ADD COLUMN memo_start_number INTEGER"))
            if "letter_number_series" not in department_cols:
                conn.execute(sql_text("ALTER TABLE departments ADD COLUMN letter_number_series VARCHAR(100)"))
            if "memo_number_series" not in department_cols:
                conn.execute(sql_text("ALTER TABLE departments ADD COLUMN memo_number_series VARCHAR(100)"))

        cols = {row[1] for row in conn.execute(sql_text("PRAGMA table_info(activities)")).fetchall()}
        if "outgoing_document" not in cols:
            conn.execute(sql_text("ALTER TABLE activities ADD COLUMN outgoing_document VARCHAR(255)"))
        if "assigned_department_id" not in cols:
            conn.execute(sql_text("ALTER TABLE activities ADD COLUMN assigned_department_id INTEGER"))
        if "status_comment" not in cols:
            conn.execute(sql_text("ALTER TABLE activities ADD COLUMN status_comment TEXT"))
        if "postponed_to" not in cols:
            conn.execute(sql_text("ALTER TABLE activities ADD COLUMN postponed_to DATETIME"))
        if "archived_at" not in cols:
            conn.execute(sql_text("ALTER TABLE activities ADD COLUMN archived_at DATETIME"))
        if "archived_by_id" not in cols:
            conn.execute(sql_text("ALTER TABLE activities ADD COLUMN archived_by_id INTEGER"))
        if "completion_memo_id" not in cols:
            conn.execute(sql_text("ALTER TABLE activities ADD COLUMN completion_memo_id INTEGER"))
        if "created_by_id" not in cols:
            conn.execute(sql_text("ALTER TABLE activities ADD COLUMN created_by_id INTEGER"))
        if "plan_requested" not in cols:
            conn.execute(sql_text("ALTER TABLE activities ADD COLUMN plan_requested BOOLEAN NOT NULL DEFAULT 0"))
        if "complexity_level" not in cols:
            conn.execute(sql_text("ALTER TABLE activities ADD COLUMN complexity_level INTEGER NOT NULL DEFAULT 3"))
        conn.execute(sql_text("""
            CREATE TABLE IF NOT EXISTS protocols (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                title VARCHAR(255) NOT NULL,
                protocol_date DATE NOT NULL,
                description TEXT,
                created_by_id INTEGER REFERENCES users(id),
                created_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP
            )
        """))
        conn.execute(sql_text("""
            CREATE TABLE IF NOT EXISTS protocol_items (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                protocol_id INTEGER NOT NULL REFERENCES protocols(id) ON DELETE CASCADE,
                activity_id INTEGER NOT NULL REFERENCES activities(id),
                position INTEGER NOT NULL DEFAULT 1,
                due_date DATE,
                comment TEXT,
                created_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP,
                UNIQUE(protocol_id, activity_id)
            )
        """))
        conn.execute(sql_text("""
            CREATE INDEX IF NOT EXISTS ix_protocol_items_protocol_position
            ON protocol_items (protocol_id, position)
        """))
        conn.execute(sql_text("""
            CREATE TABLE IF NOT EXISTS schedule_entries (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                schedule_date DATE NOT NULL,
                text TEXT NOT NULL,
                responsible_id INTEGER NOT NULL REFERENCES users(id),
                status VARCHAR(32) NOT NULL DEFAULT 'planned',
                completion_comment TEXT,
                completed_at DATETIME,
                activity_id INTEGER REFERENCES activities(id),
                memo_id INTEGER REFERENCES letters(id),
                created_by_id INTEGER REFERENCES users(id),
                created_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP,
                updated_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP
            )
        """))
        conn.execute(sql_text("""
            CREATE INDEX IF NOT EXISTS ix_schedule_entries_date_responsible
            ON schedule_entries (schedule_date, responsible_id)
        """))
        conn.execute(sql_text("""
            CREATE TABLE IF NOT EXISTS schedules (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                title VARCHAR(255) NOT NULL,
                schedule_date DATE NOT NULL,
                description TEXT,
                created_by_id INTEGER REFERENCES users(id),
                created_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP
            )
        """))
        conn.execute(sql_text("""
            CREATE TABLE IF NOT EXISTS schedule_items (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                schedule_id INTEGER NOT NULL REFERENCES schedules(id) ON DELETE CASCADE,
                position INTEGER NOT NULL DEFAULT 1,
                item_date DATE NOT NULL,
                text TEXT NOT NULL,
                responsible_id INTEGER NOT NULL REFERENCES users(id),
                status VARCHAR(32) NOT NULL DEFAULT 'planned',
                completion_comment TEXT,
                completed_at DATETIME,
                activity_id INTEGER REFERENCES activities(id),
                memo_id INTEGER REFERENCES letters(id),
                created_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP,
                updated_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP
            )
        """))
        conn.execute(sql_text("""
            CREATE INDEX IF NOT EXISTS ix_schedule_items_schedule_position
            ON schedule_items (schedule_id, position)
        """))
        document_cols = {row[1] for row in conn.execute(sql_text("PRAGMA table_info(activity_documents)")).fetchall()}
        if "uploaded_by_id" not in document_cols:
            conn.execute(sql_text("ALTER TABLE activity_documents ADD COLUMN uploaded_by_id INTEGER"))
        if "ocr_status" not in document_cols:
            conn.execute(sql_text("ALTER TABLE activity_documents ADD COLUMN ocr_status VARCHAR(20) NOT NULL DEFAULT 'completed'"))
        if "ocr_error" not in document_cols:
            conn.execute(sql_text("ALTER TABLE activity_documents ADD COLUMN ocr_error TEXT"))
        if "ocr_started_at" not in document_cols:
            conn.execute(sql_text("ALTER TABLE activity_documents ADD COLUMN ocr_started_at DATETIME"))
        if "ocr_completed_at" not in document_cols:
            conn.execute(sql_text("ALTER TABLE activity_documents ADD COLUMN ocr_completed_at DATETIME"))
        letter_document_cols = {row[1] for row in conn.execute(sql_text("PRAGMA table_info(letter_documents)")).fetchall()}
        if "ocr_status" not in letter_document_cols:
            conn.execute(sql_text("ALTER TABLE letter_documents ADD COLUMN ocr_status VARCHAR(20) NOT NULL DEFAULT 'completed'"))
        if "ocr_error" not in letter_document_cols:
            conn.execute(sql_text("ALTER TABLE letter_documents ADD COLUMN ocr_error TEXT"))
        if "ocr_started_at" not in letter_document_cols:
            conn.execute(sql_text("ALTER TABLE letter_documents ADD COLUMN ocr_started_at DATETIME"))
        if "ocr_completed_at" not in letter_document_cols:
            conn.execute(sql_text("ALTER TABLE letter_documents ADD COLUMN ocr_completed_at DATETIME"))
        conn.execute(sql_text("""CREATE TABLE IF NOT EXISTS activity_history (
            id INTEGER PRIMARY KEY AUTOINCREMENT, activity_id INTEGER NOT NULL REFERENCES activities(id),
            author_id INTEGER REFERENCES users(id), action VARCHAR(100) NOT NULL, details TEXT NOT NULL,
            created_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP)"""))
        conn.execute(sql_text("""
            UPDATE activities
            SET created_by_id = (
                SELECT h.author_id FROM activity_history h
                WHERE h.activity_id = activities.id
                  AND h.action = 'Создание задачи'
                  AND h.author_id IS NOT NULL
                ORDER BY h.created_at ASC
                LIMIT 1
            )
            WHERE created_by_id IS NULL
        """))
        conn.execute(sql_text("""
            CREATE TABLE IF NOT EXISTS activity_templates (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name VARCHAR(200) NOT NULL UNIQUE,
                title VARCHAR(255) NOT NULL,
                description TEXT,
                type_id INTEGER REFERENCES activity_types(id),
                direction VARCHAR(255),
                outgoing_document VARCHAR(255),
                duration_days INTEGER,
                created_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP
            )
        """))
        conn.execute(sql_text("""
            CREATE TABLE IF NOT EXISTS memo_templates (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name VARCHAR(200) NOT NULL UNIQUE,
                subject VARCHAR(255) NOT NULL,
                body TEXT,
                department_id INTEGER REFERENCES departments(id),
                executor_id INTEGER REFERENCES users(id),
                created_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP
            )
        """))
        # Correct titles made before memo-aware task creation was introduced.
        conn.execute(sql_text("""
            UPDATE activities
            SET title = 'Ответ на служебную записку №' || COALESCE(
                (SELECT reg_number FROM letters WHERE letters.id = activities.source_letter_id),
                activities.source_letter_id
            )
            WHERE title LIKE 'Ответ на письмо №%'
              AND source_letter_id IN (
                  SELECT l.id FROM letters l
                  JOIN letter_sources s ON s.id = l.source_id
                  WHERE s.name = 'Служебная записка'
              )
        """))

        letter_cols = {row[1] for row in conn.execute(sql_text("PRAGMA table_info(letters)")).fetchall()}
        if "reg_number" not in letter_cols:
            conn.execute(sql_text("ALTER TABLE letters ADD COLUMN reg_number VARCHAR(100)"))
        if "source_id" not in letter_cols:
            conn.execute(sql_text("ALTER TABLE letters ADD COLUMN source_id INTEGER"))
        if "executor_id" not in letter_cols:
            conn.execute(sql_text("ALTER TABLE letters ADD COLUMN executor_id INTEGER"))
        if "archived_at" not in letter_cols:
            conn.execute(sql_text("ALTER TABLE letters ADD COLUMN archived_at DATETIME"))
        if "archived_by_id" not in letter_cols:
            conn.execute(sql_text("ALTER TABLE letters ADD COLUMN archived_by_id INTEGER"))
        if "approval_status" not in letter_cols:
            conn.execute(sql_text("ALTER TABLE letters ADD COLUMN approval_status VARCHAR(20) NOT NULL DEFAULT 'pending'"))
        if "approved_by_id" not in letter_cols:
            conn.execute(sql_text("ALTER TABLE letters ADD COLUMN approved_by_id INTEGER"))
        if "approved_at" not in letter_cols:
            conn.execute(sql_text("ALTER TABLE letters ADD COLUMN approved_at DATETIME"))

        shared_link_cols = {row[1] for row in conn.execute(sql_text("PRAGMA table_info(shared_links)")).fetchall()}
        if shared_link_cols and "payload" not in shared_link_cols:
            conn.execute(sql_text("ALTER TABLE shared_links ADD COLUMN payload JSON"))

        plan_cols = {row[1] for row in conn.execute(sql_text("PRAGMA table_info(plans)")).fetchall()}
        if "approval_status" not in plan_cols:
            conn.execute(sql_text("ALTER TABLE plans ADD COLUMN approval_status VARCHAR(32) NOT NULL DEFAULT 'approved'"))
        if "approved_by_id" not in plan_cols:
            conn.execute(sql_text("ALTER TABLE plans ADD COLUMN approved_by_id INTEGER"))
        if "approved_at" not in plan_cols:
            conn.execute(sql_text("ALTER TABLE plans ADD COLUMN approved_at DATETIME"))
        if "revision_comment" not in plan_cols:
            conn.execute(sql_text("ALTER TABLE plans ADD COLUMN revision_comment TEXT"))

        # Plan rows are stored separately from the legacy free-text plans so old
        # plans remain available while new monthly plans get measurable results.
        conn.execute(sql_text("""
            CREATE TABLE IF NOT EXISTS plan_items (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                plan_id INTEGER NOT NULL REFERENCES plans(id) ON DELETE CASCADE,
                executor_id INTEGER NOT NULL REFERENCES users(id),
                position INTEGER NOT NULL DEFAULT 1,
                task_text TEXT NOT NULL,
                justification TEXT,
                deadline_kind VARCHAR(32) NOT NULL DEFAULT 'month',
                planned_hours REAL NOT NULL DEFAULT 0,
                status VARCHAR(32) NOT NULL DEFAULT 'planned'
                    CHECK (status IN ('planned', 'done', 'not_done', 'cancelled')),
                comment TEXT,
                completed_at DATETIME,
                created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
                updated_at DATETIME DEFAULT CURRENT_TIMESTAMP
            )
        """))
        conn.execute(sql_text("""
            CREATE INDEX IF NOT EXISTS ix_plan_items_plan_executor
            ON plan_items (plan_id, executor_id)
        """))
        plan_item_cols = {row[1] for row in conn.execute(sql_text("PRAGMA table_info(plan_items)")).fetchall()}
        if "execution_approved" not in plan_item_cols:
            conn.execute(sql_text("ALTER TABLE plan_items ADD COLUMN execution_approved BOOLEAN NOT NULL DEFAULT 0"))
        if "execution_approved_by_id" not in plan_item_cols:
            conn.execute(sql_text("ALTER TABLE plan_items ADD COLUMN execution_approved_by_id INTEGER"))
        if "execution_approved_at" not in plan_item_cols:
            conn.execute(sql_text("ALTER TABLE plan_items ADD COLUMN execution_approved_at DATETIME"))
        if "deadline_date" not in plan_item_cols:
            conn.execute(sql_text("ALTER TABLE plan_items ADD COLUMN deadline_date DATE"))
        if "linked_activity_id" not in plan_item_cols:
            conn.execute(sql_text("ALTER TABLE plan_items ADD COLUMN linked_activity_id INTEGER REFERENCES activities(id)"))
        if "planned_hours" not in plan_item_cols:
            conn.execute(sql_text("ALTER TABLE plan_items ADD COLUMN planned_hours REAL NOT NULL DEFAULT 0"))
        if "justification" not in plan_item_cols:
            conn.execute(sql_text("ALTER TABLE plan_items ADD COLUMN justification TEXT"))

        activity_cols = {row[1] for row in conn.execute(sql_text("PRAGMA table_info(activities)")).fetchall()}
        if "priority" not in activity_cols:
            conn.execute(sql_text("ALTER TABLE activities ADD COLUMN priority INTEGER NOT NULL DEFAULT 3"))

        create_sql_row = conn.execute(
            sql_text("SELECT sql FROM sqlite_master WHERE type='table' AND name='activities'")
        ).fetchone()
        create_sql = create_sql_row[0] if create_sql_row else ""
        if create_sql and "not_done" not in create_sql:
            conn.execute(sql_text("PRAGMA foreign_keys=OFF"))
            conn.execute(sql_text("ALTER TABLE activities RENAME TO activities__old_status_migration"))
            conn.execute(sql_text("""
                CREATE TABLE activities (
                  id INTEGER PRIMARY KEY AUTOINCREMENT,
                  type_id INTEGER NOT NULL,
                  owner_id INTEGER NOT NULL,
                  created_by_id INTEGER,
                  title TEXT NOT NULL,
                  description TEXT,
                  start_date TEXT NOT NULL DEFAULT (datetime('now')),
                  end_date TEXT,
                  status TEXT NOT NULL DEFAULT 'open'
                    CHECK (status IN ('open','in_progress','done','approved','rejected','not_done','postponed','cancelled')),
                  priority INTEGER NOT NULL DEFAULT 3,
                  approved_by_id INTEGER,
                  approved_at TEXT,
                  approve_comment TEXT,
                  source_letter_id INTEGER,
                  completion_memo_id INTEGER REFERENCES letters(id),
                  direction TEXT,
                  decision TEXT,
                  status_comment TEXT,
                  postponed_to TEXT,
                  extra_values TEXT,
                  plan_id INTEGER REFERENCES plans(id),
                  outgoing_document VARCHAR(255),
                  assigned_department_id INTEGER REFERENCES departments(id),
                  FOREIGN KEY (type_id) REFERENCES activity_types(id) ON UPDATE CASCADE ON DELETE RESTRICT,
                  FOREIGN KEY (owner_id) REFERENCES users(id) ON UPDATE CASCADE ON DELETE RESTRICT,
                  FOREIGN KEY (approved_by_id) REFERENCES users(id) ON UPDATE CASCADE ON DELETE SET NULL
                )
            """))
            conn.execute(sql_text("""
                INSERT INTO activities (
                  id, type_id, owner_id, created_by_id, title, description, start_date, end_date, status, priority,
                  approved_by_id, approved_at, approve_comment, source_letter_id, completion_memo_id, direction,
                   decision, status_comment, postponed_to, extra_values, plan_id, outgoing_document, assigned_department_id
                )
                SELECT
                  id, type_id, owner_id, created_by_id, title, description, start_date, end_date, status, COALESCE(priority, 3),
                  approved_by_id, approved_at, approve_comment, source_letter_id, completion_memo_id, direction,
                   decision, status_comment, postponed_to, extra_values, plan_id, outgoing_document, assigned_department_id
                FROM activities__old_status_migration
            """))
            conn.execute(sql_text("DROP TABLE activities__old_status_migration"))
            conn.execute(sql_text("CREATE INDEX IF NOT EXISTS idx_activities_type ON activities(type_id)"))
            conn.execute(sql_text("CREATE INDEX IF NOT EXISTS idx_activities_owner ON activities(owner_id)"))
            conn.execute(sql_text("CREATE INDEX IF NOT EXISTS idx_activities_status ON activities(status)"))
            conn.execute(sql_text("CREATE INDEX IF NOT EXISTS idx_activities_start_date ON activities(start_date)"))
            conn.execute(sql_text("PRAGMA foreign_keys=ON"))

def ensure_required_activity_types():
    needed = ["Задача", "Ответ на письмо", "написать письмо", "написать с/з", "ответ на с/з"]
    for name in needed:
        exists = ActivityType.query.filter(ActivityType.name == name).first()
        if not exists:
            db.session.add(ActivityType(name=name))
    db.session.commit()


def cleanup_broken_activity_types():
    canonical_pairs = [
        ("Задача", [1001, 1006]),
        ("Ответ на письмо", [1002, 1007]),
        ("написать письмо", [1003]),
        ("написать с/з", [1004]),
        ("ответ на с/з", [1005]),
    ]
    for canonical_name, broken_ids in canonical_pairs:
        canonical = ActivityType.query.filter(ActivityType.name == canonical_name).first()
        if not canonical:
            continue
        for broken_id in broken_ids:
            broken = ActivityType.query.get(broken_id)
            if broken and canonical.id != broken.id:
                Activity.query.filter(Activity.type_id == broken.id).update(
                    {"type_id": canonical.id},
                    synchronize_session=False,
                )
                db.session.delete(broken)
    db.session.commit()

with app.app_context():
    ensure_runtime_schema()
    ensure_required_activity_types()
    cleanup_broken_activity_types()

login_manager = LoginManager(app)
login_manager.login_view = "login_page"

@login_manager.user_loader
def load_user(user_id):
    if str(user_id) == "-1":
        return SuperAdmin()
    return User.query.get(int(user_id))


@app.before_request
def reject_disabled_user_session():
    """End an already open session after an administrator disables the account."""
    if (
        current_user.is_authenticated
        and getattr(current_user, "id", None) != -1
        and not current_user.is_active
    ):
        logout_user()
        flash("Учётная запись отключена администратором.", "warning")
        return redirect(url_for("login_page"))


REPORT_SHARE_ENDPOINTS = {
    "activities_analytics": "activities",
    "plans_statistics": "plans_statistics",
    "plans_monthly_report": "plans_monthly_report",
}
REPORT_SHARE_OPTIONS = {
    "activities": {"endpoint": "activities_analytics", "label": "Аналитика задач"},
    "plans_statistics": {"endpoint": "plans_statistics", "label": "Аналитика выполнения планов"},
    "plans_monthly_report": {"endpoint": "plans_monthly_report", "label": "Сводный отчет по планам"},
}


def get_active_report_share(endpoint_name):
    report_key = REPORT_SHARE_ENDPOINTS.get(endpoint_name)
    token = (request.args.get("share") or "").strip()
    if not report_key or not token:
        return None
    link = SharedLink.query.filter_by(
        token=token,
        doc_type="report:{}".format(report_key),
        doc_id=0,
    ).first()
    if not link or (link.expires_at and link.expires_at < dt.datetime.utcnow()):
        return None
    return link


def get_report_share_payload(endpoint_name):
    link = get_active_report_share(endpoint_name)
    if not link:
        return None
    payload = link.payload
    if isinstance(payload, str):
        try:
            payload = json.loads(payload)
        except (TypeError, ValueError):
            payload = {}
    return payload if isinstance(payload, dict) else {}


def get_report_workshop_id(shared_payload):
    """Read a valid workshop restriction stored in an external report link."""
    payload = shared_payload or {}
    if isinstance(payload, str):
        try:
            payload = json.loads(payload)
        except (TypeError, ValueError):
            payload = {}
    if not isinstance(payload, dict):
        return None
    raw_workshop_id = payload.get("workshop_id")
    try:
        workshop_id = int(raw_workshop_id) if raw_workshop_id else None
    except (TypeError, ValueError):
        return None
    return workshop_id if workshop_id and db.session.get(Workshop, workshop_id) else None


def role_required(*roles):
    def _decor(fn):
        @wraps(fn)
        def wrapper(*args, **kwargs):
            if get_active_report_share(fn.__name__):
                return fn(*args, **kwargs)
            if not current_user.is_authenticated:
                flash("Требуется вход в систему.", "warning")
                return redirect(url_for("login_page"))
            if current_user.role not in roles:
                abort(403)
            return fn(*args, **kwargs)
        return wrapper
    return _decor


def crm_api_required(fn):
    """Authenticate server-to-server CRM requests with the administrator token."""
    @wraps(fn)
    def wrapper(*args, **kwargs):
        expected_token = get_app_setting("integration_api_token", "")
        authorization = (request.headers.get("Authorization") or "").strip()
        supplied_token = request.headers.get("X-API-Key", "").strip()
        if authorization.lower().startswith("bearer "):
            supplied_token = authorization[7:].strip()
        if not expected_token or not supplied_token or not hmac.compare_digest(expected_token, supplied_token):
            return jsonify({"error": "unauthorized", "message": "Передайте действующий API-токен."}), 401
        return fn(*args, **kwargs)
    return wrapper


def crm_api_pagination():
    try:
        limit = int(request.args.get("limit", 100))
        offset = int(request.args.get("offset", 0))
    except ValueError:
        return None, None
    return min(max(limit, 1), 500), max(offset, 0)


def crm_api_date(value):
    if not value:
        return None
    try:
        return dt.datetime.strptime(value, "%Y-%m-%d").date()
    except ValueError:
        return False


def crm_api_user(user):
    if not user:
        return None
    return {
        "id": user.id,
        "full_name": user.full_name,
        "department": user.department.name if user.department else None,
    }


def crm_api_json_body():
    """Accept only JSON objects so integrations cannot accidentally submit form data."""
    payload = request.get_json(silent=True)
    return payload if isinstance(payload, dict) else None


def crm_api_optional_text(value):
    if value is None:
        return None
    if not isinstance(value, str):
        return False
    return value.strip() or None


def crm_api_datetime(value):
    if value in (None, ""):
        return None
    if not isinstance(value, str):
        return False
    try:
        parsed = dt.datetime.fromisoformat(value.replace("Z", "+00:00"))
        if parsed.tzinfo:
            parsed = parsed.astimezone(dt.timezone.utc).replace(tzinfo=None)
        return parsed
    except ValueError:
        return False


def crm_api_active_user(user_id):
    if not isinstance(user_id, int):
        return None
    user = db.session.get(User, user_id)
    return user if user and user.is_approved and user.is_enabled else None


def crm_api_activity_payload(activity):
    return {
        "id": activity.id,
        "title": activity.title,
        "description": activity.description,
        "type": {"id": activity.type_id, "name": activity.type.name if activity.type else None},
        "status": activity.status,
        "status_label": get_activity_status_meta(activity.status)["label"],
        "priority": activity.priority or 3,
        "complexity_level": activity.complexity_level or 3,
        "start_date": iso_or_none(activity.start_date),
        "end_date": iso_or_none(activity.end_date),
        "postponed_to": iso_or_none(activity.postponed_to),
        "status_comment": activity.status_comment,
        "direction": activity.direction,
        "outgoing_document": activity.outgoing_document,
        "owner": crm_api_user(activity.owner),
        "assigned_department": {
            "id": activity.assigned_department.id,
            "name": activity.assigned_department.name,
        } if activity.assigned_department else None,
        "plan_id": activity.plan_id,
        "source_document": {
            "id": activity.source_letter.id,
            "number": activity.source_letter.reg_number,
            "subject": activity.source_letter.subject,
        } if activity.source_letter else None,
        "completion_memo": {
            "id": activity.completion_memo.id,
            "number": activity.completion_memo.reg_number,
            "subject": activity.completion_memo.subject,
        } if activity.completion_memo else None,
    }


def crm_api_activity_full_payload(activity):
    """Return the fullest safe integration representation of one activity."""
    payload = crm_api_activity_payload(activity)

    logs = ActivityLog.query.filter_by(activity_id=activity.id).order_by(
        ActivityLog.entry_date.asc(), ActivityLog.id.asc()
    ).all()
    payload["logs"] = [{
        "id": row.id,
        "parent_id": row.parent_id,
        "text": row.text,
        "entry_date": iso_or_none(row.entry_date),
        "is_done": bool(row.is_done),
    } for row in logs]

    history = ActivityHistory.query.filter_by(activity_id=activity.id).options(
        joinedload(ActivityHistory.author)
    ).order_by(ActivityHistory.created_at.desc(), ActivityHistory.id.desc()).all()
    payload["history"] = [{
        "id": row.id,
        "action": row.action,
        "details": row.details,
        "created_at": iso_or_none(row.created_at),
        "author": crm_api_user(row.author),
    } for row in history]

    documents = ActivityDocument.query.filter_by(activity_id=activity.id).order_by(
        ActivityDocument.uploaded_at.desc(), ActivityDocument.id.desc()
    ).all()
    payload["documents"] = [{
        "id": row.id,
        "filename": row.filename,
        "uploaded_by_id": row.uploaded_by_id,
        "uploaded_at": iso_or_none(row.uploaded_at),
        "ocr_status": row.ocr_status,
        "ocr_error": row.ocr_error,
        "ocr_started_at": iso_or_none(row.ocr_started_at),
        "ocr_completed_at": iso_or_none(row.ocr_completed_at),
        "recognized_text": row.doc_rec,
    } for row in documents]

    extras = activity.extra_values
    if isinstance(extras, str):
        try:
            extras = json.loads(extras)
        except (TypeError, ValueError):
            extras = []
    payload["extra_values"] = extras if isinstance(extras, list) else []

    payload["linked_letters"] = [{
        "id": letter.id,
        "number": letter.reg_number,
        "subject": letter.subject,
        "date": iso_or_none(letter.letter_date),
        "approval_status": letter.approval_status,
        "author": crm_api_user(letter.author),
        "executor": crm_api_user(letter.executor),
    } for letter in activity.letters]

    payload["created_by"] = crm_api_user(activity.created_by)
    payload["approved_by"] = crm_api_user(activity.approved_by)
    payload["approved_at"] = iso_or_none(activity.approved_at)
    payload["approve_comment"] = activity.approve_comment
    payload["decision"] = activity.decision
    payload["archived_at"] = iso_or_none(activity.archived_at)
    payload["archived_by_id"] = activity.archived_by_id
    payload["plan_requested"] = bool(activity.plan_requested)

    if activity.plan_id:
        plan = Plan.query.options(
            joinedload(Plan.creator),
            joinedload(Plan.items).joinedload(PlanItem.executor).joinedload(User.department),
            joinedload(Plan.items).joinedload(PlanItem.linked_activity),
        ).filter_by(id=activity.plan_id).first()
        payload["plan"] = crm_api_plan_payload(plan) if plan else None
    else:
        payload["plan"] = None

    return payload


def crm_api_plan_payload(plan):
    return {
        "id": plan.id,
        "start_date": iso_or_none(plan.start_date),
        "end_date": iso_or_none(plan.end_date),
        "text": plan.plan_text,
        "approval_status": plan.approval_status,
        "revision_comment": plan.revision_comment,
        "created_at": iso_or_none(plan.created_at),
        "creator": crm_api_user(plan.creator),
        "items": [{
            "id": item.id,
            "position": item.position,
            "text": item.task_text,
            "justification": item.justification,
            "deadline_kind": item.deadline_kind,
            "deadline_date": iso_or_none(item.deadline_date),
            "planned_hours": item.planned_hours or 0,
            "status": item.status,
            "comment": item.comment,
            "completed_at": iso_or_none(item.completed_at),
            "executor": crm_api_user(item.executor),
            "linked_task": {
                "id": item.linked_activity.id,
                "title": item.linked_activity.title,
            } if item.linked_activity else None,
        } for item in plan.items],
    }


def crm_api_plan_reporting_payload(month_value):
    """Build the report and analytics payload from one selected calendar month."""
    period = get_plan_period(month_value)
    if not period:
        return None
    period_start, period_end = period
    plans_count = Plan.query.filter(
        Plan.start_date <= period_end,
        Plan.end_date >= period_start,
    ).count()
    items = (
        PlanItem.query.options(
            joinedload(PlanItem.executor).joinedload(User.department),
            joinedload(PlanItem.linked_activity),
        )
        .join(Plan)
        .join(User, PlanItem.executor_id == User.id)
        .filter(Plan.start_date <= period_end, Plan.end_date >= period_start)
        .order_by(User.full_name, PlanItem.position)
        .all()
    )
    department_rows, executor_rows = build_plan_department_rows(items)
    totals = {status: sum(1 for item in items if item.status == status) for status in PLAN_ITEM_STATUSES}
    total = len(items)
    summary = {
        "plans_count": plans_count,
        "departments_count": len(department_rows),
        "executors_count": len(executor_rows),
        "items_total": total,
        "planned_count": totals["planned"],
        "done_count": totals["done"],
        "not_done_count": totals["not_done"],
        "cancelled_count": totals["cancelled"],
        "completion_percent": round((totals["done"] / total) * 100, 1) if total else 0,
        "planned_hours_total": round(sum(float(item.planned_hours or 0) for item in items), 2),
        "done_hours_total": round(sum(float(item.planned_hours or 0) for item in items if item.status == "done"), 2),
        "not_done_hours_total": round(sum(float(item.planned_hours or 0) for item in items if item.status == "not_done"), 2),
    }

    def aggregate(row):
        return {
            "total": row["total"], "planned": row["planned"], "done": row["done"],
            "not_done": row["not_done"], "cancelled": row["cancelled"],
            "completion_percent": row["completion"],
            "planned_hours": round(row["planned_hours"], 2),
            "done_hours": round(row["done_hours"], 2),
            "not_done_hours": round(row["not_done_hours"], 2),
        }

    return {
        "period": {
            "month": period_start.strftime("%Y-%m"),
            "start_date": period_start.isoformat(),
            "end_date": period_end.isoformat(),
        },
        "summary": summary,
        "departments": [
            {"id": row["id"], "name": row["name"], **aggregate(row)}
            for row in department_rows
        ],
        "executors": [
            {
                "id": row["user"].id if row["user"] else None,
                "full_name": row["user"].full_name if row["user"] else "Без исполнителя",
                "department": {
                    "id": row["user"].department.id,
                    "name": row["user"].department.name,
                } if row["user"] and row["user"].department else None,
                **aggregate(row),
            }
            for row in executor_rows
        ],
        "items": [{
            "id": item.id,
            "plan_id": item.plan_id,
            "position": item.position,
            "text": item.task_text,
            "justification": item.justification,
            "deadline_kind": item.deadline_kind,
            "deadline_date": iso_or_none(item.deadline_date),
            "planned_hours": item.planned_hours or 0,
            "status": item.status,
            "comment": item.comment,
            "completed_at": iso_or_none(item.completed_at),
            "executor": crm_api_user(item.executor),
            "linked_task": {
                "id": item.linked_activity.id,
                "title": item.linked_activity.title,
            } if item.linked_activity else None,
        } for item in items],
    }


def crm_api_validate_plan_item(payload, plan, position, existing=None):
    """Apply validated external plan-line values and return an error message if invalid."""
    if not isinstance(payload, dict):
        return "Каждый пункт плана должен быть JSON-объектом."

    task_text = payload.get("text", existing.task_text if existing else "")
    task_text = task_text.strip() if isinstance(task_text, str) else ""
    if not task_text:
        return "Для пункта плана укажите text."

    justification = payload.get("justification", existing.justification if existing else None)
    justification = crm_api_optional_text(justification)
    if justification is False:
        return "justification пункта плана должен быть строкой или null."

    executor_id = payload.get("executor_id", existing.executor_id if existing else None)
    executor = crm_api_active_user(executor_id)
    if not executor:
        return "Укажите существующего активного исполнителя пункта плана."

    deadline_kind = payload.get("deadline_kind", existing.deadline_kind if existing else "month")
    if deadline_kind not in PLAN_DEADLINES:
        return "Некорректный deadline_kind пункта плана."
    deadline_date = crm_api_date(payload.get("deadline_date")) if "deadline_date" in payload else (existing.deadline_date if existing else None)
    if deadline_date is False or (deadline_kind == "date" and not deadline_date):
        return "Для срока date укажите deadline_date в формате YYYY-MM-DD."

    planned_hours = parse_plan_hours(payload.get("planned_hours", existing.planned_hours if existing else None))
    if planned_hours is None:
        return "planned_hours должен быть числом от 0 до 10000; пустое значение считается 0."

    status = payload.get("status", existing.status if existing else "planned")
    if status not in PLAN_ITEM_STATUSES:
        return "Некорректный статус пункта плана."

    linked_activity_id = payload.get("linked_task_id", existing.linked_activity_id if existing else None)
    linked_activity = None
    if linked_activity_id is not None:
        linked_activity = db.session.get(Activity, linked_activity_id) if isinstance(linked_activity_id, int) else None
        if not linked_activity or linked_activity.archived_at or linked_activity.owner_id != executor.id:
            return "Связать можно только с действующей задачей указанного исполнителя."

    item = existing or PlanItem(plan_id=plan.id, position=position)
    item.task_text = task_text
    item.justification = justification
    item.executor_id = executor.id
    item.deadline_kind = deadline_kind
    item.deadline_date = deadline_date
    item.planned_hours = planned_hours
    item.status = status
    if "comment" in payload:
        comment = crm_api_optional_text(payload["comment"])
        if comment is False:
            return "comment пункта плана должен быть строкой или null."
        item.comment = comment
    item.linked_activity_id = linked_activity.id if linked_activity else None
    if status == "done" and not item.completed_at:
        item.completed_at = dt.datetime.utcnow()
    elif status != "done":
        item.completed_at = None
    return item


@app.get("/api/crm/v1/health")
@crm_api_required
def crm_api_health():
    return jsonify({"status": "ok", "api_version": "v1", "server_time": dt.datetime.utcnow().isoformat() + "Z"})


@app.get("/api/crm/v1/users")
@crm_api_required
def crm_api_users():
    """Directory for reliable CRM-to-document-workflow employee mapping."""
    limit, offset = crm_api_pagination()
    if limit is None:
        return jsonify({"error": "invalid_pagination"}), 400
    include_inactive = (request.args.get("include_inactive") or "").lower() in {"1", "true", "yes"}
    query = User.query.options(joinedload(User.department).joinedload(Department.workshop))
    if not include_inactive:
        query = query.filter(User.is_approved.is_(True), User.is_enabled.is_(True))
    total = query.count()
    rows = query.order_by(User.full_name, User.id).offset(offset).limit(limit).all()
    return jsonify({
        "data": [{
            "id": user.id,
            "full_name": user.full_name,
            "username": user.username,
            "role": user.role,
            "is_approved": bool(user.is_approved),
            "is_enabled": bool(user.is_enabled),
            "department": {
                "id": user.department.id,
                "name": user.department.name,
                "workshop": {
                    "id": user.department.workshop.id,
                    "name": user.department.workshop.name,
                } if user.department.workshop else None,
            } if user.department else None,
        } for user in rows],
        "pagination": {"total": total, "limit": limit, "offset": offset},
    })


@app.get("/api/crm/v1/departments")
@crm_api_required
def crm_api_departments():
    """Department directory, including numbering settings used by document registers."""
    limit, offset = crm_api_pagination()
    if limit is None:
        return jsonify({"error": "invalid_pagination"}), 400
    query = Department.query.options(joinedload(Department.workshop))
    total = query.count()
    rows = query.order_by(Department.name, Department.id).offset(offset).limit(limit).all()
    return jsonify({
        "data": [{
            "id": department.id,
            "name": department.name,
            "workshop": {
                "id": department.workshop.id,
                "name": department.workshop.name,
            } if department.workshop else None,
            "letter_prefix": department.letter_prefix,
            "memo_prefix": department.memo_prefix,
            "letter_number_series": department.letter_number_series,
            "memo_number_series": department.memo_number_series,
        } for department in rows],
        "pagination": {"total": total, "limit": limit, "offset": offset},
    })


@app.get("/api/crm/v1/workshops")
@crm_api_required
def crm_api_workshops():
    """Workshop directory with the appointed manager, if one is configured."""
    limit, offset = crm_api_pagination()
    if limit is None:
        return jsonify({"error": "invalid_pagination"}), 400
    query = Workshop.query.options(joinedload(Workshop.manager))
    total = query.count()
    rows = query.order_by(Workshop.name, Workshop.id).offset(offset).limit(limit).all()
    return jsonify({
        "data": [{
            "id": workshop.id,
            "name": workshop.name,
            "manager": {
                "id": workshop.manager.id,
                "full_name": workshop.manager.full_name,
            } if workshop.manager else None,
        } for workshop in rows],
        "pagination": {"total": total, "limit": limit, "offset": offset},
    })


@app.get("/api/crm/v1/task-types")
@crm_api_required
def crm_api_task_types():
    """Task-type directory used by POST /tasks type_id validation."""
    limit, offset = crm_api_pagination()
    if limit is None:
        return jsonify({"error": "invalid_pagination"}), 400
    query = ActivityType.query
    total = query.count()
    rows = query.order_by(ActivityType.name, ActivityType.id).offset(offset).limit(limit).all()
    return jsonify({
        "data": [{"id": activity_type.id, "name": activity_type.name} for activity_type in rows],
        "pagination": {"total": total, "limit": limit, "offset": offset},
    })


@app.get("/api/crm/v1/tasks")
@crm_api_required
def crm_api_tasks():
    limit, offset = crm_api_pagination()
    if limit is None:
        return jsonify({"error": "invalid_pagination"}), 400
    date_from = crm_api_date(request.args.get("date_from"))
    date_to = crm_api_date(request.args.get("date_to"))
    if date_from is False or date_to is False:
        return jsonify({"error": "invalid_date", "message": "Используйте формат YYYY-MM-DD."}), 400
    query = Activity.query.options(
        joinedload(Activity.owner).joinedload(User.department), joinedload(Activity.type),
        joinedload(Activity.assigned_department), joinedload(Activity.source_letter), joinedload(Activity.completion_memo),
    ).filter(Activity.archived_at.is_(None))
    if date_from:
        query = query.filter(or_(Activity.end_date >= dt.datetime.combine(date_from, dt.time.min), Activity.start_date >= dt.datetime.combine(date_from, dt.time.min)))
    if date_to:
        query = query.filter(or_(Activity.start_date <= dt.datetime.combine(date_to, dt.time.max), Activity.end_date <= dt.datetime.combine(date_to, dt.time.max)))
    status = (request.args.get("status") or "").strip()
    if status:
        query = query.filter(Activity.status == status)
    total = query.count()
    rows = query.order_by(Activity.id.desc()).offset(offset).limit(limit).all()
    data = []
    for activity in rows:
        data.append({
            "id": activity.id, "title": activity.title, "description": activity.description,
            "type": {"id": activity.type_id, "name": activity.type.name if activity.type else None},
            "status": activity.status, "status_label": get_activity_status_meta(activity.status)["label"],
            "priority": activity.priority or 3,
            "start_date": iso_or_none(activity.start_date), "end_date": iso_or_none(activity.end_date),
            "postponed_to": iso_or_none(activity.postponed_to), "status_comment": activity.status_comment,
            "direction": activity.direction, "outgoing_document": activity.outgoing_document,
            "owner": crm_api_user(activity.owner),
            "assigned_department": {"id": activity.assigned_department.id, "name": activity.assigned_department.name} if activity.assigned_department else None,
            "plan_id": activity.plan_id,
            "source_document": {"id": activity.source_letter.id, "number": activity.source_letter.reg_number, "subject": activity.source_letter.subject} if activity.source_letter else None,
            "completion_memo": {"id": activity.completion_memo.id, "number": activity.completion_memo.reg_number, "subject": activity.completion_memo.subject} if activity.completion_memo else None,
        })
    return jsonify({"data": data, "pagination": {"total": total, "limit": limit, "offset": offset}})


@app.get("/api/crm/v1/tasks/<int:activity_id>")
@crm_api_required
def crm_api_task_detail(activity_id):
    activity = Activity.query.options(
        joinedload(Activity.owner).joinedload(User.department),
        joinedload(Activity.created_by).joinedload(User.department),
        joinedload(Activity.approved_by).joinedload(User.department),
        joinedload(Activity.type),
        joinedload(Activity.assigned_department),
        joinedload(Activity.source_letter),
        joinedload(Activity.completion_memo),
        joinedload(Activity.letters).joinedload(Letter.author),
        joinedload(Activity.letters).joinedload(Letter.executor),
    ).filter_by(id=activity_id).first_or_404()
    return jsonify({"data": crm_api_activity_full_payload(activity)})


@app.get("/api/crm/v1/tasks/<int:activity_id>/logs")
@crm_api_required
def crm_api_task_logs(activity_id):
    activity = db.session.get(Activity, activity_id)
    if not activity:
        return jsonify({"error": "not_found", "message": "Задача не найдена."}), 404
    rows = ActivityLog.query.filter_by(activity_id=activity.id).order_by(
        ActivityLog.entry_date.asc(), ActivityLog.id.asc()
    ).all()
    return jsonify({"data": [{
        "id": row.id,
        "parent_id": row.parent_id,
        "text": row.text,
        "entry_date": iso_or_none(row.entry_date),
        "is_done": bool(row.is_done),
    } for row in rows]})


@app.post("/api/crm/v1/tasks/<int:activity_id>/logs")
@crm_api_required
def crm_api_task_log_create(activity_id):
    activity = db.session.get(Activity, activity_id)
    if not activity:
        return jsonify({"error": "not_found", "message": "Задача не найдена."}), 404
    if activity.archived_at:
        return jsonify({"error": "archived", "message": "Архивную задачу изменять нельзя."}), 409

    payload = crm_api_json_body()
    if not payload:
        return jsonify({"error": "invalid_json", "message": "Передайте JSON-объект записи журнала."}), 400

    text_value = payload.get("text")
    if not isinstance(text_value, str) or not text_value.strip():
        return jsonify({"error": "validation_error", "message": "Укажите непустой text."}), 400

    entry_date = crm_api_datetime(payload.get("entry_date")) if "entry_date" in payload else dt.datetime.utcnow()
    if entry_date is False:
        return jsonify({"error": "validation_error", "message": "entry_date должен быть датой ISO 8601."}), 400

    parent_id = payload.get("parent_id")
    parent = None
    if parent_id is not None:
        parent = ActivityLog.query.filter_by(id=parent_id, activity_id=activity.id).first() if isinstance(parent_id, int) else None
        if not parent:
            return jsonify({"error": "validation_error", "message": "Родительская запись журнала не найдена."}), 400

    row = ActivityLog(
        activity_id=activity.id,
        text=text_value.strip(),
        entry_date=entry_date or dt.datetime.utcnow(),
        parent_id=parent.id if parent else None,
        is_done=bool(payload.get("is_done", False)),
    )
    db.session.add(row)
    db.session.flush()
    add_activity_history(activity, "Журнал действий", "Через API добавлена запись журнала: {}".format(text_value.strip()))
    db.session.commit()

    return jsonify({"data": {
        "id": row.id,
        "parent_id": row.parent_id,
        "text": row.text,
        "entry_date": iso_or_none(row.entry_date),
        "is_done": bool(row.is_done),
    }}), 201


@app.post("/api/crm/v1/tasks")
@crm_api_required
def crm_api_create_task():
    payload = crm_api_json_body()
    if not payload:
        return jsonify({"error": "invalid_json", "message": "Передайте JSON-объект задачи."}), 400

    title = payload.get("title")
    owner = crm_api_active_user(payload.get("owner_id"))
    start_date = crm_api_datetime(payload.get("start_date"))
    end_date = crm_api_datetime(payload.get("end_date"))
    type_id = payload.get("type_id")
    activity_type = db.session.get(ActivityType, type_id) if isinstance(type_id, int) else None
    if not isinstance(title, str) or not title.strip() or not owner or start_date is False or not start_date or end_date is False or not activity_type:
        return jsonify({"error": "validation_error", "message": "Укажите title, type_id, owner_id и start_date. Исполнитель должен быть активен."}), 400
    if end_date and end_date < start_date:
        return jsonify({"error": "validation_error", "message": "end_date не может быть раньше start_date."}), 400

    priority = parse_activity_priority(payload.get("priority", 3))
    complexity_level = parse_activity_complexity(payload.get("complexity_level", 3))
    status = payload.get("status", "in_progress")
    if priority is None or complexity_level is None or status not in ACTIVITY_STATUS_META:
        return jsonify({"error": "validation_error", "message": "Некорректные priority, complexity_level или status."}), 400

    assigned_department_id = payload.get("assigned_department_id")
    assigned_department = None
    if assigned_department_id is not None:
        assigned_department = db.session.get(Department, assigned_department_id) if isinstance(assigned_department_id, int) else None
        if not assigned_department:
            return jsonify({"error": "validation_error", "message": "Назначенный отдел не найден."}), 400

    source_letter_id = payload.get("source_letter_id")
    completion_memo_id = payload.get("completion_memo_id")
    source_letter = db.session.get(Letter, source_letter_id) if source_letter_id is not None else None
    completion_memo = db.session.get(Letter, completion_memo_id) if completion_memo_id is not None else None
    if (source_letter_id is not None and not source_letter) or (completion_memo_id is not None and not completion_memo):
        return jsonify({"error": "validation_error", "message": "Связанный документ не найден."}), 400

    plan_id = payload.get("plan_id")
    plan = db.session.get(Plan, plan_id) if plan_id is not None else None
    if plan_id is not None and not plan:
        return jsonify({"error": "validation_error", "message": "План не найден."}), 400

    postponed_to = crm_api_datetime(payload.get("postponed_to"))
    if postponed_to is False:
        return jsonify({"error": "validation_error", "message": "postponed_to должен быть датой ISO 8601."}), 400
    if status == "postponed" and not postponed_to:
        return jsonify({"error": "validation_error", "message": "Для статуса postponed укажите postponed_to."}), 400

    description = crm_api_optional_text(payload.get("description"))
    direction = crm_api_optional_text(payload.get("direction"))
    outgoing_document = crm_api_optional_text(payload.get("outgoing_document"))
    status_comment = crm_api_optional_text(payload.get("status_comment"))
    if False in (description, direction, outgoing_document, status_comment):
        return jsonify({"error": "validation_error", "message": "Текстовые поля должны быть строкой или null."}), 400

    activity = Activity(
        title=title.strip(),
        description=description,
        type_id=activity_type.id,
        owner_id=owner.id,
        start_date=start_date,
        end_date=end_date,
        status=status,
        priority=priority,
        complexity_level=complexity_level,
        direction=direction,
        outgoing_document=outgoing_document,
        assigned_department_id=assigned_department.id if assigned_department else None,
        status_comment=status_comment,
        postponed_to=postponed_to,
        source_letter_id=source_letter.id if source_letter else None,
        completion_memo_id=completion_memo.id if completion_memo else None,
        plan_id=plan.id if plan else None,
        plan_requested=bool(payload.get("plan_requested", False)),
    )
    db.session.add(activity)
    db.session.flush()
    if plan:
        position = max((item.position for item in plan.items), default=0) + 1
        db.session.add(PlanItem(
            plan_id=plan.id,
            executor_id=owner.id,
            position=position,
            task_text=activity.title,
            deadline_kind="date" if activity.end_date else "month",
            deadline_date=activity.end_date.date() if activity.end_date else None,
            planned_hours=calculate_activity_effort_hours(activity) or 0,
            linked_activity_id=activity.id,
            status="planned",
        ))
        activity.plan_requested = False
    add_activity_history(activity, "Создание задачи", "Задача создана через внешнюю CRM.")
    db.session.commit()
    return jsonify({"data": crm_api_activity_payload(activity)}), 201


@app.patch("/api/crm/v1/tasks/<int:activity_id>")
@crm_api_required
def crm_api_update_task(activity_id):
    payload = crm_api_json_body()
    if not payload:
        return jsonify({"error": "invalid_json", "message": "Передайте JSON-объект с изменяемыми полями."}), 400
    activity = Activity.query.options(
        joinedload(Activity.owner).joinedload(User.department), joinedload(Activity.type),
        joinedload(Activity.assigned_department), joinedload(Activity.source_letter), joinedload(Activity.completion_memo),
    ).filter_by(id=activity_id).first_or_404()
    if activity.archived_at:
        return jsonify({"error": "archived", "message": "Архивную задачу изменять нельзя."}), 409

    changes = []
    if "title" in payload:
        title = payload["title"]
        if not isinstance(title, str) or not title.strip():
            return jsonify({"error": "validation_error", "message": "title не может быть пустым."}), 400
        if title.strip() != activity.title:
            activity.title = title.strip()
            changes.append("название")
    for field in ("description", "direction", "outgoing_document", "status_comment"):
        if field in payload:
            value = payload[field]
            if value is not None and not isinstance(value, str):
                return jsonify({"error": "validation_error", "message": "{} должен быть строкой или null.".format(field)}), 400
            if getattr(activity, field) != (value.strip() if isinstance(value, str) and value.strip() else None):
                setattr(activity, field, value.strip() if isinstance(value, str) and value.strip() else None)
                changes.append(field)
    if "type_id" in payload:
        activity_type = db.session.get(ActivityType, payload["type_id"]) if isinstance(payload["type_id"], int) else None
        if not activity_type:
            return jsonify({"error": "validation_error", "message": "Тип задачи не найден."}), 400
        if activity.type_id != activity_type.id:
            activity.type_id = activity_type.id
            changes.append("тип")
    if "owner_id" in payload:
        owner = crm_api_active_user(payload["owner_id"])
        if not owner:
            return jsonify({"error": "validation_error", "message": "Исполнитель не найден или отключён."}), 400
        if activity.owner_id != owner.id:
            activity.owner_id = owner.id
            changes.append("исполнитель")
    if "assigned_department_id" in payload:
        department_id = payload["assigned_department_id"]
        department = db.session.get(Department, department_id) if isinstance(department_id, int) else None
        if department_id is not None and not department:
            return jsonify({"error": "validation_error", "message": "Назначенный отдел не найден."}), 400
        if activity.assigned_department_id != (department.id if department else None):
            activity.assigned_department_id = department.id if department else None
            changes.append("назначенный отдел")
    if "priority" in payload:
        priority = parse_activity_priority(payload["priority"])
        if priority is None:
            return jsonify({"error": "validation_error", "message": "priority должен быть от 1 до 5."}), 400
        if activity.priority != priority:
            activity.priority = priority
            changes.append("приоритет")
    if "complexity_level" in payload:
        complexity_level = parse_activity_complexity(payload["complexity_level"])
        if complexity_level is None:
            return jsonify({"error": "validation_error", "message": "complexity_level должен быть от 1 до 5."}), 400
        if activity.complexity_level != complexity_level:
            activity.complexity_level = complexity_level
            changes.append("сложность")
    for field in ("start_date", "end_date", "postponed_to"):
        if field in payload:
            value = crm_api_datetime(payload[field])
            if value is False:
                return jsonify({"error": "validation_error", "message": "{} должен быть датой ISO 8601 или null.".format(field)}), 400
            if getattr(activity, field) != value:
                setattr(activity, field, value)
                changes.append(field)
    if activity.end_date and activity.start_date and activity.end_date < activity.start_date:
        return jsonify({"error": "validation_error", "message": "end_date не может быть раньше start_date."}), 400
    if "status" in payload:
        status = payload["status"]
        if status not in ACTIVITY_STATUS_META:
            return jsonify({"error": "validation_error", "message": "Некорректный статус задачи."}), 400
        if status == "postponed" and not activity.postponed_to:
            return jsonify({"error": "validation_error", "message": "Для статуса postponed укажите postponed_to."}), 400
        if activity.status != status:
            activity.status = status
            changes.append("статус")
    if "plan_id" in payload:
        plan_id = payload["plan_id"]
        plan = db.session.get(Plan, plan_id) if isinstance(plan_id, int) else None
        if plan_id is not None and not plan:
            return jsonify({"error": "validation_error", "message": "План не найден."}), 400
        if activity.plan_id != (plan.id if plan else None):
            activity.plan_id = plan.id if plan else None
            changes.append("план")
    if "plan_requested" in payload:
        activity.plan_requested = bool(payload["plan_requested"])
        changes.append("ожидание плана")

    if changes:
        add_activity_history(activity, "Изменение задачи", "Изменено через внешнюю CRM: {}.".format(", ".join(changes)))
        db.session.commit()
    return jsonify({"data": crm_api_activity_payload(activity)})


@app.get("/api/crm/v1/plans")
@crm_api_required
def crm_api_plans():
    limit, offset = crm_api_pagination()
    if limit is None:
        return jsonify({"error": "invalid_pagination"}), 400
    date_from = crm_api_date(request.args.get("date_from"))
    date_to = crm_api_date(request.args.get("date_to"))
    if date_from is False or date_to is False:
        return jsonify({"error": "invalid_date", "message": "Используйте формат YYYY-MM-DD."}), 400
    query = Plan.query.options(
        joinedload(Plan.creator),
        joinedload(Plan.items).joinedload(PlanItem.executor).joinedload(User.department),
        joinedload(Plan.items).joinedload(PlanItem.linked_activity),
    )
    if date_from:
        query = query.filter(Plan.end_date >= date_from)
    if date_to:
        query = query.filter(Plan.start_date <= date_to)
    total = query.count()
    rows = query.order_by(Plan.start_date.desc(), Plan.id.desc()).offset(offset).limit(limit).all()
    data = []
    for plan in rows:
        data.append({
            "id": plan.id, "start_date": iso_or_none(plan.start_date), "end_date": iso_or_none(plan.end_date),
            "text": plan.plan_text, "approval_status": plan.approval_status,
            "created_at": iso_or_none(plan.created_at), "creator": crm_api_user(plan.creator),
            "items": [{
            "id": item.id, "position": item.position, "text": item.task_text,
            "deadline_kind": item.deadline_kind, "deadline_date": iso_or_none(item.deadline_date),
            "planned_hours": item.planned_hours or 0,
            "status": item.status, "comment": item.comment,
                "completed_at": iso_or_none(item.completed_at), "executor": crm_api_user(item.executor),
                "linked_task": {"id": item.linked_activity.id, "title": item.linked_activity.title} if item.linked_activity else None,
            } for item in plan.items],
        })
    return jsonify({"data": data, "pagination": {"total": total, "limit": limit, "offset": offset}})


@app.get("/api/crm/v1/plans/report")
@crm_api_required
def crm_api_plans_report():
    """Monthly plan report with source lines and aggregates for an external CRM."""
    payload = crm_api_plan_reporting_payload((request.args.get("month") or "").strip())
    if payload is None:
        return jsonify({"error": "invalid_month", "message": "Используйте month в формате YYYY-MM."}), 400
    return jsonify({"data": payload})


@app.get("/api/crm/v1/plans/analytics")
@crm_api_required
def crm_api_plans_analytics():
    """Monthly aggregate plan metrics without the per-line report data."""
    payload = crm_api_plan_reporting_payload((request.args.get("month") or "").strip())
    if payload is None:
        return jsonify({"error": "invalid_month", "message": "Используйте month в формате YYYY-MM."}), 400
    payload.pop("items")
    return jsonify({"data": payload})


@app.post("/api/crm/v1/plans")
@crm_api_required
def crm_api_create_plan():
    payload = crm_api_json_body()
    if not payload:
        return jsonify({"error": "invalid_json", "message": "Передайте JSON-объект плана."}), 400
    start_date = crm_api_date(payload.get("start_date"))
    end_date = crm_api_date(payload.get("end_date"))
    plan_text = payload.get("text")
    if not start_date or not end_date or start_date is False or end_date is False or not isinstance(plan_text, str) or not plan_text.strip():
        return jsonify({"error": "validation_error", "message": "Укажите text, start_date и end_date в формате YYYY-MM-DD."}), 400
    if end_date < start_date:
        return jsonify({"error": "validation_error", "message": "end_date не может быть раньше start_date."}), 400
    creator_id = payload.get("created_by_id")
    creator = crm_api_active_user(creator_id) if creator_id is not None else None
    if creator_id is not None and not creator:
        return jsonify({"error": "validation_error", "message": "Автор плана не найден или отключён."}), 400
    if creator and Plan.query.filter_by(created_by=creator.id, start_date=start_date).first():
        return jsonify({"error": "conflict", "message": "У автора уже есть план на этот месяц."}), 409
    approval_status = payload.get("approval_status", "submitted")
    if approval_status not in {"submitted", "revision", "approved"}:
        return jsonify({"error": "validation_error", "message": "Некорректный approval_status."}), 400
    items = payload.get("items")
    if not isinstance(items, list) or not items:
        return jsonify({"error": "validation_error", "message": "Добавьте непустой массив items."}), 400
    revision_comment = crm_api_optional_text(payload.get("revision_comment"))
    if revision_comment is False:
        return jsonify({"error": "validation_error", "message": "revision_comment должен быть строкой или null."}), 400

    plan = Plan(
        start_date=start_date,
        end_date=end_date,
        plan_text=plan_text.strip(),
        created_by=creator.id if creator else None,
        approval_status=approval_status,
        revision_comment=revision_comment,
        approved_at=dt.datetime.utcnow() if approval_status == "approved" else None,
    )
    db.session.add(plan)
    db.session.flush()
    for position, item_payload in enumerate(items, start=1):
        item = crm_api_validate_plan_item(item_payload, plan, position)
        if isinstance(item, str):
            db.session.rollback()
            return jsonify({"error": "validation_error", "message": item, "item_position": position}), 400
        db.session.add(item)
    db.session.flush()
    add_pending_activities_to_plan(plan)
    db.session.commit()
    return jsonify({"data": crm_api_plan_payload(plan)}), 201


@app.patch("/api/crm/v1/plans/<int:plan_id>")
@crm_api_required
def crm_api_update_plan(plan_id):
    payload = crm_api_json_body()
    if not payload:
        return jsonify({"error": "invalid_json", "message": "Передайте JSON-объект с изменяемыми полями."}), 400
    plan = Plan.query.options(
        joinedload(Plan.creator), joinedload(Plan.items).joinedload(PlanItem.executor).joinedload(User.department),
        joinedload(Plan.items).joinedload(PlanItem.linked_activity),
    ).filter_by(id=plan_id).first_or_404()
    if "text" in payload:
        text_value = payload["text"]
        if not isinstance(text_value, str) or not text_value.strip():
            return jsonify({"error": "validation_error", "message": "text не может быть пустым."}), 400
        plan.plan_text = text_value.strip()
    for field in ("start_date", "end_date"):
        if field in payload:
            value = crm_api_date(payload[field])
            if not value:
                return jsonify({"error": "validation_error", "message": "{} должен быть датой YYYY-MM-DD.".format(field)}), 400
            setattr(plan, field, value)
    if plan.end_date < plan.start_date:
        return jsonify({"error": "validation_error", "message": "end_date не может быть раньше start_date."}), 400
    if "approval_status" in payload:
        status = payload["approval_status"]
        if status not in {"submitted", "revision", "approved"}:
            return jsonify({"error": "validation_error", "message": "Некорректный approval_status."}), 400
        plan.approval_status = status
        plan.approved_at = dt.datetime.utcnow() if status == "approved" else None
    if "revision_comment" in payload:
        value = payload["revision_comment"]
        if value is not None and not isinstance(value, str):
            return jsonify({"error": "validation_error", "message": "revision_comment должен быть строкой или null."}), 400
        plan.revision_comment = value.strip() if isinstance(value, str) and value.strip() else None
    db.session.commit()
    return jsonify({"data": crm_api_plan_payload(plan)})


@app.post("/api/crm/v1/plans/<int:plan_id>/items")
@crm_api_required
def crm_api_create_plan_item(plan_id):
    payload = crm_api_json_body()
    if not payload:
        return jsonify({"error": "invalid_json", "message": "Передайте JSON-объект пункта плана."}), 400
    plan = db.session.get(Plan, plan_id)
    if not plan:
        return jsonify({"error": "not_found", "message": "План не найден."}), 404
    position = max((item.position for item in plan.items), default=0) + 1
    item = crm_api_validate_plan_item(payload, plan, position)
    if isinstance(item, str):
        return jsonify({"error": "validation_error", "message": item}), 400
    db.session.add(item)
    db.session.commit()
    return jsonify({"data": crm_api_plan_payload(plan)}), 201


@app.patch("/api/crm/v1/plans/<int:plan_id>/items/<int:item_id>")
@crm_api_required
def crm_api_update_plan_item(plan_id, item_id):
    payload = crm_api_json_body()
    if not payload:
        return jsonify({"error": "invalid_json", "message": "Передайте JSON-объект с изменяемыми полями."}), 400
    item = PlanItem.query.filter_by(id=item_id, plan_id=plan_id).first()
    if not item:
        return jsonify({"error": "not_found", "message": "Пункт плана не найден."}), 404
    updated_item = crm_api_validate_plan_item(payload, item.plan, item.position, existing=item)
    if isinstance(updated_item, str):
        return jsonify({"error": "validation_error", "message": updated_item}), 400
    db.session.commit()
    return jsonify({"data": crm_api_plan_payload(item.plan)})


@app.get("/api/crm/v1/memos")
@crm_api_required
def crm_api_memos():
    limit, offset = crm_api_pagination()
    if limit is None:
        return jsonify({"error": "invalid_pagination"}), 400
    date_from = crm_api_date(request.args.get("date_from"))
    date_to = crm_api_date(request.args.get("date_to"))
    if date_from is False or date_to is False:
        return jsonify({"error": "invalid_date", "message": "Используйте формат YYYY-MM-DD."}), 400
    query = Letter.query.options(
        joinedload(Letter.author).joinedload(User.department), joinedload(Letter.executor).joinedload(User.department),
        joinedload(Letter.departments),
    ).filter(Letter.source_id == ensure_memo_source(), Letter.archived_at.is_(None))
    if date_from:
        query = query.filter(or_(Letter.letter_date >= date_from, Letter.created_at >= dt.datetime.combine(date_from, dt.time.min)))
    if date_to:
        query = query.filter(or_(Letter.letter_date <= date_to, Letter.created_at <= dt.datetime.combine(date_to, dt.time.max)))
    approval_status = (request.args.get("approval_status") or "").strip()
    if approval_status:
        query = query.filter(Letter.approval_status == approval_status)
    total = query.count()
    rows = query.order_by(Letter.letter_date.desc().nullslast(), Letter.id.desc()).offset(offset).limit(limit).all()
    data = []
    for memo in rows:
        data.append({
            "id": memo.id, "number": memo.reg_number, "subject": memo.subject, "body": memo.body,
            "date": iso_or_none(memo.letter_date), "created_at": iso_or_none(memo.created_at),
            "approval_status": memo.approval_status, "approved_at": iso_or_none(memo.approved_at),
            "author": crm_api_user(memo.author), "executor": crm_api_user(memo.executor),
            "departments": [{"id": department.id, "name": department.name} for department in memo.departments],
        })
    return jsonify({"data": data, "pagination": {"total": total, "limit": limit, "offset": offset}})


@app.route("/reports/external-links", methods=["GET", "POST"])
@role_required("admin", "superadmin")
def report_external_links():
    created_url = None
    if request.method == "POST":
        report_key = (request.form.get("report_key") or "").strip()
        expires_raw = (request.form.get("expires_at") or "").strip()
        workshop_id = (request.form.get("workshop_id") or "").strip()
        if report_key not in REPORT_SHARE_OPTIONS:
            flash("Выберите отчет.", "warning")
            return redirect(request.url)
        try:
            expires_at = dt.datetime.strptime(expires_raw, "%Y-%m-%dT%H:%M")
        except ValueError:
            flash("Укажите дату и время окончания доступа.", "warning")
            return redirect(request.url)
        if expires_at <= dt.datetime.utcnow():
            flash("Дата окончания доступа должна быть в будущем.", "warning")
            return redirect(request.url)
        payload = {}
        if report_key in {"plans_statistics", "plans_monthly_report"} and workshop_id:
            if not workshop_id.isdigit() or not db.session.get(Workshop, int(workshop_id)):
                flash("Выберите существующий цех.", "warning")
                return redirect(request.url)
            payload["workshop_id"] = int(workshop_id)
        link = SharedLink(
            doc_type="report:{}".format(report_key), doc_id=0, token=uuid.uuid4().hex,
            created_by=current_user.id, expires_at=expires_at, payload=payload,
        )
        db.session.add(link)
        db.session.commit()
        endpoint = REPORT_SHARE_OPTIONS[report_key]["endpoint"]
        created_url = url_for(endpoint, share=link.token, _external=True)
        flash("Внешняя ссылка создана.", "success")

    links = []
    for link in SharedLink.query.filter(SharedLink.doc_type.like("report:%")).order_by(SharedLink.created_at.desc()).all():
        report_key = link.doc_type.split(":", 1)[1]
        option = REPORT_SHARE_OPTIONS.get(report_key)
        if not option:
            continue
        links.append({
            "id": link.id, "label": option["label"], "expires_at": link.expires_at,
            "created_at": link.created_at, "url": url_for(option["endpoint"], share=link.token, _external=True),
            "workshop": db.session.get(Workshop, get_report_workshop_id(link.payload)) if get_report_workshop_id(link.payload) else None,
        })
    return render_template(
        "report_external_links.html", title="Внешние ссылки на отчеты", report_options=REPORT_SHARE_OPTIONS,
        links=links, workshops=Workshop.query.order_by(Workshop.name).all(), created_url=created_url,
        default_expires=(dt.datetime.now() + dt.timedelta(days=7)).strftime("%Y-%m-%dT%H:%M"),
    )


@app.post("/reports/external-links/<int:link_id>/delete")
@role_required("admin", "superadmin")
def delete_report_external_link(link_id):
    link = SharedLink.query.filter(SharedLink.id == link_id, SharedLink.doc_type.like("report:%")).first_or_404()
    db.session.delete(link)
    db.session.commit()
    flash("Внешняя ссылка отозвана.", "success")
    return redirect(url_for("report_external_links"))


def iso_or_none(dtobj):
    return dtobj.isoformat() if dtobj else None

@app.route("/")
def index():
    if current_user.is_authenticated:
        return redirect(url_for("dashboard"))
    return redirect(url_for("login_page"))

@app.get("/login")
def login_page():
    login_users = (
        User.query
        .options(joinedload(User.department).joinedload(Department.workshop))
        .filter_by(is_approved=True, is_enabled=True)
        .order_by(User.full_name.asc())
        .all()
    )
    return render_template(
        "login.html",
        title="Вход",
        login_users=login_users,
        workshops=Workshop.query.order_by(Workshop.name.asc()).all(),
        has_unassigned_users=any(
            not user.department or not user.department.workshop_id
            for user in login_users
        ),
        superadmin_login=SUPERADMIN_LOGIN,
    )
    return render_template("login.html", title="Вход")


@app.get("/api/chat/tray/login-users")
def chat_tray_login_users():
    login_users = (
        User.query
        .filter_by(is_approved=True, is_enabled=True)
        .order_by(User.full_name.asc())
        .all()
    )

    lines = []
    if SUPERADMIN_LOGIN:
        lines.append(
            "\t".join([
                "USER",
                quote(SUPERADMIN_LOGIN, safe=""),
                quote("Суперадминистратор", safe=""),
            ])
        )

    for user in login_users:
        lines.append(
            "\t".join([
                "USER",
                quote(user.username or "", safe=""),
                quote(user.full_name or user.username or "", safe=""),
            ])
        )

    payload = "\n".join(lines)
    return current_app.response_class(payload, mimetype="text/plain; charset=utf-8")

@app.get("/logout")
@login_required
def logout():
    logout_user()
    return redirect(url_for("login_page"))

@app.get("/dashboard")
@login_required
def dashboard():
    from datetime import date, datetime, timedelta
    normalize_activity_statuses()
    my_acts = Activity.query.filter_by(owner_id=current_user.id).count()
    memo_source_id = ensure_memo_source()
    my_letters = (
        Letter.query
        .filter(
            Letter.author_id == current_user.id,
            Letter.archived_at.is_(None),
            or_(Letter.source_id.is_(None), Letter.source_id != memo_source_id),
        )
        .count()
    )
    my_memos = (
        Letter.query
        .filter(
            Letter.author_id == current_user.id,
            Letter.archived_at.is_(None),
            Letter.source_id == memo_source_id,
        )
        .count()
    )
    pending_approval = 0
    if current_user.role in ("manager", "admin", "superadmin", "deputy"):
        pending_approval = Activity.query.filter_by(status="done").count()

    today = date.today()
    quarter_start_month = ((today.month - 1) // 3) * 3 + 1
    first_day = date(today.year, quarter_start_month, 1)
    if quarter_start_month == 10:
        next_quarter = date(today.year + 1, 1, 1)
    else:
        next_quarter = date(today.year, quarter_start_month + 3, 1)
    last_day = next_quarter - timedelta(days=1)
    start_str = request.args.get("date_from", first_day.strftime("%Y-%m-%d"))
    end_str = request.args.get("date_to", last_day.strftime("%Y-%m-%d"))
    my_only = (request.args.get("my_only") or "").lower() in ("1", "true", "on")
    owner_id = request.args.get("owner_id", "").strip()
    selected_status = request.args.get("status", "active_and_not_done").strip() or "active_and_not_done"

    try:
        date_from = datetime.strptime(start_str, "%Y-%m-%d")
        date_to = datetime.combine(datetime.strptime(end_str, "%Y-%m-%d").date(), datetime.max.time())
    except ValueError:
        date_from = datetime.combine(first_day, datetime.min.time())
        date_to = datetime.combine(last_day, datetime.max.time())

    tasks_query = Activity.query.options(
        joinedload(Activity.type),
        joinedload(Activity.owner),
    ).filter(
        Activity.start_date >= date_from,
        Activity.start_date <= date_to,
    ).order_by(
        Activity.start_date.desc(),
        Activity.id.desc(),
    )
    if selected_status not in {value for value, _label in ACTIVITY_STATUS_FILTERS}:
        selected_status = "active_and_not_done"
    tasks_query = apply_activity_status_filter(tasks_query, selected_status)
    privileged_roles = ("admin", "manager", "superadmin", "deputy")
    filter_users = []
    if current_user.role in privileged_roles:
        if my_only:
            tasks_query = tasks_query.filter(Activity.owner_id == current_user.id)
        elif owner_id:
            try:
                tasks_query = tasks_query.filter(Activity.owner_id == int(owner_id))
            except ValueError:
                pass
        filter_users = User.query.filter_by(is_approved=True).order_by(User.full_name).all()
        dashboard_tasks = tasks_query.limit(10).all()
    else:
        dashboard_tasks = tasks_query.filter(Activity.owner_id == current_user.id).limit(10).all()

    return render_template(
        "dashboard.html",
        title="Главная",
        my_acts=my_acts,
        my_letters=my_letters,
        my_memos=my_memos,
        pending_approval=pending_approval,
        dashboard_tasks=dashboard_tasks,
        date_from=start_str,
        date_to=end_str,
        my_only=my_only,
        owner_id=owner_id,
        selected_status=selected_status,
        status_filters=ACTIVITY_STATUS_FILTERS,
        filter_users=filter_users,
        get_status_meta=get_activity_status_meta,
    )


@app.get("/help")
@login_required
def help_page():
    return render_template("help.html", title="Справка")


@app.route("/admin/document-templates", methods=["GET", "POST"])
@role_required("admin", "superadmin")
def document_templates_page():
    if request.method == "POST":
        template_kind = (request.form.get("template_kind") or "").strip()
        name = (request.form.get("name") or "").strip()
        if not name:
            flash("Укажите название шаблона.", "warning")
            return redirect(url_for("document_templates_page"))

        if template_kind == "activity":
            title = (request.form.get("title") or "").strip()
            type_id = (request.form.get("type_id") or "").strip()
            if not title or not type_id.isdigit() or not db.session.get(ActivityType, int(type_id)):
                flash("Для шаблона задачи укажите название задачи и её тип.", "warning")
                return redirect(url_for("document_templates_page"))
            duration_raw = (request.form.get("duration_days") or "").strip()
            try:
                duration_days = int(duration_raw) if duration_raw else None
            except ValueError:
                duration_days = None
            if duration_days is not None and duration_days < 0:
                flash("Срок шаблона не может быть отрицательным.", "warning")
                return redirect(url_for("document_templates_page"))
            template = ActivityTemplate(
                name=name, title=title, description=(request.form.get("description") or "").strip(),
                type_id=int(type_id), direction=(request.form.get("direction") or "").strip(),
                outgoing_document=(request.form.get("outgoing_document") or "").strip(),
                duration_days=duration_days,
            )
        elif template_kind == "memo":
            subject = (request.form.get("subject") or "").strip()
            department_id = (request.form.get("department_id") or "").strip()
            executor_id = (request.form.get("executor_id") or "").strip()
            if not subject:
                flash("Для шаблона служебной записки укажите краткое содержание.", "warning")
                return redirect(url_for("document_templates_page"))
            if department_id and (not department_id.isdigit() or not db.session.get(Department, int(department_id))):
                flash("Выбран некорректный отдел.", "warning")
                return redirect(url_for("document_templates_page"))
            if executor_id and (not executor_id.isdigit() or not db.session.get(User, int(executor_id))):
                flash("Выбран некорректный исполнитель.", "warning")
                return redirect(url_for("document_templates_page"))
            template = MemoTemplate(
                name=name, subject=subject, body=(request.form.get("body") or "").strip(),
                department_id=int(department_id) if department_id else None,
                executor_id=int(executor_id) if executor_id else None,
            )
        else:
            abort(400)

        try:
            db.session.add(template)
            db.session.commit()
        except Exception:
            db.session.rollback()
            flash("Шаблон с таким названием уже существует.", "warning")
            return redirect(url_for("document_templates_page"))
        flash("Шаблон сохранён.", "success")
        return redirect(url_for("document_templates_page"))

    return render_template(
        "document_templates.html", title="Шаблоны документов",
        activity_templates=ActivityTemplate.query.options(joinedload(ActivityTemplate.type)).order_by(ActivityTemplate.name).all(),
        memo_templates=MemoTemplate.query.options(joinedload(MemoTemplate.department), joinedload(MemoTemplate.executor)).order_by(MemoTemplate.name).all(),
        activity_types=ActivityType.query.order_by(ActivityType.name).all(),
        departments=Department.query.order_by(Department.name).all(),
        users=User.query.filter_by(is_approved=True).order_by(User.full_name).all(),
    )


@app.post("/admin/document-templates/<string:template_kind>/<int:template_id>/delete")
@role_required("admin", "superadmin")
def delete_document_template(template_kind, template_id):
    model = ActivityTemplate if template_kind == "activity" else MemoTemplate if template_kind == "memo" else None
    if model is None:
        abort(404)
    template = db.session.get(model, template_id)
    if not template:
        abort(404)
    db.session.delete(template)
    db.session.commit()
    flash("Шаблон удалён.", "success")
    return redirect(url_for("document_templates_page"))


@app.route("/activities/create", methods=["GET", "POST"])
@login_required
def create_activity():
    from datetime import datetime
    types = ActivityType.query.order_by(ActivityType.name).all()
    templates = ActivityTemplate.query.options(joinedload(ActivityTemplate.type)).order_by(ActivityTemplate.name).all()
    template_id = (request.args.get("template_id") or "").strip()
    selected_template = db.session.get(ActivityTemplate, int(template_id)) if template_id.isdigit() else None

    can_assign_executors = current_user.role in ("admin", "superadmin", "deputy")
    if current_user.role in ("admin", "superadmin"):
        users = User.query.filter_by(is_approved=True).order_by(User.full_name).all()
    elif current_user.role == "deputy":
        users = []
        if current_user.department_id:
            users = (
                User.query.filter_by(
                    is_approved=True,
                    department_id=current_user.department_id,
                )
                .order_by(User.full_name)
                .all()
            )
    else:
        users = [current_user]

    existing_dirs = [d[0] for d in db.session.query(Activity.direction).filter(Activity.direction.isnot(None)).distinct().all()]

    if request.method == "POST":
        title = request.form.get("title", "").strip()
        description = request.form.get("description", "").strip()
        type_id = request.form.get("type_id")
        start_date = request.form.get("start_date")
        end_date = request.form.get("end_date")
        priority = parse_activity_priority(request.form.get("priority", "3"))
        complexity_level = parse_activity_complexity(request.form.get("complexity_level", "3"))
        direction = request.form.get("direction", "").strip()
        outgoing_document = request.form.get("outgoing_document", "").strip()
        add_to_plan = request.form.get("add_to_plan") == "1"

        executor_ids = (
            request.form.getlist("executors")
            if can_assign_executors
            else [str(current_user.id)]
        )

        if not title:
            flash("Введите название задачи.")
            return redirect(request.url)
        selected_type = None
        if type_id and type_id.isdigit():
            selected_type = db.session.get(ActivityType, int(type_id))
        if not selected_type:
            flash("Выберите тип задачи.", "warning")
            return redirect(request.url)
        if priority is None:
            flash("Выберите приоритет от 1 до 5.", "warning")
            return redirect(request.url)
        if complexity_level is None:
            flash("Выберите уровень сложности от 1 до 5.", "warning")
            return redirect(request.url)
        if can_assign_executors and not executor_ids:
            message = "Заместителю должен быть назначен отдел." if current_user.role == "deputy" and not current_user.department_id else "Выберите хотя бы одного ответственного."
            flash(message, "warning")
            return redirect(request.url)
        try:
            executor_ids = [int(user_id) for user_id in executor_ids]
        except ValueError:
            flash("Некорректный исполнитель.", "warning")
            return redirect(request.url)
        allowed_executor_ids = {user.id for user in users}
        if not executor_ids or not set(executor_ids).issubset(allowed_executor_ids):
            flash("Можно назначать задачи только сотрудникам доступного вам отдела.", "warning")
            return redirect(request.url)

        task_start = datetime.strptime(start_date, "%Y-%m-%d") if start_date else datetime.now()
        task_end = datetime.strptime(end_date, "%Y-%m-%d") if end_date else task_start
        created = []
        linked_to_plan = 0
        for executor_id in executor_ids:
            plan_id = None
            matched_plan = None
            if add_to_plan:
                # A task may be added only to the selected executor's own plan,
                # never to a plan belonging to another employee.
                matched_plan = (
                    Plan.query
                    .filter(
                        Plan.start_date <= task_start.date(),
                        Plan.end_date >= task_start.date(),
                        or_(
                            Plan.created_by == executor_id,
                            Plan.items.any(PlanItem.executor_id == executor_id),
                        ),
                    )
                    .order_by(
                        (Plan.items.any(PlanItem.executor_id == executor_id)).desc(),
                        Plan.start_date.desc(),
                        Plan.id.desc(),
                    )
                    .first()
                )
                if matched_plan:
                    plan_id = matched_plan.id
                    linked_to_plan += 1
            a = Activity(
                title=title,
                description=description,
                owner_id=executor_id,
                created_by_id=current_user.id if current_user.id != -1 else None,
                type_id=selected_type.id,
                start_date=task_start,
                end_date=datetime.strptime(end_date, "%Y-%m-%d") if end_date else None,
                status="in_progress",
                priority=priority,
                complexity_level=complexity_level,
                direction=direction,
                outgoing_document=outgoing_document,
                plan_id=plan_id,
                plan_requested=add_to_plan and not bool(matched_plan),
            )
            db.session.add(a)
            db.session.flush()
            if matched_plan:
                last_position = (
                    db.session.query(func.max(PlanItem.position))
                    .filter(PlanItem.plan_id == matched_plan.id)
                    .scalar()
                    or 0
                )
                db.session.add(PlanItem(
                    plan_id=matched_plan.id,
                    executor_id=executor_id,
                    position=last_position + 1,
                    task_text=title,
                    deadline_kind="date" if end_date else "month",
                    deadline_date=task_end.date() if end_date else None,
                    planned_hours=calculate_activity_effort_hours(a) or 0,
                    linked_activity_id=a.id,
                ))
                a.plan_requested = False
            created.append(a)
        db.session.commit()
        for activity in created:
            add_activity_history(
                activity,
                "Создание задачи",
                "Задача создана и назначена исполнителю: {}.".format(
                    activity.owner.full_name if activity.owner else "-"
                ),
            )
        db.session.commit()

        message = f"Создано {len(created)} задач."
        if add_to_plan:
            if linked_to_plan:
                message += f" В планы добавлено: {linked_to_plan}."
            else:
                message += " Подходящего плана пока нет: задача будет добавлена автоматически после его создания."
        flash(message)
        return redirect(url_for("activities_page"))

    return render_template(
        "activity_create.html",
        title="Создание задачи",
        types=types,
        users=users,
        existing_directions=existing_dirs,
        templates=templates,
        selected_template=selected_template,
        can_assign_executors=can_assign_executors,
        executor_scope_label=("Сотрудники вашего отдела" if current_user.role == "deputy" else "Ответственный"),
        initial_start_date=datetime.now().strftime("%Y-%m-%d"),
        initial_end_date=(datetime.now().date() + dt.timedelta(days=selected_template.duration_days)).strftime("%Y-%m-%d") if selected_template and selected_template.duration_days is not None else "",
    )

@app.get("/activities")
@login_required
def activities_page():
    from datetime import datetime, date, timedelta
    normalize_activity_statuses()

    today = date.today()
    first_day = date(today.year, today.month, 1)
    next_month = first_day.replace(day=28) + timedelta(days=4)
    last_day = next_month - timedelta(days=next_month.day)

    start_str = request.args.get("date_from", first_day.strftime("%Y-%m-%d"))
    end_str = request.args.get("date_to", last_day.strftime("%Y-%m-%d"))
    my_only = request.args.get("my_only") == "1"
    created_by_me = request.args.get("created_by_me") == "1"
    owner_id = request.args.get("owner_id", "").strip()
    search = request.args.get("search", "").strip()
    selected_status = request.args.get("status", "active_and_not_done").strip() or "active_and_not_done"
    sort = request.args.get("sort", "date_desc").strip()
    page = request.args.get("page", 1, type=int) or 1

    try:
        date_from = datetime.strptime(start_str, "%Y-%m-%d")
        date_to = datetime.strptime(end_str, "%Y-%m-%d")
    except ValueError:
        date_from = datetime.combine(first_day, datetime.min.time())
        date_to = datetime.combine(last_day, datetime.max.time())

    q = Activity.query.options(
        joinedload(Activity.owner), joinedload(Activity.created_by), joinedload(Activity.type)
    ).filter(Activity.archived_at.is_(None))

    privileged_roles = ("admin", "manager", "superadmin", "deputy")
    if current_user.role not in privileged_roles:
        department_id = getattr(current_user, "department_id", None)
        q = q.filter(
            or_(
                Activity.owner_id == current_user.id,
                Activity.assigned_department_id == department_id if department_id else False,
            )
        )
    if my_only:
        q = q.filter(Activity.owner_id == current_user.id)
    elif current_user.role in privileged_roles and owner_id:
        try:
            q = q.filter(Activity.owner_id == int(owner_id))
        except ValueError:
            pass

    if created_by_me:
        q = q.filter(Activity.created_by_id == current_user.id)

    q = q.filter(Activity.start_date >= date_from, Activity.start_date <= date_to)
    if selected_status not in {value for value, _label in ACTIVITY_STATUS_FILTERS}:
        selected_status = "active_and_not_done"
    q = apply_activity_status_filter(q, selected_status)
    if search:
        like_pattern = f"%{search}%"
        q = q.filter(
            or_(
                Activity.title.ilike(like_pattern),
                Activity.description.ilike(like_pattern),
                Activity.direction.ilike(like_pattern),
                Activity.outgoing_document.ilike(like_pattern),
            )
        )
    sort_orders = {
        "date_desc": (Activity.start_date.desc(), Activity.id.desc()),
        "date_asc": (Activity.start_date.asc(), Activity.id.asc()),
        "priority_desc": (Activity.priority.desc(), Activity.start_date.desc(), Activity.id.desc()),
        "priority_asc": (Activity.priority.asc(), Activity.start_date.desc(), Activity.id.desc()),
    }
    if sort not in sort_orders:
        sort = "date_desc"
    pagination = q.order_by(*sort_orders[sort]).paginate(page=page, per_page=50, error_out=False)
    activities = pagination.items
    filter_users = []
    if current_user.role in privileged_roles:
        filter_users = User.query.filter_by(is_approved=True).order_by(User.full_name).all()

    return render_template(
        "activities.html",
        title="Список задач",
        activities=activities,
        date_from=start_str,
        date_to=end_str,
        my_only=my_only,
        created_by_me=created_by_me,
        owner_id=owner_id,
        selected_status=selected_status,
        status_filters=ACTIVITY_STATUS_FILTERS,
        search=search,
        sort=sort,
        filter_users=filter_users,
        get_status_meta=get_activity_status_meta,
        get_priority_meta=get_activity_priority_meta,
        pagination=pagination,
    )


@app.get("/activities/analytics")
@role_required("manager", "admin", "superadmin")
def activities_analytics():
    normalize_activity_statuses()
    shared_payload = get_report_share_payload("activities_analytics")
    today = dt.date.today()
    quarter_start_month = ((today.month - 1) // 3) * 3 + 1
    default_from = dt.date(today.year, quarter_start_month, 1)
    if quarter_start_month == 10:
        default_to = dt.date(today.year + 1, 1, 1) - dt.timedelta(days=1)
    else:
        default_to = dt.date(today.year, quarter_start_month + 3, 1) - dt.timedelta(days=1)

    date_from_value = str((shared_payload or {}).get("date_from") or request.args.get("date_from") or default_from.isoformat()).strip()
    date_to_value = str((shared_payload or {}).get("date_to") or request.args.get("date_to") or default_to.isoformat()).strip()
    selected_department_id = str((shared_payload or {}).get("department_id") or request.args.get("department_id") or "").strip()
    selected_owner_id = str((shared_payload or {}).get("owner_id") or request.args.get("owner_id") or "").strip()
    try:
        date_from = datetime.strptime(date_from_value, "%Y-%m-%d")
        date_to = datetime.combine(datetime.strptime(date_to_value, "%Y-%m-%d").date(), datetime.max.time())
    except ValueError:
        date_from = datetime.combine(default_from, datetime.min.time())
        date_to = datetime.combine(default_to, datetime.max.time())
        date_from_value = default_from.isoformat()
        date_to_value = default_to.isoformat()

    query = Activity.query.options(joinedload(Activity.owner).joinedload(User.department)).filter(
        Activity.start_date >= date_from,
        Activity.start_date <= date_to,
    )
    is_external_share = shared_payload is not None
    is_manager = getattr(current_user, "role", None) == "manager" and not is_external_share
    current_department_id = getattr(current_user, "department_id", None)
    if is_manager:
        # A manager can assess only their own department; without a department,
        # the safe fallback is their personal tasks.
        if current_department_id:
            query = query.join(Activity.owner).filter(User.department_id == current_department_id)
            selected_department_id = str(current_department_id)
        else:
            query = query.filter(Activity.owner_id == current_user.id)
            selected_department_id = ""
    elif selected_department_id:
        try:
            query = query.join(Activity.owner).filter(User.department_id == int(selected_department_id))
        except ValueError:
            selected_department_id = ""
    if selected_owner_id:
        try:
            query = query.filter(Activity.owner_id == int(selected_owner_id))
        except ValueError:
            selected_owner_id = ""

    activities = query.order_by(Activity.start_date.desc(), Activity.id.desc()).all()
    departments_query = Department.query.order_by(Department.name)
    users_query = User.query.filter_by(is_approved=True).order_by(User.full_name)
    if is_manager:
        if current_department_id:
            departments_query = departments_query.filter(Department.id == current_department_id)
            users_query = users_query.filter(User.department_id == current_department_id)
        else:
            departments_query = departments_query.filter(Department.id == -1)
            users_query = users_query.filter(User.id == current_user.id)
    departments = departments_query.all()
    users = users_query.all()

    buckets = {
        "done": {"label": "Выполнено", "color": "#198754"},
        "not_done": {"label": "Не выполнено", "color": "#dc3545"},
        "in_progress": {"label": "В работе", "color": "#ffc107"},
        "postponed": {"label": "Перенос", "color": "#fd7e14"},
        "cancelled": {"label": "Отменено", "color": "#6f4e37"},
    }

    def status_bucket(status):
        if status in ("done", "approved"):
            return "done"
        if status in ("not_done",):
            return "not_done"
        if status in ("cancelled", "rejected"):
            return "cancelled"
        if status == "postponed":
            return "postponed"
        return "in_progress"

    complexity_hours = {
        level: get_activity_complexity_meta(level)["daily_hours"]
        for level in ACTIVITY_COMPLEXITY_META
    }

    def effort_hours_for(activity, complexity_level):
        if not activity.start_date or not activity.end_date:
            return 0
        start_day = activity.start_date.date()
        end_day = activity.end_date.date()
        if end_day < start_day:
            return 0
        working_days = sum(
            1 for offset in range((end_day - start_day).days + 1)
            if (start_day + dt.timedelta(days=offset)).weekday() < 5
        )
        return working_days * complexity_hours[complexity_level]

    def empty_row(name, department_name=""):
        return {
            "name": name,
            "department": department_name,
            "total": 0,
            "complexity_total": 0,
            "complexity_done": 0,
            "workload_hours": 0,
            **{key: 0 for key in buckets},
        }

    department_rows_by_id = {}
    user_rows_by_id = {}
    summary = empty_row("Все задачи")
    for activity in activities:
        bucket = status_bucket(activity.status)
        owner = activity.owner
        department = owner.department if owner else None
        department_id = department.id if department else 0
        department_name = department.name if department else "Без отдела"
        if department_id not in department_rows_by_id:
            department_rows_by_id[department_id] = empty_row(department_name)
        owner_id = owner.id if owner else 0
        owner_name = owner.full_name if owner else "Удаленный пользователь"
        if owner_id not in user_rows_by_id:
            user_rows_by_id[owner_id] = empty_row(owner_name, department_name)
        complexity_level = parse_activity_complexity(activity.complexity_level) or 3
        activity_effort_hours = effort_hours_for(activity, complexity_level)
        for row in (summary, department_rows_by_id[department_id], user_rows_by_id[owner_id]):
            row["total"] += 1
            row[bucket] += 1
            if bucket != "cancelled":
                row["complexity_total"] += complexity_level
                row["workload_hours"] += activity_effort_hours
                if bucket == "done":
                    row["complexity_done"] += complexity_level

    def add_completion(row):
        applicable = row["total"] - row["cancelled"]
        row["completion"] = round(row["done"] * 100 / applicable) if applicable else 0
        row["complexity_completion"] = round(
            row["complexity_done"] * 100 / row["complexity_total"]
        ) if row["complexity_total"] else 0
        # 50% is completion by number of tasks, 50% is completion by complexity.
        row["quality"] = round((row["completion"] + row["complexity_completion"]) / 2, 1)
        return row

    summary = add_completion(summary)
    department_rows = sorted(
        (add_completion(row) for row in department_rows_by_id.values()),
        key=lambda row: (-row["total"], row["name"]),
    )
    user_rows = sorted(
        (add_completion(row) for row in user_rows_by_id.values()),
        key=lambda row: (-row["total"], row["name"]),
    )
    chart_data = {
        "status_labels": [value["label"] for value in buckets.values()],
        "status_values": [summary[key] for key in buckets],
        "status_colors": [value["color"] for value in buckets.values()],
        "department_labels": [row["name"] for row in department_rows],
        "department_done": [row["done"] for row in department_rows],
        "department_not_done": [row["not_done"] for row in department_rows],
        "department_work": [row["in_progress"] + row["postponed"] for row in department_rows],
        "department_completion": [row["completion"] for row in department_rows],
        "department_complexity_completion": [row["complexity_completion"] for row in department_rows],
        "department_quality": [row["quality"] for row in department_rows],
        "user_labels": [row["name"] for row in user_rows],
        "user_done": [row["done"] for row in user_rows],
        "user_not_done": [row["not_done"] for row in user_rows],
        "user_work": [row["in_progress"] + row["postponed"] for row in user_rows],
    }

    extra_rows_by_key = {}
    for activity in activities:
        raw_values = activity.extra_values
        if isinstance(raw_values, str):
            try:
                raw_values = json.loads(raw_values)
            except (TypeError, ValueError):
                raw_values = []
        if not isinstance(raw_values, list):
            continue
        for value in raw_values:
            if not isinstance(value, dict) or value.get("value") in (None, ""):
                continue
            key = str(value.get("id") or value.get("name") or value.get("rusname") or "unknown")
            row = extra_rows_by_key.setdefault(
                key,
                {
                    "name": value.get("rusname") or value.get("name") or "Дополнительное значение",
                    "type": value.get("type") or "string",
                    "filled": 0,
                    "task_ids": set(),
                    "numeric_count": 0,
                    "numeric_sum": 0.0,
                    "true_count": 0,
                    "false_count": 0,
                },
            )
            row["filled"] += 1
            row["task_ids"].add(activity.id)
            if row["type"] in ("int", "float"):
                try:
                    row["numeric_sum"] += float(str(value["value"]).replace(",", "."))
                    row["numeric_count"] += 1
                except (TypeError, ValueError):
                    pass
            elif row["type"] == "bool":
                if str(value["value"]).lower() in ("true", "1", "yes", "да"):
                    row["true_count"] += 1
                else:
                    row["false_count"] += 1

    extra_rows = []
    for row in extra_rows_by_key.values():
        row["tasks_count"] = len(row.pop("task_ids"))
        row["numeric_average"] = (
            round(row["numeric_sum"] / row["numeric_count"], 2)
            if row["numeric_count"]
            else None
        )
        extra_rows.append(row)
    extra_rows.sort(key=lambda row: (-row["tasks_count"], row["name"]))
    chart_data["extra_labels"] = [row["name"] for row in extra_rows[:10]]
    chart_data["extra_tasks_count"] = [row["tasks_count"] for row in extra_rows[:10]]
    template_name = "activities_analytics_print.html" if request.args.get("print") == "1" else "activities_analytics.html"
    return render_template(
        template_name,
        title="Аналитика задач",
        summary=summary,
        department_rows=department_rows,
        user_rows=user_rows,
        chart_data=chart_data,
        departments=departments,
        users=users,
        date_from=date_from_value,
        date_to=date_to_value,
        selected_department_id=selected_department_id,
        selected_owner_id=selected_owner_id,
        is_manager=is_manager,
        extra_rows=extra_rows,
    )


# --- Журнал графиков ---
SCHEDULE_STATUSES = {
    "planned": {"label": "Запланировано", "class": "bg-secondary"},
    "in_progress": {"label": "В работе", "class": "bg-warning text-dark"},
    "done": {"label": "Выполнено", "class": "bg-success"},
    "not_done": {"label": "Не выполнено", "class": "bg-danger"},
    "cancelled": {"label": "Отменено", "class": "bg-dark"},
}


def schedule_deadline_meta(item, today=None):
    """Return an urgency colour for a schedule item based on its effective due date."""
    today = today or dt.date.today()
    if item.status == "done":
        return {"label": "Выполнено", "color": "#198754", "text_color": "#ffffff"}
    if item.status == "cancelled":
        return {"label": "Отменено", "color": "#6c757d", "text_color": "#ffffff"}
    if item.status == "not_done":
        return {"label": "Не выполнено", "color": "#dc3545", "text_color": "#ffffff"}

    deadline = item.item_date
    if item.activity and item.activity.end_date:
        deadline = item.activity.end_date.date()
    days_left = (deadline - today).days
    if days_left <= 0:
        return {"label": "Просрочено" if days_left < 0 else "Срок сегодня", "color": "#dc3545", "text_color": "#ffffff"}
    if days_left == 1:
        return {"label": "1 день", "color": "#e85d04", "text_color": "#ffffff"}
    if days_left <= 3:
        return {"label": "{} дн.".format(days_left), "color": "#f59f00", "text_color": "#212529"}
    if days_left <= 6:
        return {"label": "{} дн.".format(days_left), "color": "#b6c600", "text_color": "#212529"}
    return {"label": "{} дн.".format(days_left), "color": "#198754", "text_color": "#ffffff"}


app.jinja_env.globals["schedule_deadline_meta"] = schedule_deadline_meta


def can_edit_schedule_entry(entry):
    return current_user.role in ("admin", "superadmin", "deputy") or entry.created_by_id == current_user.id or entry.responsible_id == current_user.id


def schedule_entry_options():
    """Keep the link pickers useful without loading an unlimited journal into a form."""
    users = User.query.filter_by(is_approved=True, is_enabled=True).order_by(User.full_name).all()
    activities = Activity.query.options(joinedload(Activity.owner)).filter(
        Activity.archived_at.is_(None)
    ).order_by(Activity.start_date.desc(), Activity.id.desc()).limit(500).all()
    memos = Letter.query.filter(
        Letter.source_id == ensure_memo_source(),
        Letter.archived_at.is_(None),
    ).order_by(Letter.letter_date.desc(), Letter.id.desc()).limit(500).all()
    return users, activities, memos


@app.get("/schedule-entries")
@login_required
def schedules_page():
    today = dt.date.today()
    first_day = today.replace(day=1)
    next_month = (first_day.replace(day=28) + dt.timedelta(days=4)).replace(day=1)
    last_day = next_month - dt.timedelta(days=1)
    date_from_value = (request.args.get("date_from") or first_day.isoformat()).strip()
    date_to_value = (request.args.get("date_to") or last_day.isoformat()).strip()
    responsible_id = (request.args.get("responsible_id") or "").strip()
    selected_status = (request.args.get("status") or "").strip()
    try:
        date_from = dt.datetime.strptime(date_from_value, "%Y-%m-%d").date()
        date_to = dt.datetime.strptime(date_to_value, "%Y-%m-%d").date()
    except ValueError:
        date_from, date_to = first_day, last_day
        date_from_value, date_to_value = first_day.isoformat(), last_day.isoformat()

    query = ScheduleEntry.query.options(
        joinedload(ScheduleEntry.responsible), joinedload(ScheduleEntry.creator),
        joinedload(ScheduleEntry.activity), joinedload(ScheduleEntry.memo),
    ).filter(ScheduleEntry.schedule_date >= date_from, ScheduleEntry.schedule_date <= date_to)
    if responsible_id.isdigit():
        query = query.filter(ScheduleEntry.responsible_id == int(responsible_id))
    else:
        responsible_id = ""
    if selected_status in SCHEDULE_STATUSES:
        query = query.filter(ScheduleEntry.status == selected_status)
    else:
        selected_status = ""
    pagination = query.order_by(ScheduleEntry.schedule_date.desc(), ScheduleEntry.id.desc()).paginate(
        page=request.args.get("page", 1, type=int) or 1, per_page=50, error_out=False
    )
    users = User.query.filter_by(is_approved=True, is_enabled=True).order_by(User.full_name).all()
    return render_template(
        "schedules.html", title="Журнал графиков", entries=pagination.items, pagination=pagination,
        users=users, date_from=date_from_value, date_to=date_to_value,
        selected_responsible_id=responsible_id, selected_status=selected_status,
        statuses=SCHEDULE_STATUSES,
    )


@app.route("/schedules/create", methods=["GET", "POST"])
@login_required
def schedule_create():
    users, activities, memos = schedule_entry_options()
    if request.method == "POST":
        date_raw = (request.form.get("schedule_date") or "").strip()
        text_value = (request.form.get("text") or "").strip()
        responsible_id = request.form.get("responsible_id", type=int)
        activity_id = request.form.get("activity_id", type=int)
        memo_id = request.form.get("memo_id", type=int)
        try:
            schedule_date = dt.datetime.strptime(date_raw, "%Y-%m-%d").date()
        except ValueError:
            schedule_date = None
        responsible = db.session.get(User, responsible_id) if responsible_id else None
        activity = db.session.get(Activity, activity_id) if activity_id else None
        memo = db.session.get(Letter, memo_id) if memo_id else None
        if not schedule_date or not text_value or not responsible or not responsible.is_approved or not responsible.is_enabled:
            flash("Укажите дату, текст и действующего ответственного.", "warning")
        elif activity and activity.archived_at:
            flash("Нельзя привязать архивную задачу.", "warning")
        elif memo and (memo.archived_at or memo.source_id != ensure_memo_source()):
            flash("Выберите действующую служебную записку.", "warning")
        else:
            entry = ScheduleEntry(
                schedule_date=schedule_date,
                text=text_value,
                responsible_id=responsible.id,
                activity_id=activity.id if activity else None,
                memo_id=memo.id if memo else None,
                created_by_id=current_user.id if current_user.id != -1 else None,
            )
            db.session.add(entry)
            db.session.commit()
            flash("Запись графика создана.", "success")
            return redirect(url_for("schedule_detail", schedule_id=entry.id))
    return render_template(
        "schedule_form.html", title="Создание графика", entry=None, users=users,
        activities=activities, memos=memos, statuses=SCHEDULE_STATUSES, today=dt.date.today(),
    )


@app.route("/schedules/<int:schedule_id>", methods=["GET", "POST"])
@login_required
def schedule_detail(schedule_id):
    entry = ScheduleEntry.query.options(
        joinedload(ScheduleEntry.responsible), joinedload(ScheduleEntry.creator),
        joinedload(ScheduleEntry.activity), joinedload(ScheduleEntry.memo),
    ).get_or_404(schedule_id)
    can_edit = can_edit_schedule_entry(entry)
    users, activities, memos = schedule_entry_options()
    if request.method == "POST":
        if not can_edit:
            abort(403)
        date_raw = (request.form.get("schedule_date") or "").strip()
        text_value = (request.form.get("text") or "").strip()
        responsible_id = request.form.get("responsible_id", type=int)
        status = (request.form.get("status") or "").strip()
        completion_comment = (request.form.get("completion_comment") or "").strip()
        activity_id = request.form.get("activity_id", type=int)
        memo_id = request.form.get("memo_id", type=int)
        try:
            schedule_date = dt.datetime.strptime(date_raw, "%Y-%m-%d").date()
        except ValueError:
            schedule_date = None
        responsible = db.session.get(User, responsible_id) if responsible_id else None
        activity = db.session.get(Activity, activity_id) if activity_id else None
        memo = db.session.get(Letter, memo_id) if memo_id else None
        if not schedule_date or not text_value or not responsible or not responsible.is_approved or not responsible.is_enabled:
            flash("Укажите дату, текст и действующего ответственного.", "warning")
        elif status not in SCHEDULE_STATUSES:
            flash("Выберите корректный статус выполнения.", "warning")
        elif activity and activity.archived_at:
            flash("Нельзя привязать архивную задачу.", "warning")
        elif memo and (memo.archived_at or memo.source_id != ensure_memo_source()):
            flash("Выберите действующую служебную записку.", "warning")
        else:
            entry.schedule_date = schedule_date
            entry.text = text_value
            entry.responsible_id = responsible.id
            entry.status = status
            entry.completion_comment = completion_comment or None
            entry.activity_id = activity.id if activity else None
            entry.memo_id = memo.id if memo else None
            entry.completed_at = dt.datetime.utcnow() if status == "done" else None
            db.session.commit()
            flash("Запись графика сохранена.", "success")
            return redirect(url_for("schedule_detail", schedule_id=entry.id))
    return render_template(
        "schedule_form.html", title="График", entry=entry, users=users, activities=activities,
        memos=memos, statuses=SCHEDULE_STATUSES, can_edit=can_edit, today=dt.date.today(),
        activity_types=ActivityType.query.order_by(ActivityType.name).all(),
        departments=Department.query.order_by(Department.name).all(),
    )


@app.post("/schedules/<int:schedule_id>/create-task")
@login_required
def schedule_create_task(schedule_id):
    entry = db.session.get(ScheduleEntry, schedule_id)
    if not entry:
        abort(404)
    if not can_edit_schedule_entry(entry):
        abort(403)
    title = (request.form.get("title") or "").strip()
    description = (request.form.get("description") or "").strip()
    type_id = request.form.get("type_id", type=int)
    end_raw = (request.form.get("end_date") or "").strip()
    activity_type = db.session.get(ActivityType, type_id) if type_id else None
    if not title or not activity_type:
        flash("Укажите название и тип новой задачи.", "warning")
        return redirect(url_for("schedule_detail", schedule_id=entry.id))
    try:
        end_date = dt.datetime.strptime(end_raw, "%Y-%m-%d").date() if end_raw else entry.schedule_date
    except ValueError:
        flash("Некорректная дата окончания задачи.", "warning")
        return redirect(url_for("schedule_detail", schedule_id=entry.id))
    if end_date < entry.schedule_date:
        flash("Срок задачи не может быть раньше даты графика.", "warning")
        return redirect(url_for("schedule_detail", schedule_id=entry.id))
    activity = Activity(
        type_id=activity_type.id,
        owner_id=entry.responsible_id,
        created_by_id=current_user.id if current_user.id != -1 else None,
        title=title,
        description=description or None,
        start_date=dt.datetime.combine(entry.schedule_date, dt.time.min),
        end_date=dt.datetime.combine(end_date, dt.time.max),
        status="in_progress",
        priority=3,
        complexity_level=3,
    )
    db.session.add(activity)
    db.session.flush()
    entry.activity_id = activity.id
    add_activity_history(activity, "Создание задачи", "Задача создана из графика №{}.".format(entry.id))
    db.session.commit()
    flash("Задача создана и привязана к графику.", "success")
    return redirect(url_for("schedule_detail", schedule_id=entry.id))


@app.post("/schedules/<int:schedule_id>/create-memo")
@login_required
def schedule_create_memo(schedule_id):
    entry = db.session.get(ScheduleEntry, schedule_id)
    if not entry:
        abort(404)
    if not can_edit_schedule_entry(entry):
        abort(403)
    subject = (request.form.get("subject") or "").strip()
    body = (request.form.get("body") or "").strip()
    department_id = request.form.get("department_id", type=int)
    department = db.session.get(Department, department_id) if department_id else None
    if not subject or not department:
        flash("Укажите тему и отдел служебной записки.", "warning")
        return redirect(url_for("schedule_detail", schedule_id=entry.id))
    memo = Letter(
        reg_number=generate_next_memo_reg_number(department.id),
        subject=subject,
        body=body,
        author_id=current_user.id if current_user.id != -1 else None,
        executor_id=entry.responsible_id,
        source_id=ensure_memo_source(),
        letter_date=entry.schedule_date,
    )
    if not validate_reg_number_uniqueness(memo.reg_number, is_memo=True):
        flash("Не удалось подобрать уникальный номер служебной записки. Повторите попытку.", "warning")
        return redirect(url_for("schedule_detail", schedule_id=entry.id))
    memo.departments.append(department)
    db.session.add(memo)
    db.session.flush()
    db.session.add(LetterRecipient(letter_id=memo.id, user_id=entry.responsible_id, role="executor"))
    entry.memo_id = memo.id
    db.session.commit()
    flash("Служебная записка создана и привязана к графику.", "success")
    return redirect(url_for("schedule_detail", schedule_id=entry.id))


@app.post("/schedules/<int:schedule_id>/delete")
@login_required
def schedule_delete(schedule_id):
    entry = db.session.get(ScheduleEntry, schedule_id)
    if not entry:
        abort(404)
    if current_user.role not in ("admin", "superadmin", "deputy") and entry.created_by_id != current_user.id:
        abort(403)
    db.session.delete(entry)
    db.session.commit()
    flash("Запись графика удалена.", "success")
    return redirect(url_for("schedules_page"))


def can_manage_schedule(schedule):
    return current_user.role in ("admin", "superadmin", "deputy") or schedule.created_by_id == current_user.id


def can_view_schedule(schedule):
    if current_user.role in ("admin", "superadmin", "deputy") or schedule.created_by_id == current_user.id:
        return True
    return any(item.responsible_id == current_user.id for item in schedule.items)


def can_edit_schedule_item(schedule, item):
    return can_manage_schedule(schedule) or item.responsible_id == current_user.id


def can_assign_schedule_responsible(user):
    """Workers create schedule-related documents only for themselves."""
    if not user or not user.is_approved or not user.is_enabled:
        return False
    return current_user.role in ("deputy", "admin", "superadmin") or user.id == current_user.id


def schedule_item_form_values():
    users_query = User.query.filter_by(is_approved=True, is_enabled=True)
    if current_user.role not in ("deputy", "admin", "superadmin"):
        users_query = users_query.filter(User.id == current_user.id)
    users = users_query.order_by(User.full_name).all()
    activities = Activity.query.options(joinedload(Activity.owner)).filter(Activity.archived_at.is_(None)).order_by(
        Activity.start_date.desc(), Activity.id.desc()
    ).limit(500).all()
    memos = Letter.query.filter(Letter.source_id == ensure_memo_source(), Letter.archived_at.is_(None)).order_by(
        Letter.letter_date.desc(), Letter.id.desc()
    ).limit(500).all()
    return users, activities, memos


def build_schedule_gantt(schedule):
    """Build lightweight Gantt coordinates without an external JavaScript library."""
    rows = []
    boundaries = [schedule.schedule_date]
    for item in schedule.items:
        start = item.item_date
        end = item.item_date
        if item.activity:
            if item.activity.start_date:
                start = item.activity.start_date.date()
            if item.activity.end_date:
                end = item.activity.end_date.date()
        if end < start:
            end = start
        boundaries.extend((start, end))
        rows.append({"item": item, "start": start, "end": end, "deadline": schedule_deadline_meta(item)})
    range_start, range_end = min(boundaries), max(boundaries)
    total_days = max((range_end - range_start).days + 1, 1)
    for row in rows:
        row["left"] = round(((row["start"] - range_start).days / total_days) * 100, 3)
        row["width"] = max(round((((row["end"] - row["start"]).days + 1) / total_days) * 100, 3), 2.5)
        row["status"] = SCHEDULE_STATUSES.get(row["item"].status, SCHEDULE_STATUSES["planned"])
    labels = []
    cursor = range_start
    while cursor <= range_end:
        if cursor == range_start or cursor.day == 1 or cursor.weekday() == 0:
            labels.append({
                "left": round(((cursor - range_start).days / total_days) * 100, 3),
                "label": cursor.strftime("%d.%m"),
            })
        cursor += dt.timedelta(days=1)
    return {"start": range_start, "end": range_end, "days": total_days, "rows": rows, "labels": labels}


@app.route("/schedules/journal", methods=["GET", "POST"])
@login_required
def schedules_journal_page():
    if request.method == "POST":
        title = (request.form.get("title") or "").strip()
        description = (request.form.get("description") or "").strip()
        date_raw = (request.form.get("schedule_date") or "").strip()
        try:
            schedule_date = dt.datetime.strptime(date_raw, "%Y-%m-%d").date() if date_raw else dt.date.today()
        except ValueError:
            schedule_date = None
        if not title or not schedule_date:
            flash("Укажите название и дату графика.", "warning")
        else:
            schedule = Schedule(
                title=title, schedule_date=schedule_date, description=description or None,
                created_by_id=current_user.id if current_user.id != -1 else None,
            )
            db.session.add(schedule)
            db.session.commit()
            flash("График создан. Добавьте в него пункты.", "success")
            return redirect(url_for("schedule_document_detail", schedule_id=schedule.id))

    date_from_value = (request.args.get("date_from") or "").strip()
    date_to_value = (request.args.get("date_to") or "").strip()
    query = Schedule.query.options(joinedload(Schedule.creator), joinedload(Schedule.items))
    if current_user.role not in ("admin", "superadmin", "deputy"):
        query = query.filter(or_(
            Schedule.created_by_id == current_user.id,
            Schedule.items.any(ScheduleItem.responsible_id == current_user.id),
        ))
    if date_from_value:
        try:
            query = query.filter(Schedule.schedule_date >= dt.datetime.strptime(date_from_value, "%Y-%m-%d").date())
        except ValueError:
            date_from_value = ""
    if date_to_value:
        try:
            query = query.filter(Schedule.schedule_date <= dt.datetime.strptime(date_to_value, "%Y-%m-%d").date())
        except ValueError:
            date_to_value = ""
    pagination = query.order_by(Schedule.schedule_date.desc(), Schedule.id.desc()).paginate(
        page=request.args.get("page", 1, type=int) or 1, per_page=30, error_out=False
    )
    return render_template(
        "schedules_journal.html", title="Журнал графиков", schedules=pagination.items, pagination=pagination,
        date_from=date_from_value, date_to=date_to_value, today=dt.date.today(), statuses=SCHEDULE_STATUSES,
    )


@app.route("/schedules/journal/<int:schedule_id>", methods=["GET", "POST"])
@login_required
def schedule_document_detail(schedule_id):
    schedule = Schedule.query.options(
        joinedload(Schedule.creator),
        joinedload(Schedule.items).joinedload(ScheduleItem.responsible),
        joinedload(Schedule.items).joinedload(ScheduleItem.activity),
        joinedload(Schedule.items).joinedload(ScheduleItem.memo),
    ).get_or_404(schedule_id)
    if not can_view_schedule(schedule):
        abort(403)
    can_manage = can_manage_schedule(schedule)
    users, activities, memos = schedule_item_form_values()

    if request.method == "POST":
        if not can_manage:
            abort(403)
        title = (request.form.get("title") or "").strip()
        description = (request.form.get("description") or "").strip()
        date_raw = (request.form.get("schedule_date") or "").strip()
        try:
            schedule_date = dt.datetime.strptime(date_raw, "%Y-%m-%d").date()
        except ValueError:
            schedule_date = None
        if not title or not schedule_date:
            flash("Укажите название и дату графика.", "warning")
        else:
            schedule.title, schedule.description, schedule.schedule_date = title, description or None, schedule_date
            db.session.commit()
            flash("График сохранён.", "success")
            return redirect(url_for("schedule_document_detail", schedule_id=schedule.id))
    return render_template(
        "schedule_document_detail.html", title="График", schedule=schedule, can_manage=can_manage,
        users=users, activities=activities, memos=memos, activity_types=ActivityType.query.order_by(ActivityType.name).all(),
        departments=Department.query.order_by(Department.name).all(), statuses=SCHEDULE_STATUSES,
        gantt=build_schedule_gantt(schedule),
    )


@app.post("/schedules/journal/<int:schedule_id>/items")
@login_required
def schedule_document_add_item(schedule_id):
    schedule = db.session.get(Schedule, schedule_id)
    if not schedule:
        abort(404)
    if not can_manage_schedule(schedule):
        abort(403)
    item_date_raw = (request.form.get("item_date") or "").strip()
    text_value = (request.form.get("text") or "").strip()
    responsible_id = request.form.get("responsible_id", type=int)
    activity_id = request.form.get("activity_id", type=int)
    memo_id = request.form.get("memo_id", type=int)
    try:
        item_date = dt.datetime.strptime(item_date_raw, "%Y-%m-%d").date()
    except ValueError:
        item_date = None
    responsible = db.session.get(User, responsible_id) if responsible_id else None
    activity = db.session.get(Activity, activity_id) if activity_id else None
    memo = db.session.get(Letter, memo_id) if memo_id else None
    if not item_date or not text_value or not can_assign_schedule_responsible(responsible):
        flash("Укажите дату, текст и доступного вам ответственного.", "warning")
    elif activity and activity.archived_at:
        flash("Нельзя привязать архивную задачу.", "warning")
    elif memo and (memo.archived_at or memo.source_id != ensure_memo_source()):
        flash("Выберите действующую служебную записку.", "warning")
    else:
        position = max((item.position for item in schedule.items), default=0) + 1
        db.session.add(ScheduleItem(
            schedule_id=schedule.id, position=position, item_date=item_date, text=text_value,
            responsible_id=responsible.id, activity_id=activity.id if activity else None,
            memo_id=memo.id if memo else None,
        ))
        db.session.commit()
        flash("Пункт добавлен в график.", "success")
    return redirect(url_for("schedule_document_detail", schedule_id=schedule.id))


@app.post("/schedules/journal/<int:schedule_id>/items/<int:item_id>")
@login_required
def schedule_document_update_item(schedule_id, item_id):
    schedule = db.session.get(Schedule, schedule_id)
    item = ScheduleItem.query.filter_by(id=item_id, schedule_id=schedule_id).first_or_404()
    if not schedule or not can_edit_schedule_item(schedule, item):
        abort(403)
    status = (request.form.get("status") or "").strip()
    comment = (request.form.get("completion_comment") or "").strip()
    if status not in SCHEDULE_STATUSES:
        abort(400)
    item.status = status
    item.completion_comment = comment or None
    item.completed_at = dt.datetime.utcnow() if status == "done" else None
    db.session.commit()
    flash("Результат выполнения сохранён.", "success")
    return redirect(url_for("schedule_document_detail", schedule_id=schedule.id))


@app.post("/schedules/journal/<int:schedule_id>/items/<int:item_id>/delete")
@login_required
def schedule_document_delete_item(schedule_id, item_id):
    schedule = db.session.get(Schedule, schedule_id)
    item = ScheduleItem.query.filter_by(id=item_id, schedule_id=schedule_id).first_or_404()
    if not schedule or not can_manage_schedule(schedule):
        abort(403)
    db.session.delete(item)
    db.session.commit()
    flash("Пункт удалён из графика.", "success")
    return redirect(url_for("schedule_document_detail", schedule_id=schedule.id))


@app.post("/schedules/journal/<int:schedule_id>/create-task")
@login_required
def schedule_document_create_task(schedule_id):
    schedule = db.session.get(Schedule, schedule_id)
    if not schedule or not can_manage_schedule(schedule):
        abort(403)
    item_id = request.form.get("item_id", type=int)
    item = ScheduleItem.query.filter_by(id=item_id, schedule_id=schedule.id).first() if item_id else None
    if not item:
        flash("Выберите пункт графика, для которого создаётся задача.", "warning")
        return redirect(url_for("schedule_document_detail", schedule_id=schedule.id))
    if item.activity_id:
        flash("К этому пункту уже привязана задача.", "warning")
        return redirect(url_for("schedule_document_detail", schedule_id=schedule.id))
    title = (request.form.get("title") or "").strip()
    type_id = request.form.get("type_id", type=int)
    responsible_id = request.form.get("responsible_id", type=int)
    end_raw = (request.form.get("end_date") or "").strip()
    activity_type = db.session.get(ActivityType, type_id) if type_id else None
    responsible = db.session.get(User, responsible_id) if responsible_id else None
    try:
        end_date = dt.datetime.strptime(end_raw, "%Y-%m-%d").date() if end_raw else item.item_date
    except ValueError:
        end_date = None
    if not title or not activity_type or not can_assign_schedule_responsible(responsible) or not end_date:
        flash("Укажите название, тип, доступного вам ответственного и срок задачи.", "warning")
        return redirect(url_for("schedule_document_detail", schedule_id=schedule.id))
    activity = Activity(type_id=activity_type.id, owner_id=responsible.id, created_by_id=current_user.id,
        title=title, start_date=dt.datetime.combine(item.item_date, dt.time.min),
        end_date=dt.datetime.combine(end_date, dt.time.max), status="in_progress", priority=3, complexity_level=3)
    db.session.add(activity)
    db.session.flush()
    item.activity_id = activity.id
    add_activity_history(activity, "Создание задачи", "Задача создана из графика №{}.".format(schedule.id))
    db.session.commit()
    flash("Задача создана и привязана к пункту графика.", "success")
    return redirect(url_for("schedule_document_detail", schedule_id=schedule.id))


@app.post("/schedules/journal/<int:schedule_id>/create-memo")
@login_required
def schedule_document_create_memo(schedule_id):
    schedule = db.session.get(Schedule, schedule_id)
    if not schedule or not can_manage_schedule(schedule):
        abort(403)
    item_id = request.form.get("item_id", type=int)
    item = ScheduleItem.query.filter_by(id=item_id, schedule_id=schedule.id).first() if item_id else None
    if not item:
        flash("Выберите пункт графика, для которого создаётся служебная записка.", "warning")
        return redirect(url_for("schedule_document_detail", schedule_id=schedule.id))
    if item.memo_id:
        flash("К этому пункту уже привязана служебная записка.", "warning")
        return redirect(url_for("schedule_document_detail", schedule_id=schedule.id))
    subject = (request.form.get("subject") or "").strip()
    department_id = request.form.get("department_id", type=int)
    responsible_id = request.form.get("responsible_id", type=int)
    department = db.session.get(Department, department_id) if department_id else None
    responsible = db.session.get(User, responsible_id) if responsible_id else None
    if not subject or not department or not can_assign_schedule_responsible(responsible):
        flash("Укажите тему, отдел и доступного вам ответственного служебной записки.", "warning")
        return redirect(url_for("schedule_document_detail", schedule_id=schedule.id))
    memo = Letter(reg_number=generate_next_memo_reg_number(department.id), subject=subject,
        author_id=current_user.id, executor_id=responsible.id, source_id=ensure_memo_source(), letter_date=item.item_date)
    if not validate_reg_number_uniqueness(memo.reg_number, is_memo=True):
        flash("Не удалось подобрать уникальный номер служебной записки. Повторите попытку.", "warning")
        return redirect(url_for("schedule_document_detail", schedule_id=schedule.id))
    memo.departments.append(department)
    db.session.add(memo)
    db.session.flush()
    db.session.add(LetterRecipient(letter_id=memo.id, user_id=responsible.id, role="executor"))
    item.memo_id = memo.id
    db.session.commit()
    flash("Служебная записка создана и привязана к пункту графика.", "success")
    return redirect(url_for("schedule_document_detail", schedule_id=schedule.id))


@app.post("/schedules/journal/<int:schedule_id>/delete")
@login_required
def schedule_document_delete(schedule_id):
    schedule = db.session.get(Schedule, schedule_id)
    if not schedule:
        abort(404)
    if not can_manage_schedule(schedule):
        abort(403)
    db.session.delete(schedule)
    db.session.commit()
    flash("График удалён.", "success")
    return redirect(url_for("schedules_journal_page"))


@app.get("/schedules")
@login_required
def schedules_redirect():
    return redirect(url_for("schedules_journal_page"))


# --- Протоколы: административные документы с перечнем обычных задач ---
def can_manage_protocol(protocol):
    return current_user.role == "superadmin" or protocol.created_by_id == current_user.id


@app.route("/protocols", methods=["GET", "POST"])
@role_required("admin", "superadmin")
def protocols_page():
    if request.method == "POST":
        title = (request.form.get("title") or "").strip()
        description = (request.form.get("description") or "").strip()
        date_raw = (request.form.get("protocol_date") or "").strip()
        if not title:
            flash("Укажите название протокола.", "warning")
            return redirect(url_for("protocols_page"))
        try:
            protocol_date = dt.datetime.strptime(date_raw, "%Y-%m-%d").date() if date_raw else dt.date.today()
        except ValueError:
            flash("Некорректная дата протокола.", "warning")
            return redirect(url_for("protocols_page"))
        protocol = Protocol(
            title=title,
            description=description or None,
            protocol_date=protocol_date,
            created_by_id=current_user.id if current_user.id != -1 else None,
        )
        db.session.add(protocol)
        db.session.commit()
        flash("Протокол создан. Теперь добавьте в него задачи.", "success")
        return redirect(url_for("protocol_detail", protocol_id=protocol.id))

    query = Protocol.query.options(joinedload(Protocol.creator), joinedload(Protocol.items))
    if current_user.role != "superadmin":
        query = query.filter(Protocol.created_by_id == current_user.id)
    protocols = query.order_by(Protocol.protocol_date.desc(), Protocol.id.desc()).all()
    return render_template("protocols.html", title="Протоколы", protocols=protocols, today=dt.date.today())


@app.get("/protocols/<int:protocol_id>")
@role_required("admin", "superadmin")
def protocol_detail(protocol_id):
    protocol = Protocol.query.options(
        joinedload(Protocol.creator),
        joinedload(Protocol.items).joinedload(ProtocolItem.activity).joinedload(Activity.owner),
    ).get_or_404(protocol_id)
    if not can_manage_protocol(protocol):
        abort(403)

    attached_activity_ids = [item.activity_id for item in protocol.items]
    available_query = Activity.query.options(joinedload(Activity.owner)).filter(
        Activity.archived_at.is_(None)
    )
    if attached_activity_ids:
        available_query = available_query.filter(~Activity.id.in_(attached_activity_ids))
    available_activities = available_query.order_by(Activity.start_date.desc(), Activity.id.desc()).limit(500).all()
    return render_template(
        "protocol_detail.html",
        title="Протокол №{}".format(protocol.id),
        protocol=protocol,
        available_activities=available_activities,
        activity_types=ActivityType.query.order_by(ActivityType.name).all(),
        users=User.query.filter_by(is_approved=True).order_by(User.full_name).all(),
        get_status_meta=get_activity_status_meta,
    )


@app.post("/protocols/<int:protocol_id>/items")
@role_required("admin", "superadmin")
def protocol_add_item(protocol_id):
    protocol = Protocol.query.get_or_404(protocol_id)
    if not can_manage_protocol(protocol):
        abort(403)

    activity_id = request.form.get("activity_id", type=int)
    activity = db.session.get(Activity, activity_id) if activity_id else None
    if not activity or activity.archived_at:
        flash("Выберите действующую задачу.", "warning")
        return redirect(url_for("protocol_detail", protocol_id=protocol.id))
    if ProtocolItem.query.filter_by(protocol_id=protocol.id, activity_id=activity.id).first():
        flash("Эта задача уже добавлена в протокол.", "info")
        return redirect(url_for("protocol_detail", protocol_id=protocol.id))

    position = max((item.position for item in protocol.items), default=0) + 1
    db.session.add(ProtocolItem(
        protocol_id=protocol.id,
        activity_id=activity.id,
        position=position,
        due_date=activity.end_date.date() if activity.end_date else None,
        comment=(request.form.get("comment") or "").strip() or None,
    ))
    db.session.commit()
    flash("Задача добавлена в протокол.", "success")
    return redirect(url_for("protocol_detail", protocol_id=protocol.id))


@app.post("/protocols/<int:protocol_id>/create-task")
@role_required("admin", "superadmin")
def protocol_create_task(protocol_id):
    """Create a regular task and immediately attach it to this protocol."""
    protocol = Protocol.query.get_or_404(protocol_id)
    if not can_manage_protocol(protocol):
        abort(403)

    title = (request.form.get("title") or "").strip()
    description = (request.form.get("description") or "").strip()
    type_id = request.form.get("type_id", type=int)
    owner_id = request.form.get("owner_id", type=int)
    start_raw = (request.form.get("start_date") or "").strip()
    end_raw = (request.form.get("end_date") or "").strip()
    priority = parse_activity_priority(request.form.get("priority")) or 3
    complexity_level = parse_activity_complexity(request.form.get("complexity_level")) or 3
    activity_type = db.session.get(ActivityType, type_id) if type_id else None
    owner = db.session.get(User, owner_id) if owner_id else None
    if not title or not activity_type or not owner or not end_raw:
        flash("Укажите название, тип, исполнителя и срок задачи.", "warning")
        return redirect(url_for("protocol_detail", protocol_id=protocol.id))
    try:
        start_date = dt.datetime.strptime(start_raw, "%Y-%m-%d").date() if start_raw else dt.date.today()
        end_date = dt.datetime.strptime(end_raw, "%Y-%m-%d").date()
    except ValueError:
        flash("Некорректная дата задачи.", "warning")
        return redirect(url_for("protocol_detail", protocol_id=protocol.id))
    if end_date < start_date:
        flash("Срок выполнения не может быть раньше даты начала.", "warning")
        return redirect(url_for("protocol_detail", protocol_id=protocol.id))

    activity = Activity(
        type_id=activity_type.id,
        owner_id=owner.id,
        created_by_id=current_user.id if current_user.id != -1 else None,
        title=title,
        description=description or None,
        start_date=dt.datetime.combine(start_date, dt.time.min),
        end_date=dt.datetime.combine(end_date, dt.time.max),
        status="in_progress",
        priority=priority,
        complexity_level=complexity_level,
    )
    db.session.add(activity)
    db.session.flush()
    position = max((item.position for item in protocol.items), default=0) + 1
    db.session.add(ProtocolItem(
        protocol_id=protocol.id,
        activity_id=activity.id,
        position=position,
        due_date=end_date,
        comment="Создана из протокола.",
    ))
    add_activity_history(activity, "Создание задачи", "Задача создана из протокола №{}.".format(protocol.id))
    db.session.commit()
    flash("Задача создана и добавлена в протокол.", "success")
    return redirect(url_for("protocol_detail", protocol_id=protocol.id))


@app.post("/protocols/<int:protocol_id>/items/<int:item_id>/delete")
@role_required("admin", "superadmin")
def protocol_delete_item(protocol_id, item_id):
    protocol = Protocol.query.get_or_404(protocol_id)
    if not can_manage_protocol(protocol):
        abort(403)
    item = ProtocolItem.query.filter_by(id=item_id, protocol_id=protocol.id).first_or_404()
    db.session.delete(item)
    db.session.commit()
    flash("Задача исключена из протокола.", "success")
    return redirect(url_for("protocol_detail", protocol_id=protocol.id))


@app.post("/protocols/<int:protocol_id>/delete")
@role_required("admin", "superadmin")
def protocol_delete(protocol_id):
    protocol = Protocol.query.get_or_404(protocol_id)
    if not can_manage_protocol(protocol):
        abort(403)
    db.session.delete(protocol)
    db.session.commit()
    flash("Протокол удалён.", "success")
    return redirect(url_for("protocols_page"))


@app.post("/activities")
@login_required
def  add_activity():
    title = request.form.get("title", "Активность")
    type_id = request.form.get("type_id", type=int)
    description = request.form.get("description")

    if not type_id:
        flash("Выберите тип активности")
        return redirect(url_for("activities_page"))

    a = Activity(
        type_id=type_id,
        owner_id=current_user.id,
        created_by_id=current_user.id if current_user.id != -1 else None,
        title=title,
        description=description,
        start_date=dt.datetime.utcnow(),
        status="in_progress"
    )
    db.session.add(a)
    db.session.commit()
    flash(f"Активность #{a.id} создана")
    return redirect(url_for("activities_page"))


@app.get("/activities/<int:activity_id>")
@login_required
def activity_view(activity_id: int):
    normalize_activity_statuses()
    a = (
        Activity.query
        .options(
            joinedload(Activity.logs).joinedload(ActivityLog.children),
            joinedload(Activity.letters),
            joinedload(Activity.assigned_department),
            joinedload(Activity.completion_memo).joinedload(Letter.source),
        )
        .get_or_404(activity_id)
    )
    # --- проверка прав доступа ---
    is_owner = current_user.id == a.owner_id
    is_admin = current_user.role in ("manager", "admin", "superadmin", "deputy")
    is_department_assignee = (
        a.assigned_department_id is not None
        and a.assigned_department_id == getattr(current_user, "department_id", None)
    )

    if not (is_owner or is_admin or is_department_assignee):
        flash("❌ У вас нет доступа к этой активности", "danger")
        return redirect(url_for("index"))
    # --- проверка прав ---
    can_edit = (
        current_user.id == a.owner_id
        or current_user.role in ("manager", "admin", "superadmin", "deputy")
        or is_department_assignee
    )
    status_locked = activity_work_is_locked(a)
    can_extend_deadline = bool(
        can_edit and a.status == "not_done" and a.end_date
        and a.end_date.date() < dt.date.today()
    )
    can_modify_work = can_edit and not status_locked
    can_change_status = can_edit and (
        not status_locked
        or current_user.role in ("admin", "superadmin", "deputy")
        or can_extend_deadline
    )
    can_delegate = current_user.role in ("deputy", "admin", "superadmin")
    closure_memos_query = Letter.query.filter(
        Letter.source_id == ensure_memo_source(),
        Letter.archived_at.is_(None),
    )
    if current_user.role not in ("manager", "admin", "superadmin"):
        closure_memos_query = (
            closure_memos_query
            .outerjoin(Letter.recipients)
            .outerjoin(Letter.departments)
            .filter(or_(
                Letter.author_id == current_user.id,
                Letter.executor_id == current_user.id,
                LetterRecipient.user_id == current_user.id,
                Department.id == getattr(current_user, "department_id", None),
            ))
            .distinct()
        )
    closure_memos = closure_memos_query.order_by(
        Letter.letter_date.desc().nullslast(), Letter.id.desc()
    ).all()
    linked_plan = None
    if a.plan_id:
        linked_plan = db.session.get(Plan, a.plan_id)
    if linked_plan is None:
        linked_plan = (
            Plan.query.join(PlanItem)
            .filter(PlanItem.linked_activity_id == a.id)
            .first()
        )
    plan_add_candidate = None
    if is_owner and not a.archived_at and linked_plan is None:
        plan_add_candidate = find_matching_plan_for_activity(a)

    return render_template(
        "activity_view.html",
        title=f"Активность #{a.id}",
        a=a,
        can_edit=can_edit,   # 👈 передаём флаг в шаблон
        can_modify_work=can_modify_work,
        can_change_status=can_change_status,
        can_extend_deadline=can_extend_deadline,
        status_locked=status_locked,
        datetime=datetime,
        status_meta=get_activity_status_meta(a.status),
        can_delegate=can_delegate,
        closure_memos=closure_memos,
        linked_plan=linked_plan,
        plan_add_candidate=plan_add_candidate,
        has_pending_ocr=any(document.ocr_status in ("queued", "processing") for document in a.documents),
        delegation_users=User.query.filter_by(is_approved=True).order_by(User.full_name).all() if can_delegate else [],
        delegation_departments=Department.query.order_by(Department.name).all() if can_delegate else [],
        history=ActivityHistory.query.filter_by(activity_id=a.id).order_by(ActivityHistory.created_at.desc()).all(),
    )


@app.post("/activities/<int:activity_id>/add-to-plan")
@login_required
def add_activity_to_plan(activity_id):
    """Lets an executor attach an existing task to their monthly plan."""
    activity = Activity.query.get_or_404(activity_id)
    if activity.owner_id != current_user.id or activity.archived_at:
        abort(403)

    plan = find_matching_plan_for_activity(activity)
    if plan is None:
        flash("План исполнителя на месяц даты начала задачи пока не создан.", "warning")
        return redirect(url_for("activity_view", activity_id=activity.id))

    existing_item = PlanItem.query.filter_by(
        plan_id=plan.id, linked_activity_id=activity.id
    ).first()
    if existing_item:
        activity.plan_id = plan.id
        activity.plan_requested = False
        db.session.commit()
        flash("Задача уже находится в этом плане.", "info")
        return redirect(url_for("plan_detail", plan_id=plan.id))

    position = (
        db.session.query(func.max(PlanItem.position))
        .filter(PlanItem.plan_id == plan.id)
        .scalar()
        or 0
    )
    deadline_date = activity.end_date.date() if activity.end_date else None
    db.session.add(PlanItem(
        plan_id=plan.id,
        executor_id=activity.owner_id,
        position=position + 1,
        task_text=activity.title,
        deadline_kind="date" if deadline_date else "month",
        deadline_date=deadline_date,
        planned_hours=calculate_activity_effort_hours(activity) or 0,
        linked_activity_id=activity.id,
        status="planned",
    ))
    activity.plan_id = plan.id
    activity.plan_requested = False
    add_activity_history(
        activity,
        "Добавление в план",
        "Исполнитель добавил задачу в план «{}».".format(plan.plan_text),
    )
    db.session.commit()
    flash("Задача добавлена в план.", "success")
    return redirect(url_for("plan_detail", plan_id=plan.id))


@app.get("/activities/archive")
@role_required("admin", "superadmin", "deputy")
def activities_archive_page():
    search = (request.args.get("search") or "").strip()
    query = Activity.query.options(joinedload(Activity.owner), joinedload(Activity.type)).filter(
        Activity.archived_at.is_not(None)
    )
    if search:
        like_pattern = "%{}%".format(search)
        query = query.filter(or_(
            Activity.title.ilike(like_pattern),
            Activity.description.ilike(like_pattern),
            Activity.direction.ilike(like_pattern),
        ))
    activities = query.order_by(Activity.archived_at.desc(), Activity.id.desc()).all()
    return render_template(
        "activities_archive.html", title="Архив задач", activities=activities, search=search,
        get_status_meta=get_activity_status_meta, get_priority_meta=get_activity_priority_meta,
    )


@app.post("/activities/<int:activity_id>/archive")
@role_required("admin", "superadmin", "deputy")
def archive_activity(activity_id):
    activity = Activity.query.get_or_404(activity_id)
    activity.archived_at = dt.datetime.utcnow()
    activity.archived_by_id = current_user.id if current_user.id != -1 else None
    add_activity_history(activity, "Архивирование", "Задача перемещена в архив.")
    db.session.commit()
    flash("Задача архивирована.", "success")
    return redirect(request.referrer or url_for("activities_page"))


@app.post("/activities/<int:activity_id>/restore")
@role_required("admin", "superadmin", "deputy")
def restore_activity(activity_id):
    activity = Activity.query.get_or_404(activity_id)
    activity.archived_at = None
    activity.archived_by_id = None
    add_activity_history(activity, "Восстановление", "Задача восстановлена из архива.")
    db.session.commit()
    flash("Задача восстановлена из архива.", "success")
    return redirect(request.referrer or url_for("activity_view", activity_id=activity.id))


@app.get("/activities/audit")
@role_required("manager", "admin", "superadmin", "deputy")
def activities_audit():
    today = dt.date.today()
    default_from = today.replace(day=1)
    date_from_raw = (request.args.get("date_from") or default_from.isoformat()).strip()
    date_to_raw = (request.args.get("date_to") or today.isoformat()).strip()
    activity_id = (request.args.get("activity_id") or "").strip()
    author_id = (request.args.get("author_id") or "").strip()
    action = (request.args.get("action") or "").strip()
    search = (request.args.get("search") or "").strip()
    try:
        date_from = dt.datetime.strptime(date_from_raw, "%Y-%m-%d")
        date_to = dt.datetime.strptime(date_to_raw, "%Y-%m-%d") + dt.timedelta(days=1)
    except ValueError:
        date_from = dt.datetime.combine(default_from, dt.time.min)
        date_to = dt.datetime.combine(today + dt.timedelta(days=1), dt.time.min)
        date_from_raw, date_to_raw = default_from.isoformat(), today.isoformat()

    query = (
        ActivityHistory.query
        .join(Activity, Activity.id == ActivityHistory.activity_id)
        .outerjoin(User, User.id == ActivityHistory.author_id)
        .options(joinedload(ActivityHistory.author))
        .filter(ActivityHistory.created_at >= date_from, ActivityHistory.created_at < date_to)
    )
    if current_user.role == "manager":
        if current_user.department_id:
            query = query.filter(Activity.owner.has(User.department_id == current_user.department_id))
        else:
            query = query.filter(Activity.owner_id == current_user.id)
    if activity_id.isdigit():
        query = query.filter(ActivityHistory.activity_id == int(activity_id))
    if author_id.isdigit():
        query = query.filter(ActivityHistory.author_id == int(author_id))
    if action:
        query = query.filter(ActivityHistory.action == action)
    if search:
        query = query.filter(or_(Activity.title.ilike("%{}%".format(search)), ActivityHistory.details.ilike("%{}%".format(search))))

    rows = query.order_by(ActivityHistory.created_at.desc()).limit(1000).all()
    action_options = [row[0] for row in db.session.query(ActivityHistory.action).distinct().order_by(ActivityHistory.action).all()]
    users_query = User.query.filter_by(is_approved=True).order_by(User.full_name)
    if current_user.role == "manager" and current_user.department_id:
        users_query = users_query.filter(User.department_id == current_user.department_id)
    return render_template(
        "activities_audit.html", title="Аудит задач", rows=rows,
        date_from=date_from_raw, date_to=date_to_raw, activity_id=activity_id,
        author_id=author_id, action=action, search=search,
        action_options=action_options, users=users_query.all(),
    )


@app.post("/activities/<int:activity_id>/redirect")
@role_required("deputy", "admin", "superadmin")
def redirect_activity(activity_id):
    activity = Activity.query.options(joinedload(Activity.owner), joinedload(Activity.assigned_department)).get_or_404(activity_id)
    user_id = (request.form.get("target_user_id") or "").strip()
    department_id = (request.form.get("target_department_id") or "").strip()
    if bool(user_id) == bool(department_id):
        flash("Выберите одного сотрудника или один отдел для перенаправления.", "warning")
        return redirect(url_for("activity_view", activity_id=activity.id))

    previous_target = activity.owner.full_name if activity.owner else "не назначен"
    if activity.assigned_department:
        previous_target += f" (отдел {activity.assigned_department.name})"
    if user_id:
        try:
            target_user = db.session.get(User, int(user_id))
        except ValueError:
            target_user = None
        if not target_user or not target_user.is_approved:
            flash("Выбранный сотрудник недоступен.", "warning")
            return redirect(url_for("activity_view", activity_id=activity.id))
        activity.owner_id = target_user.id
        activity.assigned_department_id = target_user.department_id
        new_target = target_user.full_name
        if target_user.department:
            new_target += f" (отдел {target_user.department.name})"
    else:
        try:
            target_department = db.session.get(Department, int(department_id))
        except ValueError:
            target_department = None
        if not target_department:
            flash("Выбранный отдел не найден.", "warning")
            return redirect(url_for("activity_view", activity_id=activity.id))
        activity.assigned_department_id = target_department.id
        new_target = f"отдел {target_department.name}"

    db.session.add(ActivityLog(
        activity_id=activity.id,
        text=f"Задача перенаправлена: {previous_target} -> {new_target}. Выполнил: {current_user.full_name}.",
        entry_date=dt.datetime.utcnow(),
    ))
    add_activity_history(
        activity,
        "Перенаправление",
        "Задача перенаправлена: {} -> {}.".format(previous_target, new_target),
    )
    db.session.commit()
    flash("Задача перенаправлена. Запись добавлена в журнал.", "success")
    return redirect(url_for("activity_view", activity_id=activity.id))


@app.route("/activities/<int:activity_id>/edit", methods=["GET", "POST"])
@role_required("admin", "superadmin")
def edit_activity(activity_id):
    activity = Activity.query.get_or_404(activity_id)
    users = User.query.filter_by(is_approved=True).order_by(User.full_name).all()
    types = ActivityType.query.order_by(ActivityType.name).all()
    if request.method == "POST":
        previous_values = {
            "title": activity.title or "",
            "description": activity.description or "",
            "owner": activity.owner.full_name if activity.owner else "-",
            "type": activity.type.name if activity.type else "-",
            "start_date": activity.start_date.strftime("%d.%m.%Y") if activity.start_date else "-",
            "end_date": activity.end_date.strftime("%d.%m.%Y") if activity.end_date else "-",
            "status": get_activity_status_meta(activity.status)["label"],
            "priority": activity.priority or 3,
            "complexity": activity.complexity_level or 3,
            "direction": activity.direction or "",
            "outgoing_document": activity.outgoing_document or "",
        }
        title = (request.form.get("title") or "").strip()
        if not title:
            flash("Введите название задачи.", "warning")
            return redirect(request.url)
        try:
            owner_id = int(request.form.get("owner_id") or "")
        except ValueError:
            flash("Выберите исполнителя.", "warning")
            return redirect(request.url)
        owner = db.session.get(User, owner_id)
        if not owner:
            flash("Исполнитель не найден.", "warning")
            return redirect(request.url)
        try:
            start_date_raw = (request.form.get("start_date") or "").strip()
            end_date_raw = (request.form.get("end_date") or "").strip()
            start_date = datetime.strptime(start_date_raw, "%Y-%m-%d") if start_date_raw else activity.start_date
            end_date = datetime.strptime(end_date_raw, "%Y-%m-%d") if end_date_raw else None
        except ValueError:
            flash("Некорректный формат даты.", "warning")
            return redirect(request.url)
        status = (request.form.get("status") or "in_progress").strip()
        if status not in ACTIVITY_STATUS_META:
            flash("Некорректный статус задачи.", "warning")
            return redirect(request.url)
        priority = parse_activity_priority(request.form.get("priority", activity.priority))
        if priority is None:
            flash("Выберите приоритет от 1 до 5.", "warning")
            return redirect(request.url)
        complexity_level = parse_activity_complexity(
            request.form.get("complexity_level", activity.complexity_level)
        )
        if complexity_level is None:
            flash("Выберите уровень сложности от 1 до 5.", "warning")
            return redirect(request.url)
        type_id = (request.form.get("type_id") or "").strip()
        selected_type = db.session.get(ActivityType, int(type_id)) if type_id.isdigit() else None
        activity.title = title
        activity.description = (request.form.get("description") or "").strip()
        activity.owner_id = owner.id
        activity.type_id = selected_type.id if selected_type else None
        activity.start_date = start_date
        activity.end_date = end_date
        activity.status = status
        activity.priority = priority
        activity.complexity_level = complexity_level
        activity.direction = (request.form.get("direction") or "").strip()
        activity.outgoing_document = (request.form.get("outgoing_document") or "").strip()
        updated_values = {
            "Название": (previous_values["title"], activity.title or ""),
            "Описание": (previous_values["description"], activity.description or ""),
            "Исполнитель": (previous_values["owner"], owner.full_name),
            "Тип": (previous_values["type"], selected_type.name if selected_type else "-"),
            "Дата начала": (previous_values["start_date"], activity.start_date.strftime("%d.%m.%Y") if activity.start_date else "-"),
            "Дата окончания": (previous_values["end_date"], activity.end_date.strftime("%d.%m.%Y") if activity.end_date else "-"),
            "Статус": (previous_values["status"], get_activity_status_meta(activity.status)["label"]),
            "Приоритет": (get_activity_priority_meta(previous_values["priority"])["label"], get_activity_priority_meta(activity.priority)["label"]),
            "Сложность": (get_activity_complexity_meta(previous_values["complexity"])["label"], get_activity_complexity_meta(activity.complexity_level)["label"]),
            "Тема": (previous_values["direction"], activity.direction or ""),
            "Обоснования": (previous_values["outgoing_document"], activity.outgoing_document or ""),
        }
        changes = ["{}: «{}» -> «{}»".format(label, before or "-", after or "-") for label, (before, after) in updated_values.items() if before != after]
        if changes:
            add_activity_history(activity, "Редактирование задачи", "; ".join(changes))
        db.session.commit()
        flash("Задача обновлена.", "success")
        return redirect(url_for("activity_view", activity_id=activity.id))
    return render_template(
        "activity_edit.html",
        title=f"Редактирование задачи #{activity.id}",
        activity=activity,
        users=users,
        types=types,
        status_meta=ACTIVITY_STATUS_META,
    )


@app.post("/activities/<int:activity_id>/logs")
@login_required
def add_activity_log(activity_id: int):
    a = Activity.query.get_or_404(activity_id)
    if a.owner_id != current_user.id and current_user.role not in ("manager", "admin", "superadmin", "deputy"):
        abort(403)
    if activity_work_is_locked(a):
        abort(403)

    text = request.form.get("text", "").strip()
    entry_date_str = request.form.get("entry_date")  # <-- берём из формы

    if not text:
        return {"success": False, "error": "Текст пустой"}, 400

    # Парсим дату/время, если передана
    from datetime import datetime as _dt
    entry_dt = None
    if entry_date_str:
        try:
            entry_dt = _dt.strptime(entry_date_str, "%Y-%m-%dT%H:%M")
        except ValueError:
            entry_dt = _dt.utcnow()

    log = ActivityLog(
        activity_id=a.id,
        text=text,
        entry_date=entry_dt or _dt.utcnow()
    )
    db.session.add(log)
    add_activity_history(a, "Журнал действий", "Добавлено действие: {}".format(text[:300]))
    db.session.commit()

    return redirect(url_for("activity_view", activity_id=a.id))

import uuid
from werkzeug.utils import secure_filename
import mimetypes



@app.post("/activities/<int:activity_id>/upload")
@login_required
def upload_activity_document(activity_id):
    a = Activity.query.get_or_404(activity_id)
    if activity_work_is_locked(a) or not can_modify_activity_work(a):
        abort(403)
    file = request.files.get("file")

    if not file or not file.filename:
        flash("⚠️ Файл не выбран", "warning")
        return redirect(url_for("activity_view", activity_id=a.id))

    # Проверяем расширение
    allowed_ext = {"png", "jpg", "jpeg", "pdf", "txt", "docx"}
    ext = file.filename.rsplit(".", 1)[-1].lower()
    if ext not in allowed_ext:
        flash(f"🚫 Недопустимый формат файла: .{ext}", "danger")
        return redirect(url_for("activity_view", activity_id=a.id))

    # Проверяем MIME-тип
    mime, _ = mimetypes.guess_type(file.filename)
    if mime and not mime.startswith(("image/", "application/", "text/")):
        flash("🚫 Неверный тип файла", "danger")
        return redirect(url_for("activity_view", activity_id=a.id))

    # Создаём папку если нет
    upload_dir = os.path.join(app.root_path, "static", "uploads")
    os.makedirs(upload_dir, exist_ok=True)

    # Генерируем безопасное имя
    filename = secure_filename_rus(file.filename)
    random_name = f"{uuid.uuid4().hex}_{filename}"
    save_path = os.path.join(upload_dir, random_name)
    file.save(save_path)

    # Ограничиваем размер до 10 МБ
    max_size = 10 * 1024 * 1024
    if os.path.getsize(save_path) > max_size:
        os.remove(save_path)
        flash("🚫 Файл слишком большой (макс. 10 МБ)", "danger")
        return redirect(url_for("activity_view", activity_id=a.id))

    # Save first, then let the single OCR worker process it in the background.
    doc = ActivityDocument(
        activity_id=a.id,
        filename=filename,
        filepath=f"uploads/{random_name}",  # 🔗 путь внутри static/
        ocr_status="queued",
        uploaded_by_id=current_user.id if current_user.id != -1 else None,
    )
    db.session.add(doc)
    add_activity_history(a, "Документ", "Добавлен документ: {}.".format(filename))
    db.session.commit()
    enqueue_document_ocr("activity", doc.id)

    flash("Документ загружен и поставлен в очередь распознавания.", "success")
    return redirect(url_for("activity_view", activity_id=a.id))
def has_document_access(user_id: int, doc_type: str, doc_id: int) -> bool:
    """Проверка прав на документ активности или письма"""
    if current_user.role in ("admin", "manager"):
        return True

    if doc_type == "activity":
        doc = ActivityDocument.query.get(doc_id)
        if not doc:
            return False
        return (doc.activity.owner_id == user_id) or (
            DocumentAccess.query.filter_by(doc_type="activity", doc_id=doc_id, user_id=user_id).count() > 0
        )

    if doc_type == "letter":
        doc = LetterDocument.query.get(doc_id)
        if not doc:
            return False
        return (doc.letter.author_id == user_id) or (
            DocumentAccess.query.filter_by(doc_type="letter", doc_id=doc_id, user_id=user_id).count() > 0
        )

    return False

@app.get("/documents/<int:doc_id>")
@login_required
def download_document(doc_id: int):
    d = ActivityDocument.query.get_or_404(doc_id)
    return send_from_directory(app.config["UPLOAD_FOLDER"], os.path.basename(d.filepath), as_attachment=True)

@app.post("/activities/<int:activity_id>/done")
@login_required
def mark_activity_done(activity_id: int):
    a = Activity.query.get_or_404(activity_id)
    if a.owner_id != current_user.id and current_user.role not in ("manager", "admin"):
        abort(403)
    a.status = "done"
    if not a.end_date:
        a.end_date = dt.datetime.utcnow()
    db.session.commit()
    flash("Отмечено как выполнено")
    return redirect(url_for("activity_view", activity_id=activity_id))

@app.post("/activities/<int:activity_id>/status")
@login_required
def update_activity_status(activity_id: int):
    a = Activity.query.get_or_404(activity_id)
    if not can_modify_activity_work(a):
        abort(403)

    new_status = request.form.get("status", "").strip()
    comment = request.form.get("comment", "").strip()
    completion_memo_id = (request.form.get("completion_memo_id") or "").strip()
    if a.status in ("done", "not_done") and current_user.role not in ("admin", "superadmin", "deputy"):
        if new_status != "extend":
            abort(403)
    allowed_statuses = {"in_progress", "done", "not_done", "postponed", "cancelled", "extend"}
    if new_status not in allowed_statuses:
        flash("Недопустимый статус", "warning")
        return redirect(url_for("activity_view", activity_id=activity_id))

    postponed_to = None
    if new_status == "not_done" and not comment:
        flash("Укажите причину невыполнения", "warning")
        return redirect(url_for("activity_view", activity_id=activity_id))
    if new_status in ("postponed", "extend"):
        postponed_to_raw = request.form.get("postponed_to", "").strip()
        if not postponed_to_raw or not comment:
            flash("Для продления укажите новую дату и причину", "warning")
            return redirect(url_for("activity_view", activity_id=activity_id))
        try:
            postponed_to = dt.datetime.strptime(postponed_to_raw, "%Y-%m-%d")
        except ValueError:
            flash("Некорректная дата переноса", "warning")
            return redirect(url_for("activity_view", activity_id=activity_id))
        if new_status == "extend":
            deadline = a.end_date.date() if a.end_date else None
            if a.status != "not_done" or not deadline or deadline >= dt.date.today():
                flash("Продлить можно только просроченную задачу.", "warning")
                return redirect(url_for("activity_view", activity_id=activity_id))
            if postponed_to.date() < dt.date.today():
                flash("Новый срок не может быть в прошлом.", "warning")
                return redirect(url_for("activity_view", activity_id=activity_id))

    completion_memo = None
    if new_status == "done" and completion_memo_id:
        if not completion_memo_id.isdigit():
            flash("Выберите служебную записку из списка.", "warning")
            return redirect(url_for("activity_view", activity_id=activity_id))
        completion_memo = Letter.query.filter(
            Letter.id == int(completion_memo_id),
            Letter.source_id == ensure_memo_source(),
            Letter.archived_at.is_(None),
        ).first()
        if not completion_memo:
            flash("Выбранная служебная записка не найдена или находится в архиве.", "warning")
            return redirect(url_for("activity_view", activity_id=activity_id))
        memo_access = (
            current_user.role in ("manager", "admin", "superadmin")
            or completion_memo.author_id == current_user.id
            or completion_memo.executor_id == current_user.id
            or any(recipient.user_id == current_user.id for recipient in completion_memo.recipients)
            or any(department.id == getattr(current_user, "department_id", None) for department in completion_memo.departments)
        )
        if not memo_access:
            abort(403)

    previous_status = a.status
    stored_status = "postponed" if new_status == "extend" else new_status
    a.status = stored_status
    a.completion_memo_id = completion_memo.id if new_status == "done" and completion_memo else None
    if new_status == "done" and not a.end_date:
        a.end_date = dt.datetime.utcnow()
    if stored_status == "postponed":
        a.postponed_to = postponed_to
        a.end_date = postponed_to
    elif new_status == "in_progress":
        a.postponed_to = None
    a.status_comment = comment or None
    status_text = "Продление срока" if new_status == "extend" else get_activity_status_meta(stored_status)["label"]
    log_text = "Статус изменён: {} -> {}.".format(
        get_activity_status_meta(previous_status)["label"],
        status_text,
    )
    if stored_status == "postponed":
        log_text += " Новый срок: {}.".format(postponed_to.strftime("%d.%m.%Y"))
    if comment:
        log_text += " Комментарий: {}".format(comment)
    if completion_memo:
        log_text += " Задача закрыта служебной запиской №{}.".format(
            completion_memo.reg_number or completion_memo.id
        )
    db.session.add(ActivityLog(
        activity_id=a.id,
        text=log_text,
        entry_date=dt.datetime.utcnow(),
    ))
    add_activity_history(a, "Смена статуса", log_text)
    db.session.commit()
    flash("Статус задачи обновлён", "success")
    return redirect(url_for("activity_view", activity_id=activity_id))

@app.post("/activities/<int:activity_id>/approve")
@role_required("manager", "admin")
def approve_activity(activity_id: int):
    a = Activity.query.get_or_404(activity_id)
    approve = request.form.get("approve", "true").lower() != "false"
    comment = (request.form.get("comment") or "").strip()
    a.status = "approved" if approve else "in_progress"
    a.approved_by_id = current_user.id
    a.approved_at = dt.datetime.utcnow()
    a.approve_comment = comment
    if not approve:
        log_text = "Утверждение отклонено: задача возвращена в работу."
        if comment:
            log_text += " Комментарий руководителя: {}".format(comment)
        db.session.add(ActivityLog(
            activity_id=a.id,
            text=log_text,
            entry_date=dt.datetime.utcnow(),
        ))
    db.session.commit()
    flash("Задача утверждена" if approve else "Задача возвращена в работу")
    return redirect(url_for("activity_view", activity_id=activity_id))

# =========================
# Письма: HTML
# =========================
from sqlalchemy import or_

@app.get("/letters")
@login_required
def letters_page():
    from datetime import datetime, date, timedelta
    memo_source_id = ensure_memo_source()
    filter_type = request.args.get("filter", "all")
    source_id = (request.args.get("source_id") or "").strip()
    executor_id = (request.args.get("executor_id") or "").strip()
    author_id = (request.args.get("author_id") or "").strip()
    search = request.args.get("search", "").strip()
    page = request.args.get("page", 1, type=int) or 1

    # даты фильтра
    today = date.today()
    first_day = date(today.year, today.month, 1)
    last_day = (first_day.replace(month=today.month % 12 + 1, day=1) - timedelta(days=1))

    start_str = request.args.get("date_from", first_day.strftime("%Y-%m-%d"))
    end_str = request.args.get("date_to", last_day.strftime("%Y-%m-%d"))

    try:
        date_from = datetime.strptime(start_str, "%Y-%m-%d").date()
        date_to = datetime.strptime(end_str, "%Y-%m-%d").date()
    except ValueError:
        date_from = first_day
        date_to = last_day

    load_options = [
        joinedload(Letter.author),
        joinedload(Letter.recipients).joinedload(LetterRecipient.user),
        joinedload(Letter.departments),
    ]
    if hasattr(Letter, "executor"):
        load_options.append(joinedload(Letter.executor))
    if hasattr(Letter, "source"):
        load_options.append(joinedload(Letter.source))

    q = (
        Letter.query
        .options(*load_options)
        .outerjoin(Letter.recipients)
        .outerjoin(Letter.departments)
        .filter(
            Letter.archived_at.is_(None),
            or_(Letter.source_id.is_(None), Letter.source_id != memo_source_id),
        )
    )

    # По умолчанию журнал общий. Ограничения применяются только по явному фильтру.
    if filter_type == "my":
        q = q.filter(LetterRecipient.user_id == current_user.id)
    elif filter_type == "dept":
        q = q.filter(Department.id == current_user.department_id)
    elif filter_type == "sent":
        q = q.filter(Letter.author_id == current_user.id)

    # === фильтрация по датам ===
    q = q.filter(
        or_(
            Letter.letter_date.between(date_from, date_to),
            Letter.created_at.between(date_from, date_to)
        )
    )

    if search:
        like_pattern = f"%{search}%"
        letter_filters = [
            Letter.subject.ilike(like_pattern),
            Letter.body.ilike(like_pattern),
        ]
        if hasattr(Letter, "reg_number"):
            letter_filters.append(Letter.reg_number.ilike(like_pattern))
        q = q.filter(or_(*letter_filters))

    if source_id and hasattr(Letter, "source_id"):
        try:
            q = q.filter(Letter.source_id == int(source_id))
        except ValueError:
            source_id = ""

    if executor_id:
        try:
            executor_int = int(executor_id)
            executor_filters = [
                and_(
                    LetterRecipient.user_id == executor_int,
                    LetterRecipient.role == "executor",
                )
            ]
            if hasattr(Letter, "executor_id"):
                executor_filters.insert(0, Letter.executor_id == executor_int)
            q = q.filter(
                or_(
                    *executor_filters,
                )
            )
        except ValueError:
            executor_id = ""

    if author_id:
        try:
            q = q.filter(Letter.author_id == int(author_id))
        except ValueError:
            author_id = ""

    pagination = q.distinct().order_by(
        Letter.letter_date.desc().nullslast(),
        Letter.created_at.desc(),
        Letter.id.desc(),
    ).paginate(page=page, per_page=50, error_out=False)
    letters = pagination.items

    sources = LetterSource.query.order_by(LetterSource.name).all()
    source_lookup = {source.id: source for source in sources}
    executor_ids = sorted({
        le.executor_id for le in letters
        if getattr(le, "executor_id", None)
    })
    executor_lookup = {}
    if executor_ids:
        executor_lookup = {
            user.id: user
            for user in User.query.filter(User.id.in_(executor_ids)).all()
        }
    filter_users = []
    if current_user.role in ("admin", "manager", "superadmin"):
        filter_users = User.query.filter_by(is_approved=True).order_by(User.full_name).all()
    authors = User.query.filter_by(is_approved=True).order_by(User.full_name).all()

    return render_template(
        "letters.html",
        title="Письма",
        letters=letters,
        filter_type=filter_type,
        date_from=start_str,
        date_to=end_str,
        sources=sources,
        source_lookup=source_lookup,
        filter_users=filter_users,
        source_id=source_id,
        executor_id=executor_id,
        author_id=author_id,
        authors=authors,
        search=search,
        pagination=pagination,
        executor_lookup=executor_lookup,
    )

@app.get("/letters/new")
@login_required
def new_letter_page():
    deps = Department.query.order_by(Department.name).all()
    users = User.query.filter_by(is_approved=True).order_by(User.full_name).all()
    sources = LetterSource.query.order_by(LetterSource.name).all()
    import datetime
    initial_reg_number = generate_next_letter_reg_number(getattr(current_user, "department_id", None))

    return render_template(
        "letter_create.html",
        title="Создание письма",
        deps=deps,
        users=users,
        sources=sources,
        letter_prefix=get_app_setting("letter_prefix", ""),
        initial_reg_number=initial_reg_number,
        number_department_id=getattr(current_user, "department_id", None),
        can_manage_reg_numbers=user_can_manage_reg_numbers(),
        now=datetime.datetime.now
    )


@app.get("/api/registration-number")
@login_required
def registration_number_preview():
    document_type = (request.args.get("document_type") or "").strip()
    department_id = (request.args.get("department_id") or "").strip()
    try:
        department_id = int(department_id) if department_id else getattr(current_user, "department_id", None)
    except ValueError:
        return {"error": "Некорректный отдел."}, 400

    if document_type == "letter":
        number = generate_next_letter_reg_number(department_id)
    elif document_type == "memo":
        number = generate_next_memo_reg_number(department_id)
    else:
        return {"error": "Неизвестный тип документа."}, 400
    return {"number": number}


@app.post("/letters")
@login_required
def add_letter():
    subject = request.form.get("subject", "").strip()
    body = request.form.get("body", "").strip()
    source_id = (request.form.get("source_id") or "").strip()
    executor_id = (request.form.get("executor_id") or request.form.get("executors") or "").strip()
    requested_reg_number = (request.form.get("reg_number") or "").strip()

    # --- получаем строку даты из формы ---
    letter_date_str = request.form.get("letter_date", "").strip()

    # --- преобразуем строку в объект datetime.date ---
    letter_date = None
    if letter_date_str:
        try:
            letter_date = datetime.strptime(letter_date_str, "%Y-%m-%d").date()
        except ValueError:
            flash("Некорректный формат даты письма.")
            return redirect(url_for("letters_page"))

    executor = None
    if executor_id:
        try:
            executor = db.session.get(User, int(executor_id))
        except ValueError:
            executor = None
        if not executor:
            flash("Некорректный исполнитель.", "warning")
            return redirect(url_for("new_letter_page"))

    number_department_id = executor.department_id if executor else getattr(current_user, "department_id", None)

    # --- создаём объект письма ---
    le = Letter(
        subject=subject,
        body=body,
        author_id=current_user.id,
        letter_date=letter_date
    )
    le.reg_number = (
        requested_reg_number
        if user_can_manage_reg_numbers() and requested_reg_number
        else generate_next_letter_reg_number(number_department_id)
    )
    if not validate_reg_number_uniqueness(le.reg_number, is_memo=False):
        flash("Письмо с таким номером уже существует.", "warning")
        return redirect(url_for("new_letter_page"))

    if source_id:
        try:
            le.source_id = int(source_id)
        except ValueError:
            flash("Некорректное предприятие.", "warning")
            return redirect(url_for("new_letter_page"))

    if executor:
        le.executor_id = executor.id

    # --- добавляем выбранные отделы ---
    for dep_id in request.form.getlist("departments"):
        dep = Department.query.get(dep_id)
        if dep:
            le.departments.append(dep)

    db.session.add(le)
    db.session.commit()

    # --- добавляем исполнителей и ознакомление ---
    # --- добавляем исполнителей и ознакомление ---
    for key, role in [("executors", "executor"), ("observers", "cc")]:
        ids_str = request.form.get(key, "")
        if ids_str:
            for uid in [x.strip() for x in ids_str.split(",") if x.strip()]:
                db.session.add(LetterRecipient(letter_id=le.id, user_id=int(uid), role=role))

    db.session.commit()
    flash("Письмо создано.")
    return redirect(url_for("letters_page"))

# === Основной маршрут ===
@app.post("/letters/<int:letter_id>/upload")
@login_required
def upload_letter_document(letter_id: int):
    letter = Letter.query.get_or_404(letter_id)
    is_memo = letter.source_id == ensure_memo_source()
    can_edit = letter.author_id == current_user.id or current_user.role in ("admin", "superadmin")
    if is_memo and current_user.role == "manager":
        can_edit = True
    if not can_edit:
        abort(403)
    f = request.files.get("file")

    if not f or not f.filename:
        flash("⚠️ Файл не выбран", "warning")
        return redirect(url_for("letter_view", letter_id=letter_id))

    # Разрешённые типы
    allowed_ext = {"png", "jpg", "jpeg", "pdf", "txt", "docx"}
    ext = f.filename.rsplit(".", 1)[-1].lower()
    if ext not in allowed_ext:
        flash(f"🚫 Недопустимый формат файла: .{ext}", "danger")
        return redirect(url_for("letter_view", letter_id=letter_id))

    # Создаём папку
    upload_dir = os.path.join(current_app.root_path, "static", "uploads")
    os.makedirs(upload_dir, exist_ok=True)

    # Безопасное оригинальное имя
    original_name = secure_filename_rus(f.filename)

    # Генерируем случайное имя
    random_name = f"{uuid.uuid4().hex}_{original_name}"
    save_path = os.path.join(upload_dir, random_name)
    f.save(save_path)

    # Save first, then let the single OCR worker process it in the background.
    doc = LetterDocument(
        letter_id=letter.id,
        filename=original_name,
        filepath=f"uploads/{random_name}",
        ocr_status="queued",
    )

    db.session.add(doc)
    db.session.commit()
    enqueue_document_ocr("letter", doc.id)

    flash("Документ загружен и поставлен в очередь распознавания.", "success")
    return redirect(url_for("letter_view", letter_id=letter.id))
@app.get("/letters/<int:letter_id>")
@login_required
def letter_view(letter_id):
    le = Letter.query.get_or_404(letter_id)
    deps = Department.query.all()
    executor_lookup = {}
    source_lookup = {}
    if getattr(le, "executor_id", None):
        executor_user = User.query.get(le.executor_id)
        if executor_user:
            executor_lookup[executor_user.id] = executor_user
    if getattr(le, "source_id", None):
        source_obj = LetterSource.query.get(le.source_id)
        if source_obj:
            source_lookup[source_obj.id] = source_obj

    is_author = le.author_id == current_user.id
    is_admin = current_user.role in ("admin", "superadmin")

    can_edit = is_author or is_admin

    return render_template(
        "letter_view.html",
        le=le,
        deps=deps,
        executor_lookup=executor_lookup,
        source_lookup=source_lookup,
        can_edit=can_edit,
        has_pending_ocr=any(document.ocr_status in ("queued", "processing") for document in le.documents),
        title=request.form.get("title", "").strip() or f"????? ?? ?????? ?{le.id}"
    )


@app.post("/letters/<int:letter_id>/archive")
@role_required("admin", "superadmin")
def archive_letter(letter_id):
    letter = Letter.query.get_or_404(letter_id)
    if letter.source_id == ensure_memo_source():
        abort(404)
    letter.archived_at = dt.datetime.utcnow()
    letter.archived_by_id = current_user.id if current_user.id != -1 else None
    db.session.commit()
    flash("Письмо перемещено в архив.", "success")
    return redirect(url_for("letters_page"))


@app.post("/letters/<int:letter_id>/restore")
@role_required("admin", "superadmin")
def restore_letter(letter_id):
    letter = Letter.query.get_or_404(letter_id)
    if letter.source_id == ensure_memo_source():
        abort(404)
    letter.archived_at = None
    letter.archived_by_id = None
    db.session.commit()
    flash("Письмо восстановлено из архива.", "success")
    return redirect(url_for("letter_view", letter_id=letter.id))


@app.get("/letters/archive")
@role_required("admin", "superadmin")
def letters_archive_page():
    memo_source_id = ensure_memo_source()
    letters = (
        Letter.query
        .options(
            joinedload(Letter.author),
            joinedload(Letter.source),
            joinedload(Letter.executor),
            joinedload(Letter.recipients).joinedload(LetterRecipient.user),
        )
        .filter(
            Letter.archived_at.isnot(None),
            or_(Letter.source_id.is_(None), Letter.source_id != memo_source_id),
        )
        .order_by(Letter.archived_at.desc(), Letter.id.desc())
        .all()
    )
    return render_template("letters_archive.html", title="Архив писем", letters=letters)


@app.get("/memos/<int:memo_id>")
@login_required
def memo_view(memo_id):
    memo = Letter.query.get_or_404(memo_id)
    if memo.source_id != ensure_memo_source():
        abort(404)

    is_author = memo.author_id == current_user.id
    is_privileged = current_user.role in ("admin", "superadmin", "manager")

    executor_lookup = {}
    if getattr(memo, "executor_id", None):
        executor_user = User.query.get(memo.executor_id)
        if executor_user:
            executor_lookup[executor_user.id] = executor_user

    can_edit = is_author or is_privileged
    linked_activities = Activity.query.filter(
        or_(Activity.source_letter_id == memo.id, Activity.completion_memo_id == memo.id)
    ).order_by(Activity.id.desc()).all()
    return render_template(
        "memo_view.html",
        memo=memo,
        executor_lookup=executor_lookup,
        can_edit=can_edit,
        linked_activities=linked_activities,
        title=f"Служебная записка {memo.reg_number or memo.id}",
    )


@app.post("/memos/<int:memo_id>/archive")
@role_required("admin", "superadmin")
def archive_memo(memo_id):
    memo = Letter.query.get_or_404(memo_id)
    if memo.source_id != ensure_memo_source():
        abort(404)
    memo.archived_at = dt.datetime.utcnow()
    memo.archived_by_id = current_user.id if current_user.id != -1 else None
    db.session.commit()
    flash("Служебная записка перемещена в архив.", "success")
    return redirect(url_for("memos_page"))


@app.post("/memos/<int:memo_id>/approve")
@role_required("manager", "admin", "superadmin")
def approve_memo(memo_id):
    memo = Letter.query.get_or_404(memo_id)
    if memo.source_id != ensure_memo_source():
        abort(404)
    memo.approval_status = "approved"
    memo.approved_by_id = current_user.id if current_user.id != -1 else None
    memo.approved_at = dt.datetime.utcnow()
    db.session.commit()
    flash("Служебная записка утверждена.", "success")
    return redirect(url_for("memo_view", memo_id=memo.id))


@app.post("/memos/<int:memo_id>/restore")
@role_required("admin", "superadmin")
def restore_memo(memo_id):
    memo = Letter.query.get_or_404(memo_id)
    if memo.source_id != ensure_memo_source():
        abort(404)
    memo.archived_at = None
    memo.archived_by_id = None
    db.session.commit()
    flash("Служебная записка восстановлена из архива.", "success")
    return redirect(url_for("memo_view", memo_id=memo.id))


@app.get("/memos/archive")
@role_required("admin", "superadmin")
def memos_archive_page():
    memo_source_id = ensure_memo_source()
    memos = (
        Letter.query
        .options(
            joinedload(Letter.author),
            joinedload(Letter.recipients).joinedload(LetterRecipient.user),
            joinedload(Letter.departments),
        )
        .filter(Letter.source_id == memo_source_id, Letter.archived_at.isnot(None))
        .order_by(Letter.archived_at.desc())
        .all()
    )
    executor_lookup = {
        user.id: user
        for user in User.query.filter_by(is_approved=True).all()
    }
    return render_template(
        "memos_archive.html",
        title="Архив служебных записок",
        memos=memos,
        executor_lookup=executor_lookup,
    )


@app.route("/memos/<int:memo_id>/edit", methods=["GET", "POST"])
@login_required
def edit_memo(memo_id):
    memo = Letter.query.get_or_404(memo_id)
    if memo.source_id != ensure_memo_source():
        abort(404)
    if memo.author_id != current_user.id and current_user.role not in ("admin", "superadmin", "manager"):
        abort(403)
    departments = Department.query.order_by(Department.name).all()
    users = User.query.filter_by(is_approved=True).order_by(User.full_name).all()
    if request.method == "POST":
        subject = (request.form.get("subject") or "").strip()
        reg_number = (request.form.get("reg_number") or "").strip()
        department_id = (request.form.get("department_id") or "").strip()
        executor_id = (request.form.get("executor_id") or "").strip()
        memo_date_raw = (request.form.get("memo_date") or "").strip()
        if not subject or not reg_number:
            flash("Заполните номер и краткое содержание.", "warning")
            return redirect(request.url)
        try:
            department = db.session.get(Department, int(department_id))
        except ValueError:
            department = None
        if not department:
            flash("Выберите отдел.", "warning")
            return redirect(request.url)
        try:
            memo_date = datetime.strptime(memo_date_raw, "%Y-%m-%d").date() if memo_date_raw else None
        except ValueError:
            flash("Некорректный формат даты.", "warning")
            return redirect(request.url)
        executor = None
        if executor_id:
            try:
                executor = db.session.get(User, int(executor_id))
            except ValueError:
                executor = None
            if not executor:
                flash("Исполнитель не найден.", "warning")
                return redirect(request.url)
        if not validate_reg_number_uniqueness(reg_number, is_memo=True, exclude_letter_id=memo.id):
            flash("Служебная записка с таким номером уже существует.", "warning")
            return redirect(request.url)
        memo.reg_number = reg_number
        memo.subject = subject
        memo.body = (request.form.get("body") or "").strip()
        memo.letter_date = memo_date
        memo.executor_id = executor.id if executor else None
        memo.departments[:] = [department]
        LetterRecipient.query.filter_by(letter_id=memo.id, role="executor").delete(synchronize_session=False)
        if executor:
            db.session.add(LetterRecipient(letter_id=memo.id, user_id=executor.id, role="executor"))
        db.session.commit()
        flash("Служебная записка обновлена.", "success")
        return redirect(url_for("memo_view", memo_id=memo.id))
    return render_template(
        "memo_edit.html",
        title=f"Редактирование служебной записки {memo.reg_number or memo.id}",
        memo=memo,
        departments=departments,
        users=users,
    )

# @app.get("/letters/<int:letter_id>")
# @login_required
# def letter_view(letter_id):
#     le = Letter.query.get_or_404(letter_id)
#     deps = Department.query.all()
#
#     recipient_ids = [r.user_id for r in le.recipients]
#     is_author = le.author_id == current_user.id
#     is_recipient = current_user.id in recipient_ids
#     is_admin = current_user.role == "admin"
#
#     # 🔒 Проверяем доступ
#     has_access = is_author or is_recipient or is_admin
#
#     if not has_access:
#         # Не автор и не получатель — показываем страницу с просьбой дать доступ
#         author = User.query.get(le.author_id)
#         return render_template(
#             "letter_request_access.html",
#             title="Доступ к письму запрещён",
#             le=le,
#             author=author,
#         ), 403
#
#     can_edit = is_author or is_admin
#
#     return render_template(
#         "letter_view.html",
#         le=le,
#         deps=deps,
#         can_edit=can_edit,
#         title=f"Письмо #{le.id}"
#     )
#

@app.post("/letters/<int:letter_id>/read")
@login_required
def mark_letter_read(letter_id: int):
    le = Letter.query.get_or_404(letter_id)
    updated = 0
    for r in le.recipients:
        if (r.user_id == current_user.id) or (r.department_id and r.department_id == current_user.department_id):
            if not r.is_read:
                r.is_read = True
                r.read_at = dt.datetime.utcnow()
                updated += 1
    if updated:
        db.session.commit()
        flash("Письмо отмечено прочитанным")
    return redirect(url_for("letter_view", letter_id=letter_id))

@app.post("/letters/<int:letter_id>/sign")
@login_required
def sign_letter(letter_id: int):
    le = Letter.query.get_or_404(letter_id)
    signed_result = request.form.get("signed_result", "")
    le.signed_result = signed_result
    db.session.commit()
    flash("Подпись-результат сохранена")
    return redirect(url_for("letter_view", letter_id=letter_id))

# Связи «из письма → активность»
@app.route("/letters/<int:letter_id>/create-activity", methods=["GET", "POST"])
@login_required
def create_activity_from_letter(letter_id):
    le = Letter.query.get_or_404(letter_id)

    if request.method == "POST":
        from datetime import datetime

        type_id = request.form.get("type_id")
        start_date = request.form.get("start_date")
        end_date = request.form.get("end_date")
        priority = parse_activity_priority(request.form.get("priority", "3"))
        selected_type = db.session.get(ActivityType, int(type_id)) if type_id and type_id.isdigit() else None
        if not selected_type:
            flash("Выберите тип задачи.", "warning")
            return redirect(request.url)
        if priority is None:
            flash("Выберите приоритет от 1 до 5.", "warning")
            return redirect(request.url)

        memo_source = LetterSource.query.filter_by(name="Служебная записка").first()
        source_label = "служебную записку" if memo_source and le.source_id == memo_source.id else "письмо"
        is_memo = source_label == "служебную записку"


        # --- создаём активность "Ответ на письмо №..." ---
        a = Activity(
            title=request.form.get("title", "").strip() or "Ответ на {} №{}".format(source_label, le.reg_number or le.id),
            description=request.form.get("description", "").strip() or (le.body[:300] + "..."),
            owner_id=current_user.id,
            created_by_id=current_user.id if current_user.id != -1 else None,
            type_id=selected_type.id,
            priority=priority,
            source_letter_id=le.id,
            status="in_progress",
            start_date=datetime.now(),
            direction=request.form.get("direction", "").strip(),
            outgoing_document=request.form.get("outgoing_document", "").strip()
        )

        if start_date:
            a.start_date = datetime.strptime(start_date, "%Y-%m-%d")
        if end_date:
            a.end_date = datetime.strptime(end_date, "%Y-%m-%d")

        db.session.add(a)
        db.session.commit()

        flash("Создана активность: ответ на {} №{}.".format(source_label, le.reg_number or le.id))
        return redirect(url_for("memo_view", memo_id=le.id) if is_memo else url_for("letter_view", letter_id=le.id))

    types = ActivityType.query.order_by(ActivityType.name).all()
    existing_dirs = [d[0] for d in db.session.query(Activity.direction).filter(Activity.direction.isnot(None)).distinct().all()]
    return render_template("activity_create_from_letter.html", le=le, types=types, existing_directions=existing_dirs)

# Связи «из активности → письмо»
@app.get("/memos")
@login_required
def memos_page():
    from datetime import datetime, date, timedelta

    memo_source_id = ensure_memo_source()
    today = date.today()
    first_day = date(today.year, today.month, 1)
    last_day = (first_day.replace(month=today.month % 12 + 1, day=1) - timedelta(days=1))

    start_str = request.args.get("date_from", first_day.strftime("%Y-%m-%d"))
    end_str = request.args.get("date_to", last_day.strftime("%Y-%m-%d"))
    department_id = (request.args.get("department_id") or "").strip()
    executor_id = (request.args.get("executor_id") or "").strip()
    author_id = (request.args.get("author_id") or "").strip()
    search = request.args.get("search", "").strip()
    page = request.args.get("page", 1, type=int) or 1

    try:
        date_from = datetime.strptime(start_str, "%Y-%m-%d").date()
        date_to = datetime.strptime(end_str, "%Y-%m-%d").date()
    except ValueError:
        date_from = first_day
        date_to = last_day

    q = (
        Letter.query
        .options(
            joinedload(Letter.author),
            joinedload(Letter.recipients).joinedload(LetterRecipient.user),
            joinedload(Letter.departments),
        )
        .outerjoin(Letter.recipients)
        .outerjoin(Letter.departments)
        .filter(Letter.source_id == memo_source_id, Letter.archived_at.is_(None))
    )

    q = q.filter(
        or_(
            Letter.letter_date.between(date_from, date_to),
            Letter.created_at.between(date_from, date_to),
        )
    )

    if department_id:
        try:
            q = q.filter(Department.id == int(department_id))
        except ValueError:
            department_id = ""

    if executor_id:
        try:
            executor_int = int(executor_id)
            q = q.filter(
                or_(
                    Letter.executor_id == executor_int,
                    and_(LetterRecipient.user_id == executor_int, LetterRecipient.role == "executor"),
                )
            )
        except ValueError:
            executor_id = ""

    if author_id:
        try:
            q = q.filter(Letter.author_id == int(author_id))
        except ValueError:
            author_id = ""

    if search:
        like_pattern = f"%{search}%"
        q = q.filter(
            or_(
                Letter.reg_number.ilike(like_pattern),
                Letter.subject.ilike(like_pattern),
                Letter.body.ilike(like_pattern),
            )
        )

    pagination = q.distinct().order_by(
        Letter.letter_date.desc().nullslast(),
        Letter.created_at.desc(),
        Letter.id.desc(),
    ).paginate(page=page, per_page=50, error_out=False)
    memos = pagination.items

    departments = Department.query.order_by(Department.name).all()
    department_lookup = {dep.id: dep for dep in departments}
    executor_ids = sorted({memo.executor_id for memo in memos if getattr(memo, "executor_id", None)})
    executor_lookup = {}
    if executor_ids:
        executor_lookup = {user.id: user for user in User.query.filter(User.id.in_(executor_ids)).all()}
    filter_users = []
    privileged_roles = ("admin", "manager", "superadmin")
    if current_user.role in privileged_roles:
        filter_users = User.query.filter_by(is_approved=True).order_by(User.full_name).all()
    authors = User.query.filter_by(is_approved=True).order_by(User.full_name).all()

    return render_template(
        "memos.html",
        title="Служебные записки",
        memos=memos,
        date_from=start_str,
        date_to=end_str,
        departments=departments,
        department_lookup=department_lookup,
        department_id=department_id,
        executor_id=executor_id,
        author_id=author_id,
        authors=authors,
        executor_lookup=executor_lookup,
        filter_users=filter_users,
        search=search,
        pagination=pagination,
    )


@app.get("/memos/analytics")
@role_required("manager", "admin", "superadmin")
def memos_analytics():
    today = dt.date.today()
    default_from = today.replace(day=1)
    default_to = (default_from.replace(day=28) + dt.timedelta(days=4)).replace(day=1) - dt.timedelta(days=1)
    date_from_value = (request.args.get("date_from") or default_from.isoformat()).strip()
    date_to_value = (request.args.get("date_to") or default_to.isoformat()).strip()
    department_id = (request.args.get("department_id") or "").strip()
    executor_id = (request.args.get("executor_id") or "").strip()
    try:
        date_from = dt.datetime.strptime(date_from_value, "%Y-%m-%d").date()
        date_to = dt.datetime.strptime(date_to_value, "%Y-%m-%d").date()
    except ValueError:
        date_from, date_to = default_from, default_to
        date_from_value, date_to_value = default_from.isoformat(), default_to.isoformat()

    query = (
        Letter.query.options(joinedload(Letter.author).joinedload(User.department), joinedload(Letter.executor), joinedload(Letter.departments))
        .outerjoin(Letter.departments)
        .filter(
            Letter.source_id == ensure_memo_source(), Letter.archived_at.is_(None),
            or_(Letter.letter_date.between(date_from, date_to), Letter.created_at.between(date_from, date_to)),
        )
    )
    is_manager = current_user.role == "manager"
    if is_manager and current_user.department_id:
        query, department_id = query.filter(Department.id == current_user.department_id), str(current_user.department_id)
    elif department_id.isdigit():
        query = query.filter(Department.id == int(department_id))
    if executor_id.isdigit():
        query = query.filter(Letter.executor_id == int(executor_id))
    memos = query.distinct().all()
    summary = {"total": len(memos), "approved": 0, "pending": 0}
    department_rows, executor_rows = {}, {}
    for memo in memos:
        status = "approved" if memo.approval_status == "approved" else "pending"
        summary[status] += 1
        department = memo.departments[0] if memo.departments else (memo.author.department if memo.author else None)
        department_name = department.name if department else "Без отдела"
        executor_name = memo.executor.full_name if memo.executor else "Не назначен"
        for bucket, name in ((department_rows, department_name), (executor_rows, executor_name)):
            row = bucket.setdefault(name, {"name": name, "total": 0, "approved": 0, "pending": 0})
            row["total"] += 1
            row[status] += 1
    for row in list(department_rows.values()) + list(executor_rows.values()):
        row["approval_percent"] = round(100 * row["approved"] / row["total"]) if row["total"] else 0
    department_rows = sorted(department_rows.values(), key=lambda row: (-row["total"], row["name"]))
    executor_rows = sorted(executor_rows.values(), key=lambda row: (-row["total"], row["name"]))
    summary["approval_percent"] = round(100 * summary["approved"] / summary["total"]) if summary["total"] else 0
    chart_data = {
        "status": [summary["approved"], summary["pending"]],
        "departments": [row["name"] for row in department_rows],
        "department_approved": [row["approved"] for row in department_rows],
        "department_pending": [row["pending"] for row in department_rows],
    }
    users = User.query.filter_by(is_approved=True).order_by(User.full_name)
    if is_manager and current_user.department_id:
        users = users.filter(User.department_id == current_user.department_id)
    return render_template(
        "memos_analytics.html", title="Аналитика служебных записок", summary=summary,
        department_rows=department_rows, executor_rows=executor_rows, chart_data=chart_data,
        departments=Department.query.order_by(Department.name).all(), users=users.all(),
        date_from=date_from_value, date_to=date_to_value, department_id=department_id,
        executor_id=executor_id, is_manager=is_manager,
    )


@app.get("/memos/new")
@login_required
def new_memo_page():
    departments = Department.query.order_by(Department.name).all()
    users = User.query.filter_by(is_approved=True).order_by(User.full_name).all()
    import datetime
    templates = MemoTemplate.query.options(joinedload(MemoTemplate.department), joinedload(MemoTemplate.executor)).order_by(MemoTemplate.name).all()
    template_id = (request.args.get("template_id") or "").strip()
    selected_template = db.session.get(MemoTemplate, int(template_id)) if template_id.isdigit() else None
    number_department_id = (
        selected_template.department_id if selected_template and selected_template.department_id
        else getattr(current_user, "department_id", None)
    )
    initial_reg_number = generate_next_memo_reg_number(number_department_id)

    return render_template(
        "memo_create.html",
        title="Создание служебной записки",
        departments=departments,
        users=users,
        letter_prefix=get_app_setting("letter_prefix", ""),
        memo_prefix=get_app_setting("memo_prefix", "421"),
        initial_reg_number=initial_reg_number,
        number_department_id=number_department_id,
        can_manage_reg_numbers=user_can_manage_reg_numbers(),
        now=datetime.datetime.now,
        templates=templates,
        selected_template=selected_template,
    )


@app.post("/memos")
@login_required
def add_memo():
    subject = request.form.get("subject", "").strip()
    body = request.form.get("body", "").strip()
    department_id = (request.form.get("department_id") or "").strip()
    executor_id = (request.form.get("executor_id") or "").strip()
    memo_date_str = request.form.get("memo_date", "").strip()
    requested_reg_number = (request.form.get("reg_number") or "").strip()

    memo_date = None
    if memo_date_str:
        try:
            memo_date = datetime.strptime(memo_date_str, "%Y-%m-%d").date()
        except ValueError:
            flash("Некорректный формат даты служебной записки.", "warning")
            return redirect(url_for("new_memo_page"))

    department = None
    if department_id:
        try:
            department = db.session.get(Department, int(department_id))
        except ValueError:
            department = None
    if not department:
        flash("Выберите отдел служебной записки.", "warning")
        return redirect(url_for("new_memo_page"))

    memo = Letter(
        reg_number=(
            requested_reg_number
            if user_can_manage_reg_numbers() and requested_reg_number
            else generate_next_memo_reg_number(department.id)
        ),
        subject=subject,
        body=body,
        author_id=current_user.id,
        source_id=ensure_memo_source(),
        letter_date=memo_date,
    )
    if not validate_reg_number_uniqueness(memo.reg_number, is_memo=True):
        flash("Служебная записка с таким номером уже существует.", "warning")
        return redirect(url_for("new_memo_page"))

    if executor_id:
        try:
            memo.executor_id = int(executor_id)
        except ValueError:
            flash("Некорректный исполнитель.", "warning")
            return redirect(url_for("new_memo_page"))

    memo.departments.append(department)

    db.session.add(memo)
    db.session.commit()

    if memo.executor_id:
        db.session.add(LetterRecipient(letter_id=memo.id, user_id=memo.executor_id, role="executor"))
        db.session.commit()

    flash(f"Служебная записка {memo.reg_number} создана.", "success")
    return redirect(url_for("memos_page"))


@app.get("/admin/settings")
@role_required("admin", "superadmin")
def admin_settings():
    complexity_levels = [
        {
            "level": level,
            "label": meta["label"],
            "hours": get_app_setting_int(
                "activity_complexity_hours_{}".format(level), meta["default_hours"]
            ),
        }
        for level, meta in ACTIVITY_COMPLEXITY_META.items()
    ]
    return render_template(
        "admin_settings.html",
        title="Настройки",
        letter_prefix=get_app_setting("letter_prefix", ""),
        memo_prefix=get_app_setting("memo_prefix", "421"),
        letter_start_number=get_app_setting_int("letter_start_number", 1),
        memo_start_number=get_app_setting_int("memo_start_number", 1),
        integration_api_token=get_app_setting("integration_api_token", ""),
        departments=Department.query.order_by(Department.name).all(),
        complexity_levels=complexity_levels,
    )


@app.post("/admin/settings")
@role_required("admin", "superadmin")
def save_admin_settings():
    letter_prefix = (request.form.get("letter_prefix") or "").strip()
    memo_prefix = (request.form.get("memo_prefix") or "").strip()
    letter_start_number = (request.form.get("letter_start_number") or "1").strip()
    memo_start_number = (request.form.get("memo_start_number") or "1").strip()
    if not memo_prefix:
        flash("Префикс номера служебной записки не может быть пустым.", "warning")
        return redirect(url_for("admin_settings"))
    try:
        letter_start_number_int = int(letter_start_number)
        memo_start_number_int = int(memo_start_number)
    except ValueError:
        flash("Стартовые номера должны быть целыми числами.", "warning")
        return redirect(url_for("admin_settings"))
    if letter_start_number_int < 1 or memo_start_number_int < 1:
        flash("Стартовые номера должны быть больше нуля.", "warning")
        return redirect(url_for("admin_settings"))

    complexity_hours = {}
    for level, meta in ACTIVITY_COMPLEXITY_META.items():
        raw_hours = (request.form.get("activity_complexity_hours_{}".format(level)) or "").strip()
        try:
            hours = int(raw_hours)
        except ValueError:
            flash("Часы по уровню сложности должны быть целыми числами.", "warning")
            return redirect(url_for("admin_settings"))
        if hours < 1 or hours > 24:
            flash("Норма часов по уровню сложности должна быть от 1 до 24.", "warning")
            return redirect(url_for("admin_settings"))
        complexity_hours[level] = hours

    departments = Department.query.order_by(Department.name).all()
    for department in departments:
        department_letter_prefix = (request.form.get(f"department_letter_prefix_{department.id}") or "").strip()
        department_memo_prefix = (request.form.get(f"department_memo_prefix_{department.id}") or "").strip()
        department_letter_start = (request.form.get(f"department_letter_start_number_{department.id}") or "").strip()
        department_memo_start = (request.form.get(f"department_memo_start_number_{department.id}") or "").strip()
        department_letter_series = (request.form.get(f"department_letter_number_series_{department.id}") or "").strip()
        department_memo_series = (request.form.get(f"department_memo_number_series_{department.id}") or "").strip()
        try:
            department_letter_start_int = int(department_letter_start) if department_letter_start else None
            department_memo_start_int = int(department_memo_start) if department_memo_start else None
        except ValueError:
            flash(f"Стартовые номера отдела «{department.name}» должны быть целыми числами.", "warning")
            return redirect(url_for("admin_settings"))
        if (
            (department_letter_start_int is not None and department_letter_start_int < 1)
            or (department_memo_start_int is not None and department_memo_start_int < 1)
        ):
            flash(f"Стартовые номера отдела «{department.name}» должны быть больше нуля.", "warning")
            return redirect(url_for("admin_settings"))
        department.letter_prefix = department_letter_prefix or None
        department.memo_prefix = department_memo_prefix or None
        department.letter_start_number = department_letter_start_int
        department.memo_start_number = department_memo_start_int
        department.letter_number_series = department_letter_series or None
        department.memo_number_series = department_memo_series or None

    set_app_setting("letter_prefix", letter_prefix)
    set_app_setting("memo_prefix", memo_prefix)
    set_app_setting("letter_start_number", str(letter_start_number_int))
    set_app_setting("memo_start_number", str(memo_start_number_int))
    for level, hours in complexity_hours.items():
        set_app_setting("activity_complexity_hours_{}".format(level), str(hours))
    db.session.commit()
    flash("Настройки сохранены.", "success")
    return redirect(url_for("admin_settings"))


@app.post("/admin/settings/integration-token")
@role_required("admin", "superadmin")
def regenerate_integration_api_token():
    set_app_setting("integration_api_token", secrets.token_urlsafe(32))
    flash("Токен интеграции CRM создан. Передайте его только администратору внешней CRM.", "success")
    return redirect(url_for("admin_settings"))


@app.get("/admin/backups")
@role_required("admin", "superadmin")
def admin_backups():
    backups = []
    for path in list_backups():
        stat = path.stat()
        backups.append({
            "name": path.name,
            "size_kb": round(stat.st_size / 1024, 1),
            "created_at": dt.datetime.fromtimestamp(stat.st_mtime),
        })
    return render_template("admin_backups.html", title="Резервные копии БД", backups=backups)


@app.post("/admin/backups/create")
@role_required("admin", "superadmin")
def create_database_backup():
    try:
        archive = create_backup()
    except (OSError, sqlite3.Error) as exc:
        flash(f"Не удалось создать резервную копию: {exc}", "danger")
    else:
        flash(f"Архив базы данных создан: {archive.name}", "success")
    return redirect(url_for("admin_backups"))


@app.get("/admin/backups/<filename>")
@role_required("admin", "superadmin")
def download_database_backup(filename):
    if Path(filename).name != filename or not filename.endswith(".zip"):
        abort(404)
    archive = BACKUP_DIR / filename
    if not archive.is_file():
        abort(404)
    return send_from_directory(BACKUP_DIR, filename, as_attachment=True, download_name=filename)


@app.get("/admin/letter-sources")
@login_required
def admin_letter_sources():
    sources = LetterSource.query.order_by(LetterSource.name).all()
    return render_template(
        "admin_letter_sources.html",
        title="Предприятия",
        sources=sources,
        can_manage_sources=current_user.role in ("admin", "superadmin"),
    )

@app.post("/admin/letter-sources")
@login_required
def add_letter_source():
    name = request.form.get("name", "").strip()
    bin_code = request.form.get("bin_code", "").strip()
    address = request.form.get("address", "").strip()
    phone = request.form.get("phone", "").strip()
    email = request.form.get("email", "").strip()
    contact_person = request.form.get("contact_person", "").strip()
    notes = request.form.get("notes", "").strip()
    if not name:
        flash("Название предприятия не может быть пустым.")
        return redirect(url_for("admin_letter_sources"))
    if LetterSource.query.filter(func.lower(LetterSource.name) == name.lower()).first():
        flash("Такое предприятие уже существует.")
        return redirect(url_for("admin_letter_sources"))

    db.session.add(
        LetterSource(
            name=name,
            bin_code=bin_code or None,
            address=address or None,
            phone=phone or None,
            email=email or None,
            contact_person=contact_person or None,
            notes=notes or None,
        )
    )
    db.session.commit()
    flash(f"Предприятие «{name}» добавлено.")
    return redirect(url_for("admin_letter_sources"))


@app.post("/letter-sources/quick-add")
@login_required
def quick_add_letter_source():
    name = (request.form.get("name") or "").strip()
    if not name:
        return jsonify(success=False, error="Введите название предприятия."), 400

    existing = LetterSource.query.filter(func.lower(LetterSource.name) == name.lower()).first()
    if existing:
        return jsonify(success=True, id=existing.id, name=existing.name, existed=True)

    source = LetterSource(name=name)
    db.session.add(source)
    db.session.commit()
    return jsonify(success=True, id=source.id, name=source.name, existed=False)


@app.post("/letters/<int:letter_id>/update-reg-number")
@role_required("admin", "superadmin")
def update_letter_reg_number(letter_id):
    letter = Letter.query.get_or_404(letter_id)
    reg_number = (request.form.get("reg_number") or "").strip()
    if not reg_number:
        flash("Номер не может быть пустым.", "warning")
        return redirect(request.referrer or url_for("letters_page"))

    is_memo = letter.source_id == ensure_memo_source()
    if not validate_reg_number_uniqueness(reg_number, is_memo=is_memo, exclude_letter_id=letter.id):
        flash("Документ с таким номером уже существует.", "warning")
        return redirect(request.referrer or url_for("letters_page"))

    letter.reg_number = reg_number
    db.session.commit()
    flash("Номер документа обновлён.", "success")

    if is_memo:
        return redirect(url_for("memo_view", memo_id=letter.id))
    return redirect(url_for("letter_view", letter_id=letter.id))

@app.post("/admin/letter-sources/<int:source_id>/edit")
@role_required("admin", "superadmin")
def edit_letter_source(source_id):
    source = LetterSource.query.get_or_404(source_id)
    name = request.form.get("name", "").strip()
    bin_code = request.form.get("bin_code", "").strip()
    address = request.form.get("address", "").strip()
    phone = request.form.get("phone", "").strip()
    email = request.form.get("email", "").strip()
    contact_person = request.form.get("contact_person", "").strip()
    notes = request.form.get("notes", "").strip()
    if not name:
        flash("Название предприятия не может быть пустым.")
        return redirect(url_for("admin_letter_sources"))

    duplicate = LetterSource.query.filter(
        func.lower(LetterSource.name) == name.lower(),
        LetterSource.id != source.id,
    ).first()
    if duplicate:
        flash("Такое предприятие уже существует.")
        return redirect(url_for("admin_letter_sources"))

    source.name = name
    source.bin_code = bin_code or None
    source.address = address or None
    source.phone = phone or None
    source.email = email or None
    source.contact_person = contact_person or None
    source.notes = notes or None
    db.session.commit()
    flash("Изменения сохранены.")
    return redirect(url_for("admin_letter_sources"))

@app.post("/admin/letter-sources/<int:source_id>/delete")
@role_required("admin", "superadmin")
def delete_letter_source(source_id):
    source = LetterSource.query.get_or_404(source_id)
    Letter.query.filter_by(source_id=source.id).update({"source_id": None})
    db.session.delete(source)
    db.session.commit()
    flash(f"Предприятие «{source.name}» удалено.")
    return redirect(url_for("admin_letter_sources"))

@app.get("/activities/<int:activity_id>/create-letter")
@login_required
def create_letter_from_activity(activity_id: int):
    a = Activity.query.get_or_404(activity_id)

    # создаём письмо
    subject = request.form.get("subject", f"Результаты активности №{a.id}: {a.title}")
    body = request.form.get("body", a.description or "")

    # дата письма берётся из формы или текущая
    letter_date_str = request.form.get("letter_date", "")
    if letter_date_str:
        try:
            letter_date = datetime.strptime(letter_date_str, "%Y-%m-%d").date()
        except ValueError:
            letter_date = datetime.now().date()
    else:
        letter_date = datetime.now().date()

    le = Letter(
        author_id=current_user.id,
        subject=subject,
        body=body,
        letter_date=letter_date
    )

    # ✅ сохраняем сам объект письма
    db.session.add(le)
    db.session.flush()
    print("📨 DEBUG: добавляем письмо к активности" )
    # ✅ устанавливаем связь письмо ↔ активность
    if hasattr(le, "activities"):
        le.activities.append(a)
    elif hasattr(a, "letters"):
        a.letters.append(le)
    print("📨 DEBUG: добавляем письмо к активности", a.id)
    print("📨 DEBUG: до добавления:", [x.id for x in a.letters])
    le.activities.append(a)
    print("📨 DEBUG: после добавления:", [x.id for x in le.activities])
    db.session.commit()
    print("📨 DEBUG: commit успешен")

    db.session.commit()

    flash(f"Создано письмо №{le.id} по активности №{a.id}")
    return redirect(url_for("activity_view", activity_id=a.id))


# =========================
# Тех. хелпер
# =========================
@app.errorhandler(403)
def _403(e): return render_template("error.html", title="Доступ запрещён", code=403), 403

@app.errorhandler(404)
def _404(e): return render_template("error.html", title="Не найдено", code=404), 404

# =========================
# Запуск
# =========================
# --- Регистрация ---
@app.get("/register")
def register_page():
    deps = Department.query.order_by(Department.name).all()
    return render_template("register.html", title="Регистрация", deps=deps)

@app.post("/register")
def register_form():
    data = request.form
    full_name = data.get("full_name")
    username = data.get("username")
    password = data.get("password")
    dep_id = data.get("department_id")

    if User.query.filter_by(username=username).first():
        flash("Пользователь с таким логином уже существует!")
        return redirect(url_for("register_page"))

    u = User(full_name=full_name, username=username,
             department_id=dep_id, role="worker", is_approved=False)
    u.set_password(password)
    db.session.add(u)
    db.session.commit()
    flash("Регистрация отправлена на подтверждение администратору.")
    return redirect(url_for("login_page"))

# --- Вход ---
from flask import request, redirect, url_for, flash
from flask_login import login_user, UserMixin
from config import SUPERADMIN_LOGIN, SUPERADMIN_PASSWORD
from sqlalchemy import text, func

@app.post("/login")
def login_form():
    username = (request.form.get("username") or "").strip()
    password = request.form.get("password") or ""
    selected_workshop = (request.form.get("workshop_id") or "").strip()

    print("=== DEBUG LOGIN ===")
    print("INPUT username:", repr(username))
    print("INPUT password:", repr(password))

    if username == SUPERADMIN_LOGIN and password == SUPERADMIN_PASSWORD:
        if selected_workshop and selected_workshop != "superadmin":
            flash("Для входа суперадминистратором выберите соответствующий пункт.", "warning")
            return redirect(url_for("login_page"))
        super_user =  SuperAdmin()
        login_user(super_user, remember=True)
        flash("🔐 Вход выполнен как суперадминистратор", "success")
        return redirect(url_for("admin_users"))
    # --- обычная проверка пользователей ---
    try:
        user = User.query.filter(func.trim(User.username) == username).first()
        print("filtered(trim) user:", user)
    except Exception as e:
        print("DB query error:", e)
        user = None

    if not user:
        flash("Пользователь не найден.", "warning")
        print("⚠️ user not found in DB")
        return redirect(url_for("login_page"))

    if not getattr(user, "is_approved", True):
        flash("Ваш аккаунт ожидает подтверждения администратора.", "info")
        return redirect(url_for("login_page"))

    if not user.is_enabled:
        flash("Учётная запись отключена администратором.", "warning")
        return redirect(url_for("login_page"))

    # The browser filters the list too, but validating here prevents selecting
    # a user from another workshop through a manually crafted request.
    if selected_workshop:
        actual_workshop = (
            str(user.department.workshop_id)
            if user.department and user.department.workshop_id
            else "unassigned"
        )
        if selected_workshop != actual_workshop:
            flash("Выбранный сотрудник не относится к указанному цеху.", "warning")
            return redirect(url_for("login_page"))

    # --- Проверка пароля (шифрованного) ---
    try:
        ok = user.check_password(password)
    except Exception as e:
        print("check_password exception:", e)
        ok = False
    print("check_password returned:", ok)

    if ok:
        login_user(user)
        flash("Вход успешен", "success")
        return redirect(url_for("dashboard"))
    else:
        flash("Неверный пароль.", "danger")
        return redirect(url_for("login_page"))

# --- Админка: список пользователей ---
def can_manage_user(user):
    """Администратор работает только с пользователями своего цеха."""
    if current_user.role == "superadmin":
        return True
    workshop_id = current_workshop_id()
    return (
        current_user.role == "admin"
        and workshop_id is not None
        and user.department is not None
        and user.department.workshop_id == workshop_id
    )


def departments_available_to_current_admin():
    query = Department.query.options(joinedload(Department.workshop)).order_by(Department.name)
    if current_user.role == "superadmin":
        return query
    workshop_id = current_workshop_id()
    return query.filter(Department.workshop_id == workshop_id) if workshop_id else query.filter(False)


@app.get("/admin/users")
@role_required("admin","superadmin")
def admin_users():
    users_query = User.query.options(
        joinedload(User.department).joinedload(Department.workshop)
    ).order_by(User.created_at.desc())
    if current_user.role != "superadmin":
        workshop_id = current_workshop_id()
        users_query = (
            users_query.join(User.department).filter(Department.workshop_id == workshop_id)
            if workshop_id else users_query.filter(User.id == current_user.id)
        )
    users = users_query.all()
    deps = departments_available_to_current_admin().all()
    return render_template("admin_users.html", title="Пользователи", users=users, deps=deps)

@app.post("/admin/users")
@role_required("admin", "superadmin")
def admin_add_user():
    full_name = (request.form.get("full_name") or "").strip()
    username = (request.form.get("username") or "").strip()
    password = request.form.get("password") or ""
    dep_id = (request.form.get("department_id") or "").strip()
    role = (request.form.get("role") or "worker").strip()

    if not full_name or not username or not password:
        flash("Заполните ФИО, логин и пароль.", "warning")
        return redirect(url_for("admin_users"))

    if User.query.filter(func.trim(User.username) == username).first():
        flash("Пользователь с таким логином уже существует.", "warning")
        return redirect(url_for("admin_users"))

    allowed_roles = {"worker", "deputy", "manager"}
    if current_user.role == "superadmin":
        allowed_roles.add("admin")
    if role not in allowed_roles:
        role = "worker"

    department_id = None
    if dep_id:
        try:
            department_id = int(dep_id)
        except ValueError:
            flash("Некорректный отдел.", "warning")
            return redirect(url_for("admin_users"))

    department = db.session.get(Department, department_id) if department_id else None
    if department_id and not department:
        flash("Выбранный отдел не найден.", "warning")
        return redirect(url_for("admin_users"))
    if department and not can_manage_department(department):
        abort(403)
    if current_user.role == "admin" and department is None:
        flash("Для пользователя необходимо выбрать отдел своего цеха.", "warning")
        return redirect(url_for("admin_users"))

    u = User(
        full_name=full_name,
        username=username,
        department_id=department_id,
        role=role,
        is_approved=True,
    )
    u.set_password(password)
    db.session.add(u)
    db.session.commit()
    flash(f"Пользователь {u.full_name} добавлен.", "success")
    return redirect(url_for("admin_users"))


@app.route("/admin/users/<int:user_id>/edit", methods=["GET", "POST"])
@role_required("admin", "superadmin")
def edit_user(user_id):
    """Изменение учетной записи без необходимости пересоздавать пользователя."""
    user = User.query.get_or_404(user_id)
    if not can_manage_user(user):
        abort(403)
    if user.role in ("admin", "superadmin") and current_user.role != "superadmin":
        abort(403)

    departments = departments_available_to_current_admin().all()
    allowed_roles = {"worker", "deputy", "manager"}
    if current_user.role == "superadmin":
        allowed_roles.add("admin")

    if request.method == "POST":
        full_name = (request.form.get("full_name") or "").strip()
        username = (request.form.get("username") or "").strip()
        password = request.form.get("password") or ""
        role = (request.form.get("role") or "worker").strip()
        department_raw = (request.form.get("department_id") or "").strip()

        if not full_name or not username:
            flash("Заполните ФИО и логин.", "warning")
            return redirect(url_for("edit_user", user_id=user.id))
        if role not in allowed_roles:
            flash("Недопустимая роль.", "warning")
            return redirect(url_for("edit_user", user_id=user.id))
        duplicate = User.query.filter(
            func.trim(User.username) == username,
            User.id != user.id,
        ).first()
        if duplicate:
            flash("Пользователь с таким логином уже существует.", "warning")
            return redirect(url_for("edit_user", user_id=user.id))

        department_id = None
        if department_raw:
            try:
                department_id = int(department_raw)
            except ValueError:
                flash("Некорректный отдел.", "warning")
                return redirect(url_for("edit_user", user_id=user.id))
            department = db.session.get(Department, department_id)
            if not department:
                flash("Выбранный отдел не найден.", "warning")
                return redirect(url_for("edit_user", user_id=user.id))
            if not can_manage_department(department):
                abort(403)

        if current_user.role == "admin" and department_id is None:
            flash("Для пользователя необходимо выбрать отдел своего цеха.", "warning")
            return redirect(url_for("edit_user", user_id=user.id))

        user.full_name = full_name
        user.username = username
        user.department_id = department_id
        user.role = role
        user.is_approved = request.form.get("is_approved") == "1"
        if password:
            user.set_password(password)
        db.session.commit()
        flash(f"Данные пользователя {user.full_name} сохранены.", "success")
        return redirect(url_for("admin_users"))

    return render_template(
        "admin_user_edit.html",
        title="Редактирование пользователя",
        user=user,
        departments=departments,
        allowed_roles=allowed_roles,
    )

@app.post("/admin/users/<int:user_id>/approve")
@role_required("admin", "superadmin")
def approve_user(user_id):
    u = User.query.get_or_404(user_id)
    if not can_manage_user(u):
        abort(403)
    u.is_approved = True
    db.session.commit()
    flash(f"Пользователь {u.full_name} подтверждён.")
    return redirect(url_for("admin_users"))


@app.post("/admin/users/<int:user_id>/toggle-enabled")
@role_required("admin", "superadmin")
def toggle_user_enabled(user_id):
    """Disable an account without deleting its documents, tasks, or audit trail."""
    user = User.query.get_or_404(user_id)
    if not can_manage_user(user):
        abort(403)
    if getattr(current_user, "id", None) == user.id:
        flash("Нельзя отключить собственную учётную запись.", "warning")
        return redirect(url_for("admin_users"))
    if user.role in ("admin", "superadmin") and current_user.role != "superadmin":
        abort(403)

    user.is_enabled = not user.is_enabled
    db.session.commit()
    state = "включена" if user.is_enabled else "отключена"
    flash(f"Учётная запись {user.full_name} {state}.", "success")
    return redirect(url_for("admin_users"))

@app.post("/admin/users/<int:user_id>/delete")
@role_required("admin", "superadmin")
def delete_user(user_id):
    u = User.query.get_or_404(user_id)
    if not can_manage_user(u):
        abort(403)
    if getattr(current_user, "id", None) == u.id:
        flash("Нельзя удалить самого себя.", "warning")
        return redirect(url_for("admin_users"))

    if u.role in ("admin", "superadmin") and current_user.role != "superadmin":
        flash("Удалять администраторов может только суперадминистратор.", "warning")
        return redirect(url_for("admin_users"))

    replacement_query = User.query.filter(User.id != u.id)
    if current_user.role != "superadmin" and u.department and u.department.workshop_id:
        replacement_query = replacement_query.join(User.department).filter(
            Department.workshop_id == u.department.workshop_id
        )
    replacement_user = replacement_query.order_by(
        User.is_approved.desc(), User.created_at.asc()
    ).first()
    replacement_user_id = replacement_user.id if replacement_user else None

    owned_records_count = (
        Activity.query.filter_by(owner_id=u.id).count()
        + Letter.query.filter_by(author_id=u.id).count()
        + Plan.query.filter_by(created_by=u.id).count()
        + Protocol.query.filter_by(created_by_id=u.id).count()
        + Task.query.filter_by(created_by=u.id).count()
        + GeneralDocument.query.filter_by(uploaded_by=u.id).count()
        + SharedLink.query.filter_by(created_by=u.id).count()
    )
    if owned_records_count and replacement_user_id is None:
        flash("Нельзя удалить пользователя: некому передать связанные записи.", "danger")
        return redirect(url_for("admin_users"))

    try:
        Activity.query.filter_by(owner_id=u.id).update({"owner_id": replacement_user_id}, synchronize_session=False)
        Activity.query.filter_by(approved_by_id=u.id).update({"approved_by_id": None}, synchronize_session=False)
        Letter.query.filter_by(author_id=u.id).update({"author_id": replacement_user_id}, synchronize_session=False)
        Plan.query.filter_by(created_by=u.id).update({"created_by": replacement_user_id}, synchronize_session=False)
        Protocol.query.filter_by(created_by_id=u.id).update({"created_by_id": replacement_user_id}, synchronize_session=False)
        Task.query.filter_by(created_by=u.id).update({"created_by": replacement_user_id}, synchronize_session=False)
        GeneralDocument.query.filter_by(uploaded_by=u.id).update({"uploaded_by": replacement_user_id}, synchronize_session=False)
        SharedLink.query.filter_by(created_by=u.id).update({"created_by": replacement_user_id}, synchronize_session=False)

        LetterRecipient.query.filter_by(user_id=u.id).delete(synchronize_session=False)
        PlanExecutor.query.filter_by(user_id=u.id).delete(synchronize_session=False)
        TaskAssignee.query.filter_by(user_id=u.id).delete(synchronize_session=False)
        DownloadLog.query.filter_by(user_id=u.id).delete(synchronize_session=False)
        DocumentAccess.query.filter(
            (DocumentAccess.user_id == u.id) | (DocumentAccess.granted_by == u.id)
        ).delete(synchronize_session=False)
        ChatMessage.query.filter(
            (ChatMessage.sender_id == u.id) | (ChatMessage.receiver_id == u.id)
        ).delete(synchronize_session=False)

        db.session.delete(u)
        db.session.commit()
    except Exception as exc:
        db.session.rollback()
        flash(f"Не удалось удалить пользователя: {exc}", "danger")
        return redirect(url_for("admin_users"))
    flash("Пользователь удалён.")
    return redirect(url_for("admin_users"))
# --- Админка: цехи и отделы ---
def current_workshop_id(user=None):
    user = user or current_user
    department = getattr(user, "department", None)
    return getattr(department, "workshop_id", None) if department else None


def can_manage_department(department):
    if current_user.role == "superadmin":
        return True
    workshop_id = current_workshop_id()
    return (
        current_user.role == "admin"
        and workshop_id is not None
        and department.workshop_id == workshop_id
    )


@app.get("/admin/workshops")
@role_required("superadmin")
def admin_workshops():
    workshops = Workshop.query.options(joinedload(Workshop.manager)).order_by(Workshop.name).all()
    managers = User.query.filter_by(is_approved=True, is_enabled=True).order_by(User.full_name).all()
    return render_template("admin_workshops.html", title="Цехи", workshops=workshops, managers=managers)


def workshop_manager_from_request():
    """Return an active confirmed user selected as a workshop manager."""
    raw_manager_id = (request.form.get("manager_id") or "").strip()
    if not raw_manager_id:
        return None
    if not raw_manager_id.isdigit():
        return False
    manager = db.session.get(User, int(raw_manager_id))
    if not manager or not manager.is_approved or not manager.is_enabled:
        return False
    return manager


@app.post("/admin/workshops")
@role_required("superadmin")
def add_workshop():
    name = (request.form.get("name") or "").strip()
    manager = workshop_manager_from_request()
    if not name:
        flash("Название цеха не может быть пустым.", "warning")
    elif manager is False:
        flash("Выберите действующего подтверждённого пользователя в качестве начальника.", "warning")
    elif Workshop.query.filter_by(name=name).first():
        flash("Такой цех уже существует.", "warning")
    else:
        db.session.add(Workshop(name=name, manager=manager))
        db.session.commit()
        flash("Цех «{}» добавлен.".format(name), "success")
    return redirect(url_for("admin_workshops"))


@app.post("/admin/workshops/<int:workshop_id>/edit")
@role_required("superadmin")
def edit_workshop(workshop_id):
    workshop = Workshop.query.get_or_404(workshop_id)
    name = (request.form.get("name") or "").strip()
    manager = workshop_manager_from_request()
    if not name:
        flash("Название цеха не может быть пустым.", "warning")
    elif manager is False:
        flash("Выберите действующего подтверждённого пользователя в качестве начальника.", "warning")
    elif Workshop.query.filter(Workshop.name == name, Workshop.id != workshop.id).first():
        flash("Такой цех уже существует.", "warning")
    else:
        workshop.name = name
        workshop.manager = manager
        db.session.commit()
        flash("Название цеха изменено.", "success")
    return redirect(url_for("admin_workshops"))


@app.post("/admin/workshops/<int:workshop_id>/delete")
@role_required("superadmin")
def delete_workshop(workshop_id):
    workshop = Workshop.query.get_or_404(workshop_id)
    if workshop.departments:
        flash("Сначала перенесите или удалите отделы этого цеха.", "warning")
        return redirect(url_for("admin_workshops"))
    db.session.delete(workshop)
    db.session.commit()
    flash("Цех удалён.", "success")
    return redirect(url_for("admin_workshops"))


@app.get("/admin/departments")
@role_required("admin", "superadmin")
def admin_departments():
    workshop_id = current_workshop_id()
    query = Department.query.options(joinedload(Department.workshop)).order_by(Department.name)
    if current_user.role != "superadmin":
        query = query.filter(Department.workshop_id == workshop_id) if workshop_id else query.filter(False)
    return render_template(
        "admin_departments.html", title="Отделы", deps=query.all(),
        workshops=Workshop.query.order_by(Workshop.name).all(),
        current_workshop_id=workshop_id,
        current_workshop=db.session.get(Workshop, workshop_id) if workshop_id else None,
        can_select_workshop=current_user.role == "superadmin",
    )

@app.post("/admin/departments")
@role_required("admin", "superadmin")
def add_department():
    name = request.form.get("name", "").strip()
    workshop_raw = (request.form.get("workshop_id") or "").strip()
    if not name:
        flash("Название отдела не может быть пустым.")
        return redirect(url_for("admin_departments"))
    if Department.query.filter_by(name=name).first():
        flash("Такой отдел уже существует.")
        return redirect(url_for("admin_departments"))

    workshop_id = int(workshop_raw) if workshop_raw.isdigit() else None
    if current_user.role != "superadmin":
        workshop_id = current_workshop_id()
    if not workshop_id or not db.session.get(Workshop, workshop_id):
        flash("Выберите цех для отдела.", "warning")
        return redirect(url_for("admin_departments"))
    d = Department(name=name, workshop_id=workshop_id)
    db.session.add(d)
    db.session.commit()
    flash(f"Отдел «{name}» добавлен.")
    return redirect(url_for("admin_departments"))

@app.post("/admin/departments/<int:dep_id>/edit")
@role_required("admin", "superadmin")
def edit_department(dep_id):
    d = Department.query.get_or_404(dep_id)
    if not can_manage_department(d):
        abort(403)
    name = request.form.get("name", "").strip()
    if not name:
        flash("Название не может быть пустым.")
        return redirect(url_for("admin_departments"))
    workshop_raw = (request.form.get("workshop_id") or "").strip()
    if current_user.role == "superadmin" and workshop_raw.isdigit():
        workshop = db.session.get(Workshop, int(workshop_raw))
        if not workshop:
            flash("Выбранный цех не найден.", "warning")
            return redirect(url_for("admin_departments"))
        d.workshop_id = workshop.id
    d.name = name
    db.session.commit()
    flash("Изменения сохранены.")
    return redirect(url_for("admin_departments"))

@app.post("/admin/departments/<int:dep_id>/delete")
@role_required("admin", "superadmin")
def delete_department(dep_id):
    d = Department.query.get_or_404(dep_id)
    if not can_manage_department(d):
        abort(403)
    db.session.delete(d)
    db.session.commit()
    flash(f"Отдел «{d.name}» удалён.")
    return redirect(url_for("admin_departments"))
# --- Админка: типы активностей ---
@app.get("/admin/activity-types")
@role_required("admin")
def admin_activity_types():
    types = ActivityType.query.order_by(ActivityType.name).all()
    return render_template("admin_activity_types.html", title="Типы активностей", types=types)

@app.post("/admin/activity-types")
@role_required("admin")
def add_activity_type():
    name = request.form.get("name", "").strip()
    if not name:
        flash("Название типа активности не может быть пустым.")
        return redirect(url_for("admin_activity_types"))
    if ActivityType.query.filter_by(name=name).first():
        flash("Такой тип уже существует.")
        return redirect(url_for("admin_activity_types"))

    t = ActivityType(name=name)
    db.session.add(t)
    db.session.commit()
    flash(f"Тип активности «{name}» добавлен.")
    return redirect(url_for("admin_activity_types"))

@app.post("/admin/activity-types/<int:type_id>/edit")
@role_required("admin")
def edit_activity_type(type_id):
    t = ActivityType.query.get_or_404(type_id)
    name = request.form.get("name", "").strip()
    if not name:
        flash("Название не может быть пустым.")
        return redirect(url_for("admin_activity_types"))
    t.name = name
    db.session.commit()
    flash("Изменения сохранены.")
    return redirect(url_for("admin_activity_types"))

@app.post("/admin/activity-types/<int:type_id>/delete")
@role_required("admin")
def delete_activity_type(type_id):
    t = ActivityType.query.get_or_404(type_id)
    db.session.delete(t)
    db.session.commit()
    flash(f"Тип активности «{t.name}» удалён.")
    return redirect(url_for("admin_activity_types"))
@app.get("/activities/<int:activity_id>/create-letter")
@login_required

def new_letter_from_activity(activity_id):
    activity = Activity.query.get_or_404(activity_id)
    deps = Department.query.order_by(Department.name).all()
    users = User.query.filter_by(is_approved=True).order_by(User.full_name).all()
    sources = LetterSource.query.order_by(LetterSource.name).all()

    # создаём черновик письма с автозаполнением
    subject = f"Письмо по активности №{activity.id} — {activity.title}"
    body = f"Описание активности:\n{activity.description or ''}\n\nСтатус: {activity.status}"

    return render_template(
        "letter_create.html",
        title="Письмо на основе активности",
        deps=deps,
        users=users,
        sources=sources,
        subject=subject,
        body=body,
        from_activity=activity,
        now=datetime.now  # ← вот это добавляем
    )
# ✅ AJAX: переключение статуса "выполнено"
@app.post("/logs/<int:log_id>/toggle_done")
@login_required
def toggle_log_done(log_id: int):
    log = ActivityLog.query.get_or_404(log_id)
    if activity_work_is_locked(log.activity) or not can_modify_activity_work(log.activity):
        abort(403)
    log.is_done = not log.is_done
    add_activity_history(
        log.activity,
        "Журнал действий",
        "{} выполнение действия: {}".format("Отмечено" if log.is_done else "Снята отметка", log.text[:300]),
    )
    db.session.commit()
    return {"success": True, "is_done": log.is_done}


# ✅ AJAX: добавление поддействия
@app.post("/logs/<int:log_id>/add_child")
@login_required
def add_child_log(log_id: int):
    parent = ActivityLog.query.get_or_404(log_id)
    if activity_work_is_locked(parent.activity) or not can_modify_activity_work(parent.activity):
        abort(403)
    text = request.form.get("text", "").strip()
    if not text:
        return {"success": False, "error": "Пустой текст"}, 400

    child = ActivityLog(activity_id=parent.activity_id, parent_id=log_id, text=text)
    db.session.add(child)
    add_activity_history(parent.activity, "Журнал действий", "Добавлено поддействие: {}".format(text[:300]))
    db.session.commit()

    return {
        "success": True,
        "child": {
            "id": child.id,
            "text": child.text,
            "is_done": child.is_done,
            "entry_date": child.entry_date.strftime("%d.%m.%Y %H:%M")
        }
    }
# ✅ AJAX: удаление действия (и всех поддействий)
@app.post("/logs/<int:log_id>/delete")
@login_required
def delete_log(log_id: int):
    log = ActivityLog.query.get_or_404(log_id)
    if activity_work_is_locked(log.activity) or not can_modify_activity_work(log.activity):
        abort(403)

    # Рекурсивное удаление всех поддействий
    def delete_children(node):
        for child in node.children:
            delete_children(child)
            db.session.delete(child)

    deleted_text = log.text
    delete_children(log)
    db.session.delete(log)
    add_activity_history(log.activity, "Журнал действий", "Удалено действие: {}".format(deleted_text[:300]))
    db.session.commit()

    return jsonify(success=True)
@app.post("/logs/<int:log_id>/toggle_done")
@login_required
def toggle_done(log_id: int):
    log = ActivityLog.query.get_or_404(log_id)
    if activity_work_is_locked(log.activity) or not can_modify_activity_work(log.activity):
        abort(403)
    log.is_done = not log.is_done
    add_activity_history(
        log.activity,
        "Журнал действий",
        "{} выполнение действия: {}".format("Отмечено" if log.is_done else "Снята отметка", log.text[:300]),
    )
    db.session.commit()
    return jsonify(success=True, is_done=log.is_done)
@app.post("/logs/<int:log_id>/edit")
@login_required
def edit_log(log_id: int):
    log = ActivityLog.query.get_or_404(log_id)
    if activity_work_is_locked(log.activity) or not can_modify_activity_work(log.activity):
        abort(403)
    text = request.form.get("text", "").strip()
    if not text:
        return jsonify(success=False, error="Пустой текст"), 400
    previous_text = log.text
    log.text = text
    add_activity_history(
        log.activity,
        "Журнал действий",
        "Изменено действие: «{}» -> «{}»".format(previous_text[:150], text[:150]),
    )
    db.session.commit()
    return jsonify(success=True)


@app.route('/reports', methods=['GET', 'POST'])
@role_required("admin")
def manage_reports():
    if request.method == 'POST':
        rid = request.form.get('id')
        name = request.form.get('name')
        sql_text = request.form.get('sql')
        json_data = request.form.get('json')
        js_code = request.form.get('js')

        if rid:
            db.session.execute(
                text("UPDATE otchets SET name=:n, sql=:s, json=:j, js=:k WHERE id=:id"),
                {"n": name, "s": sql_text, "j": json_data, "k": js_code, "id": rid},
            )
        else:
            db.session.execute(
                text("INSERT INTO otchets (name, sql, json, js) VALUES (:n, :s, :j, :k)"),
                {"n": name, "s": sql_text, "j": json_data, "k": js_code},
            )
        db.session.commit()
        flash('Отчёт сохранён', 'success')
        return redirect(url_for('manage_reports'))

    rows = db.session.execute(text("SELECT * FROM otchets ORDER BY id DESC")).fetchall()
    return render_template('otchet.html', rows=rows)

@app.context_processor
def inject_now():
    from datetime import datetime
    return {'now': datetime.now}


@app.context_processor
def inject_template_helpers():
    return {
        "has_endpoint": lambda endpoint: endpoint in app.view_functions,
    }
@app.get("/admin/values")
@login_required
def admin_values_list():
    if current_user.role != "admin":
        abort(403)
    values = ValueTemplate.query.order_by(ValueTemplate.rusname).all()
    return render_template("admin_values.html", title="Шаблоны значений", values=values)

@app.post("/admin/values")
@login_required
def admin_values_create():
    if current_user.role != "admin":
        abort(403)
    name = request.form.get("name","").strip()
    rusname = request.form.get("rusname","").strip()
    vtype = request.form.get("type","string").strip()
    if not name or not rusname:
        flash("Имя и заголовок обязательны","warning")
        return redirect(url_for("admin_values_list"))
    vt = ValueTemplate(name=name, rusname=rusname, type=vtype, is_active=True)
    db.session.add(vt)
    db.session.commit()
    flash("Шаблон добавлен","success")
    return redirect(url_for("admin_values_list"))

@app.post("/admin/values/<int:vid>/toggle")
@login_required
def admin_values_toggle(vid):
    if current_user.role != "admin":
        abort(403)
    vt = ValueTemplate.query.get_or_404(vid)
    vt.is_active = not vt.is_active
    db.session.commit()
    return redirect(url_for("admin_values_list"))

@app.post("/admin/values/<int:vid>/delete")
@login_required
def admin_values_delete(vid):
    if current_user.role != "admin":
        abort(403)
    vt = ValueTemplate.query.get_or_404(vid)
    db.session.delete(vt)
    db.session.commit()
    flash("Шаблон удалён","info")
    return redirect(url_for("admin_values_list"))

# ====== API: получить активные шаблоны в JSON (для JS на форме активности) ======



# ====== ДОРАБОТКА СОЗДАНИЯ АКТИВНОСТИ: читаем direction/decision/extra_values ======

@app.get("/api/value-templates")
@login_required
def api_value_templates():
    templates = ValueTemplate.query.filter_by(is_active=True).all()
    data = [
        {"id": v.id, "name": v.name, "rusname": v.rusname, "type": v.type}
        for v in templates
    ]
    return jsonify(data)
# -------------------------------
#  Управление шаблонами значений
# -------------------------------
@app.get("/admin/value-templates")
@login_required
def value_templates_page():
    if current_user.role != "admin":
        abort(403)
    values = ValueTemplate.query.order_by(ValueTemplate.rusname).all()
    return render_template("value_templates.html", values=values, title="Шаблоны значений")


@app.post("/admin/value-templates/add")
@login_required
def add_value_template():
    if current_user.role != "admin":
        abort(403)
    name = request.form.get("name")
    rusname = request.form.get("rusname")
    type_ = request.form.get("type")
    if not name or not rusname:
        flash("Имя и описание обязательны", "warning")
        return redirect(url_for("value_templates_page"))
    vt = ValueTemplate(name=name.strip(), rusname=rusname.strip(), type=type_)
    db.session.add(vt)
    db.session.commit()
    flash(f"Добавлен шаблон: {rusname}", "success")
    return redirect(url_for("value_templates_page"))


@app.post("/admin/value-templates/<int:vt_id>/toggle")
@login_required
def toggle_value_template(vt_id):
    if current_user.role != "admin":
        abort(403)
    vt = ValueTemplate.query.get_or_404(vt_id)
    vt.is_active = not vt.is_active
    db.session.commit()
    flash(f"Шаблон {'включён' if vt.is_active else 'отключён'}", "info")
    return redirect(url_for("value_templates_page"))


@app.post("/admin/value-templates/<int:vt_id>/delete")
@login_required
def delete_value_template(vt_id):
    if current_user.role != "admin":
        abort(403)
    vt = ValueTemplate.query.get_or_404(vt_id)
    db.session.delete(vt)
    db.session.commit()
    flash("Шаблон удалён", "danger")
    return redirect(url_for("value_templates_page"))
@app.post("/activities/<int:activity_id>/add_value")
@login_required
def add_activity_value(activity_id):
    a = Activity.query.get_or_404(activity_id)
    if activity_work_is_locked(a) or not can_modify_activity_work(a):
        abort(403)
    tpl_id = request.form.get("template_id")
    val = request.form.get("value")

    tpl = ValueTemplate.query.get_or_404(tpl_id)

    # --- если extra_values хранится как строка, преобразуем ---
    if not a.extra_values:
        values = []
    else:
        try:
            if isinstance(a.extra_values, str):
                values = json.loads(a.extra_values)
            else:
                values = a.extra_values
        except Exception:
            values = []

    # --- добавляем новое значение ---
    values.append({
        "id": tpl.id,
        "name": tpl.name,
        "rusname": tpl.rusname,
        "type": tpl.type,
        "value": val
    })

    # --- сохраняем список обратно как JSON ---
    a.extra_values = json.dumps(values, ensure_ascii=False)
    add_activity_history(a, "Дополнительное значение", "Добавлено значение: {}.".format(tpl.rusname))
    db.session.commit()

    return jsonify(success=True)

@app.post("/activities/<int:activity_id>/delete_value/<int:value_id>")
@login_required
def delete_activity_value(activity_id, value_id):
    a = Activity.query.get_or_404(activity_id)
    if activity_work_is_locked(a) or not can_modify_activity_work(a):
        abort(403)
    if isinstance(a.extra_values, str):
        a.extra_values = json.loads(a.extra_values)
    deleted = next((value for value in a.extra_values if int(value.get("id", 0)) == value_id), None)
    a.extra_values = [v for v in a.extra_values if int(v.get("id", 0)) != value_id]
    if deleted:
        add_activity_history(
            a,
            "Дополнительное значение",
            "Удалено значение: {}.".format(deleted.get("rusname") or deleted.get("name") or value_id),
        )
    db.session.commit()
    return jsonify(success=True)
@app.post("/letters/<int:letter_id>/save_extras")
@login_required
def save_letter_extras(letter_id):
    import json
    le = Letter.query.get_or_404(letter_id)
    if le.author_id != current_user.id and current_user.role not in ("admin", "superadmin"):
        abort(403)
    data = request.get_json()
    new_extras = data.get("extras", [])

    # Загружаем старые значения, если есть
    existing = []
    if le.extra_values:
        try:
            existing = json.loads(le.extra_values)
        except Exception as e:
            print("⚠️ Ошибка парсинга JSON:", e)

    # Добавляем новые к старым
    existing.extend(new_extras)

    # Сохраняем обратно
    le.extra_values = existing
    db.session.commit()

    print(f"✅ Обновлённые extras для письма #{le.id}:", le.extra_values)
    return jsonify({"success": True, "count": len(existing)})
# =========================
# Планы (создание, исполнители, активности)
# =========================
PLAN_ITEM_STATUSES = {
    "planned": {"label": "Запланировано", "class": "bg-secondary"},
    "done": {"label": "Выполнено", "class": "bg-success"},
    "not_done": {"label": "Невыполнено", "class": "bg-danger"},
    "cancelled": {"label": "Отменено", "class": "bg-dark"},
}
PLAN_DEADLINES = {
    "month": "В течение месяца",
    "q1": "1 квартал",
    "q2": "2 квартал",
    "q3": "3 квартал",
    "q4": "4 квартал",
    "date": "До указанной даты",
}


def get_plan_period(month_value):
    from calendar import monthrange
    from datetime import date, datetime

    if month_value:
        try:
            period_start = datetime.strptime(month_value, "%Y-%m").date().replace(day=1)
        except ValueError:
            return None
    else:
        period_start = date.today().replace(day=1)
    period_end = period_start.replace(day=monthrange(period_start.year, period_start.month)[1])
    return period_start, period_end


def plan_item_status_meta(status):
    return PLAN_ITEM_STATUSES.get(status, {"label": status or "-", "class": "bg-secondary"})


def parse_plan_deadline_date(deadline_kind, raw_date):
    if deadline_kind != "date":
        return None
    try:
        return dt.datetime.strptime((raw_date or "").strip(), "%Y-%m-%d").date()
    except ValueError:
        return None


def parse_plan_hours(raw_hours):
    """Normalize planned workload; an omitted value is stored as zero hours."""
    if raw_hours is None or str(raw_hours).strip() == "":
        return 0.0
    try:
        hours = float(str(raw_hours).strip().replace(",", "."))
    except (TypeError, ValueError):
        return None
    return round(hours, 2) if 0 <= hours <= 10000 else None


def build_plan_department_rows(items):
    """Return one consistent aggregation for department and employee reports."""
    departments = {}
    employees = {}
    for item in items:
        user = item.executor
        department = user.department if user else None
        department_id = department.id if department else 0
        department_name = department.name if department else "Без отдела"
        department_row = departments.setdefault(department_id, {
            "id": department_id if department else None, "name": department_name,
            "total": 0, "planned": 0, "done": 0, "not_done": 0, "cancelled": 0,
            "planned_hours": 0, "done_hours": 0, "not_done_hours": 0,
            "executor_map": {},
        })
        employee_row = department_row["executor_map"].setdefault(item.executor_id, {
            "user": user,
            "items": [],
            "total": 0, "planned": 0, "done": 0, "not_done": 0, "cancelled": 0,
            "planned_hours": 0, "done_hours": 0, "not_done_hours": 0,
        })
        planned_hours = float(item.planned_hours or 0)
        for row in (department_row, employee_row):
            row["total"] += 1
            row[item.status] = row.get(item.status, 0) + 1
            row["planned_hours"] += planned_hours
            if item.status == "done":
                row["done_hours"] += planned_hours
            elif item.status == "not_done":
                row["not_done_hours"] += planned_hours
        employee_row["items"].append(item)
        employees[item.executor_id] = employee_row

    department_rows = list(departments.values())
    for row in department_rows:
        row["completion"] = round((row["done"] / row["total"]) * 100, 1) if row["total"] else 0
        row["executors"] = list(row.pop("executor_map").values())
        for employee in row["executors"]:
            employee["completion"] = round((employee["done"] / employee["total"]) * 100, 1) if employee["total"] else 0
        row["executors"].sort(key=lambda employee: (employee["user"].full_name or "").lower())
    department_rows.sort(key=lambda row: (row["name"] or "").lower())
    executor_rows = sorted(employees.values(), key=lambda row: (row["user"].full_name or "").lower())
    return department_rows, executor_rows


def deputy_can_review_plan(plan, user=None):
    """План доступен заместителю, если он относится к его отделу.

    План может быть создан сотрудником до добавления пунктов, поэтому учитываем
    и отдел автора, и исполнителей пунктов. Чужие планы при этом не открываются.
    """
    user = user or current_user
    if getattr(user, "role", None) != "deputy" or not getattr(user, "department_id", None):
        return False
    creator = getattr(plan, "creator", None)
    if creator is not None and creator.department_id == user.department_id:
        return True
    return any(
        item.executor is not None and item.executor.department_id == user.department_id
        for item in plan.items
    )


class ListPagination:
    """Small pagination adapter for a permission-filtered in-memory collection."""
    def __init__(self, records, page, per_page=25):
        self.total = len(records)
        self.per_page = per_page
        self.pages = max(1, (self.total + per_page - 1) // per_page)
        self.page = min(max(1, page), self.pages)
        first = (self.page - 1) * per_page
        self.items = records[first:first + per_page]

    @property
    def has_prev(self):
        return self.page > 1

    @property
    def prev_num(self):
        return self.page - 1

    @property
    def has_next(self):
        return self.page < self.pages

    @property
    def next_num(self):
        return self.page + 1

    def iter_pages(self, left_edge=1, left_current=2, right_current=2, right_edge=1):
        last = 0
        for number in range(1, self.pages + 1):
            if (
                number <= left_edge
                or number > self.pages - right_edge
                or self.page - left_current <= number <= self.page + right_current
            ):
                if last + 1 != number:
                    yield None
                yield number
                last = number


@app.route("/plans", methods=["GET", "POST"])
@login_required
def plans_list():
    from datetime import datetime, date

    # --- фильтрация по дате ---
    start_str = request.args.get("date_from")
    end_str = request.args.get("date_to")
    selected_executor_id = (request.args.get("executor_id") or "").strip()
    page = request.args.get("page", 1, type=int) or 1

    q = Plan.query.options(
        joinedload(Plan.creator),
        joinedload(Plan.items).joinedload(PlanItem.executor),
    )

    if start_str:
        try:
            date_from = datetime.strptime(start_str, "%Y-%m-%d").date()
            q = q.filter(Plan.start_date >= date_from)
        except ValueError:
            pass

    if end_str:
        try:
            date_to = datetime.strptime(end_str, "%Y-%m-%d").date()
            q = q.filter(Plan.end_date <= date_to)
        except ValueError:
            pass

    is_manager = current_user.role in ("admin", "superadmin")
    is_deputy = current_user.role == "deputy"
    filter_users = []
    if is_manager:
        filter_users = User.query.filter_by(is_approved=True).order_by(User.full_name).all()
    elif is_deputy and current_user.department_id:
        q = q.filter(or_(
            Plan.creator.has(User.department_id == current_user.department_id),
            Plan.items.any(PlanItem.executor.has(User.department_id == current_user.department_id)),
        ))
        filter_users = User.query.filter_by(
            is_approved=True, department_id=current_user.department_id
        ).order_by(User.full_name).all()
    else:
        q = q.filter(Plan.items.any(PlanItem.executor_id == current_user.id))

    if selected_executor_id:
        try:
            executor_id = int(selected_executor_id)
            if any(user.id == executor_id for user in filter_users):
                q = q.filter(Plan.items.any(PlanItem.executor_id == executor_id))
            else:
                selected_executor_id = ""
        except ValueError:
            selected_executor_id = ""

    plans = q.order_by(Plan.start_date.desc(), Plan.id.desc()).all()
    if is_deputy:
        plans = [plan for plan in plans if deputy_can_review_plan(plan)]
    pagination = ListPagination(plans, page, per_page=25)
    plans = pagination.items
    plan_summary = {}
    for plan in plans:
        item_counts = {status: 0 for status in PLAN_ITEM_STATUSES}
        for item in plan.items:
            item_counts[item.status] = item_counts.get(item.status, 0) + 1
        item_counts["planned_hours"] = round(sum(float(item.planned_hours or 0) for item in plan.items), 2)
        plan_summary[plan.id] = item_counts
    return render_template(
        "plans.html", title="Планы", plans=plans, date_from=start_str, date_to=end_str,
        plan_summary=plan_summary, is_manager=is_manager, is_deputy=is_deputy,
        can_review_plans=is_manager or is_deputy, filter_users=filter_users,
        selected_executor_id=selected_executor_id, deadline_options=PLAN_DEADLINES,
        pagination=pagination,
    )


@app.get("/plans/review")
@role_required("admin", "superadmin", "deputy")
def plans_review():
    """Рабочее место администратора для принятия месячных планов."""
    month_value = (request.args.get("month") or dt.date.today().strftime("%Y-%m")).strip()
    status = (request.args.get("status") or "submitted").strip()
    period = get_plan_period(month_value)
    if not period:
        flash("Выберите корректный месяц.", "warning")
        return redirect(url_for("plans_review"))
    if status not in {"all", "submitted", "revision", "approved"}:
        status = "submitted"

    period_start, period_end = period
    query = (
        Plan.query.options(
            joinedload(Plan.creator),
            joinedload(Plan.items).joinedload(PlanItem.executor),
            joinedload(Plan.items).joinedload(PlanItem.linked_activity),
        )
        .filter(Plan.start_date <= period_end, Plan.end_date >= period_start)
    )
    if status != "all":
        query = query.filter(Plan.approval_status == status)
    plans = query.order_by(Plan.created_at.desc(), Plan.id.desc()).all()
    if current_user.role == "deputy":
        plans = [plan for plan in plans if deputy_can_review_plan(plan)]

    return render_template(
        "plans_review.html",
        title="Проверка планов",
        plans=plans,
        month_value=month_value,
        selected_status=status,
        deadline_options=PLAN_DEADLINES,
        get_plan_status_meta=plan_item_status_meta,
    )


def plan_action_redirect(plan):
    """После действия на странице проверки сохраняем выбранный фильтр."""
    if request.form.get("return_to") == "review":
        values = {"month": request.form.get("month") or plan.start_date.strftime("%Y-%m")}
        status = request.form.get("review_status")
        if status:
            values["status"] = status
        return redirect(url_for("plans_review", **values))
    return redirect(url_for("plan_detail", plan_id=plan.id))


def find_matching_plan_for_activity(activity):
    """Return the executor's plan covering the task start date, if one exists."""
    if not activity.owner_id or not activity.start_date:
        return None
    task_date = activity.start_date.date()
    return (
        Plan.query
        .filter(
            Plan.start_date <= task_date,
            Plan.end_date >= task_date,
            or_(
                Plan.created_by == activity.owner_id,
                Plan.items.any(PlanItem.executor_id == activity.owner_id),
            ),
        )
        .order_by(
            (Plan.items.any(PlanItem.executor_id == activity.owner_id)).desc(),
            Plan.start_date.desc(),
            Plan.id.desc(),
        )
        .first()
    )


def add_pending_activities_to_plan(plan):
    """Attach opted-in tasks when a plan for their executor and month appears."""
    executor_ids = {item.executor_id for item in plan.items if item.executor_id}
    if not executor_ids:
        return 0

    period_start = dt.datetime.combine(plan.start_date, dt.time.min)
    period_end = dt.datetime.combine(plan.end_date + dt.timedelta(days=1), dt.time.min)
    already_linked_ids = {
        item.linked_activity_id for item in plan.items if item.linked_activity_id
    }
    activities = (
        Activity.query
        .filter(
            Activity.plan_requested.is_(True),
            Activity.plan_id.is_(None),
            Activity.owner_id.in_(executor_ids),
            Activity.start_date >= period_start,
            Activity.start_date < period_end,
        )
        .order_by(Activity.start_date, Activity.id)
        .all()
    )
    position = max((item.position for item in plan.items), default=0)
    added = 0
    for activity in activities:
        if activity.id in already_linked_ids:
            activity.plan_requested = False
            continue
        position += 1
        deadline_date = activity.end_date.date() if activity.end_date else None
        db.session.add(PlanItem(
            plan_id=plan.id,
            executor_id=activity.owner_id,
            position=position,
            task_text=activity.title,
            deadline_kind="date" if deadline_date else "month",
            deadline_date=deadline_date,
            planned_hours=calculate_activity_effort_hours(activity) or 0,
            linked_activity_id=activity.id,
            status="planned",
        ))
        activity.plan_id = plan.id
        activity.plan_requested = False
        add_activity_history(
            activity,
            "Добавление в план",
            "Задача автоматически добавлена в план «{}».".format(plan.plan_text),
        )
        added += 1
    return added


@app.route("/plans/create", methods=["GET", "POST"])
@login_required
def plan_create():
    if request.method == "POST":
        period = get_plan_period((request.form.get("month") or "").strip())
        if not period:
            flash("Выберите месяц и год плана.", "warning")
            return redirect(url_for("plan_create"))
        start_date, end_date = period
        plan_text = (request.form.get("plan_text") or "").strip() or f"План на {start_date.strftime('%m.%Y')}"
        existing_plan = Plan.query.filter_by(created_by=current_user.id, start_date=start_date).first()
        if existing_plan:
            flash("На выбранный месяц у вас уже создан план. Откройте его для редактирования.", "warning")
            return redirect(url_for("plan_detail", plan_id=existing_plan.id))
        executor_ids = request.form.getlist("executor_id")
        tasks = request.form.getlist("task_text")
        deadlines = request.form.getlist("deadline_kind")
        deadline_dates = request.form.getlist("deadline_date")
        planned_hours_values = request.form.getlist("planned_hours")
        justifications = request.form.getlist("justification")

        rows = []
        for index, task in enumerate(tasks):
            task = (task or "").strip()
            executor_id = executor_ids[index] if index < len(executor_ids) else ""
            deadline = deadlines[index] if index < len(deadlines) else "month"
            deadline_date = parse_plan_deadline_date(
                deadline, deadline_dates[index] if index < len(deadline_dates) else ""
            )
            if not task and not executor_id:
                continue
            planned_hours = parse_plan_hours(planned_hours_values[index] if index < len(planned_hours_values) else "")
            justification = (justifications[index] if index < len(justifications) else "").strip() or None
            if not task or not executor_id or deadline not in PLAN_DEADLINES or (deadline == "date" and not deadline_date) or planned_hours is None:
                flash("В каждой строке укажите исполнителя, мероприятие и срок. Плановые часы можно оставить пустыми.", "warning")
                return redirect(url_for("plan_create"))
            rows.append((int(executor_id), task, justification, deadline, deadline_date, planned_hours))

        if not rows:
            flash("Добавьте хотя бы один пункт плана.", "warning")
            return redirect(url_for("plan_create"))

        is_manager = current_user.role in ("admin", "superadmin")
        if not is_manager and any(executor_id != current_user.id for executor_id, _, _, _, _, _ in rows):
            abort(403)
        p = Plan(
            start_date=start_date,
            end_date=end_date,
            plan_text=plan_text,
            created_by=current_user.id,
            approval_status="approved" if is_manager else "submitted",
            approved_by_id=current_user.id if is_manager else None,
            approved_at=dt.datetime.utcnow() if is_manager else None,
        )
        db.session.add(p)
        db.session.flush()
        for position, (executor_id, task, justification, deadline, deadline_date, planned_hours) in enumerate(rows, start=1):
            db.session.add(PlanItem(
                plan_id=p.id, executor_id=executor_id, position=position,
                task_text=task, justification=justification, deadline_kind=deadline, deadline_date=deadline_date,
                planned_hours=planned_hours, status="planned",
            ))
        db.session.flush()
        auto_added = add_pending_activities_to_plan(p)
        db.session.commit()
        message = "План и его пункты созданы"
        if auto_added:
            message += ". Задач автоматически добавлено: {}".format(auto_added)
        flash(message, "success")
        return redirect(url_for("plan_detail", plan_id=p.id))

    is_manager = current_user.role in ("admin", "superadmin")
    users = User.query.filter_by(is_approved=True).order_by(User.full_name).all() if is_manager else [current_user]
    return render_template(
        "plan_create.html", title="Создание плана", users=users,
        deadline_options=PLAN_DEADLINES, is_manager=is_manager,
    )


@app.route("/plans/<int:plan_id>")
@login_required
def plan_detail(plan_id):
    plan = Plan.query.get_or_404(plan_id)
    is_manager = current_user.role in ("admin", "superadmin")
    can_review = is_manager or deputy_can_review_plan(plan)
    own_items = [item for item in plan.items if item.executor_id == current_user.id]
    is_creator = plan.created_by == current_user.id
    if not can_review and not own_items and not is_creator:
        abort(403)
    items = plan.items if can_review else own_items
    users = User.query.filter_by(is_approved=True).order_by(User.full_name).all() if is_manager else []
    task_owner_ids = {item.executor_id for item in items}
    available_tasks = (
        Activity.query.options(joinedload(Activity.owner), joinedload(Activity.type))
        .filter(Activity.archived_at.is_(None), Activity.owner_id.in_(task_owner_ids or {-999999}))
        .order_by(Activity.owner_id, Activity.start_date.desc(), Activity.id.desc())
        .all()
    )
    return render_template(
        "plan_detail.html", title="План", plan=plan, items=items, users=users,
        is_manager=is_manager, can_review=can_review, is_creator=is_creator,
        can_edit_definition=is_manager or (is_creator and plan.approval_status != "approved"),
        deadline_options=PLAN_DEADLINES,
        status_options=PLAN_ITEM_STATUSES, get_plan_status_meta=plan_item_status_meta,
        available_tasks=available_tasks,
    )


@app.post("/plans/<int:plan_id>/items")
@login_required
def plan_item_create(plan_id):
    plan = Plan.query.get_or_404(plan_id)
    is_manager = current_user.role in ("admin", "superadmin")
    if not is_manager and (plan.created_by != current_user.id or plan.approval_status == "approved"):
        abort(403)
    task_text = (request.form.get("task_text") or "").strip()
    justification = (request.form.get("justification") or "").strip() or None
    executor_id = request.form.get("executor_id", type=int)
    deadline_kind = request.form.get("deadline_kind", "month")
    deadline_date = parse_plan_deadline_date(deadline_kind, request.form.get("deadline_date"))
    planned_hours = parse_plan_hours(request.form.get("planned_hours"))
    if not task_text or not executor_id or deadline_kind not in PLAN_DEADLINES or (deadline_kind == "date" and not deadline_date) or planned_hours is None:
        flash("Заполните исполнителя, мероприятие и срок. Плановые часы можно оставить пустыми.", "warning")
        return redirect(url_for("plan_detail", plan_id=plan.id))
    if not is_manager and executor_id != current_user.id:
        abort(403)
    position = (max((item.position for item in plan.items), default=0) + 1)
    db.session.add(PlanItem(
        plan_id=plan.id, executor_id=executor_id, position=position,
        task_text=task_text, justification=justification, deadline_kind=deadline_kind, deadline_date=deadline_date,
        planned_hours=planned_hours,
    ))
    db.session.commit()
    flash("Пункт плана добавлен", "success")
    return redirect(url_for("plan_detail", plan_id=plan.id))


@app.post("/plan-items/<int:item_id>/update")
@login_required
def plan_item_update(item_id):
    item = PlanItem.query.get_or_404(item_id)
    plan = item.plan
    is_manager = current_user.role in ("admin", "superadmin")
    if not is_manager and item.executor_id != current_user.id:
        abort(403)

    if item.execution_approved and not is_manager:
        flash("Выполнение уже принято администратором: пункт доступен только для просмотра.", "warning")
        return redirect(url_for("plan_detail", plan_id=item.plan_id))

    can_update_result = is_manager or plan.approval_status == "approved"
    if can_update_result:
        status = request.form.get("status", item.status)
        if status not in PLAN_ITEM_STATUSES:
            abort(400)
        item.status = status
        item.comment = (request.form.get("comment") or "").strip() or None
        linked_activity_id = request.form.get("linked_activity_id", type=int)
        if linked_activity_id:
            linked_activity = db.session.get(Activity, linked_activity_id)
            if not linked_activity or linked_activity.archived_at or linked_activity.owner_id != item.executor_id:
                flash("Можно привязать только действующую задачу исполнителя пункта плана.", "warning")
                return redirect(url_for("plan_detail", plan_id=item.plan_id))
            item.linked_activity_id = linked_activity.id
        elif request.form.get("linked_activity_id") == "":
            item.linked_activity_id = None
        if status == "done" and not item.completed_at:
            item.completed_at = dt.datetime.utcnow()
    if is_manager or (plan.created_by == current_user.id and plan.approval_status != "approved"):
        item.task_text = (request.form.get("task_text") or item.task_text).strip()
        item.justification = (request.form.get("justification") or "").strip() or None
        deadline_kind = request.form.get("deadline_kind", item.deadline_kind)
        if deadline_kind in PLAN_DEADLINES:
            deadline_date = parse_plan_deadline_date(deadline_kind, request.form.get("deadline_date"))
            if deadline_kind == "date" and not deadline_date:
                flash("Для срока «До указанной даты» выберите дату.", "warning")
                return redirect(url_for("plan_detail", plan_id=item.plan_id))
            item.deadline_kind = deadline_kind
            item.deadline_date = deadline_date
        planned_hours = parse_plan_hours(request.form.get("planned_hours"))
        if planned_hours is None:
            flash("Плановые часы должны быть числом от 0 до 10000.", "warning")
            return redirect(url_for("plan_detail", plan_id=item.plan_id))
        item.planned_hours = planned_hours
        executor_id = request.form.get("executor_id", type=int)
        if executor_id and is_manager:
            item.executor_id = executor_id
    db.session.commit()
    flash("Результат по пункту плана сохранен", "success")
    return redirect(url_for("plan_detail", plan_id=item.plan_id))


@app.post("/plan-items/<int:item_id>/delete")
@login_required
def plan_item_delete(item_id):
    item = PlanItem.query.get_or_404(item_id)
    plan = item.plan
    is_manager = current_user.role in ("admin", "superadmin")
    if not is_manager and (plan.created_by != current_user.id or plan.approval_status == "approved" or item.executor_id != current_user.id):
        abort(403)
    plan_id = item.plan_id
    db.session.delete(item)
    db.session.commit()
    flash("Пункт плана удален", "success")
    return redirect(url_for("plan_detail", plan_id=plan_id))


@app.post("/plans/<int:plan_id>/delete")
@role_required("admin", "superadmin")
def plan_delete(plan_id):
    plan = Plan.query.get_or_404(plan_id)
    # Legacy activities can remain in the task register after their plan is removed.
    Activity.query.filter_by(plan_id=plan.id).update({"plan_id": None}, synchronize_session=False)
    db.session.delete(plan)
    db.session.commit()
    flash("План удален.", "success")
    return redirect(url_for("plans_list"))


@app.post("/plans/<int:plan_id>/approve")
@role_required("admin", "superadmin", "deputy")
def plan_approve(plan_id):
    plan = Plan.query.get_or_404(plan_id)
    if current_user.role == "deputy" and not deputy_can_review_plan(plan):
        abort(403)
    plan.approval_status = "approved"
    plan.approved_by_id = current_user.id if current_user.id != -1 else None
    plan.approved_at = dt.datetime.utcnow()
    plan.revision_comment = None
    db.session.commit()
    flash("План принят. Исполнитель теперь может указывать только результат выполнения.", "success")
    return plan_action_redirect(plan)


@app.post("/plans/<int:plan_id>/return-for-revision")
@role_required("admin", "superadmin", "deputy")
def plan_return_for_revision(plan_id):
    plan = Plan.query.get_or_404(plan_id)
    if current_user.role == "deputy" and not deputy_can_review_plan(plan):
        abort(403)
    comment = (request.form.get("revision_comment") or "").strip()
    if not comment:
        flash("Укажите, что требуется доработать.", "warning")
        return plan_action_redirect(plan)
    plan.approval_status = "revision"
    plan.approved_by_id = None
    plan.approved_at = None
    plan.revision_comment = comment
    db.session.commit()
    flash("План возвращен на доработку.", "warning")
    return plan_action_redirect(plan)


@app.post("/plans/<int:plan_id>/resubmit")
@login_required
def plan_resubmit(plan_id):
    plan = Plan.query.get_or_404(plan_id)
    if plan.created_by != current_user.id or plan.approval_status != "revision":
        abort(403)
    plan.approval_status = "submitted"
    db.session.commit()
    flash("План повторно направлен администратору на принятие.", "success")
    return redirect(url_for("plan_detail", plan_id=plan.id))


@app.post("/plan-items/<int:item_id>/approve-execution")
@role_required("admin", "superadmin")
def plan_item_approve_execution(item_id):
    item = PlanItem.query.get_or_404(item_id)
    if item.status == "planned":
        flash("Сначала исполнитель должен указать результат выполнения.", "warning")
        return redirect(url_for("plan_detail", plan_id=item.plan_id))
    item.execution_approved = True
    item.execution_approved_by_id = current_user.id if current_user.id != -1 else None
    item.execution_approved_at = dt.datetime.utcnow()
    db.session.commit()
    flash("Выполнение принято. Пункт переведен в режим только для чтения.", "success")
    return redirect(url_for("plan_detail", plan_id=item.plan_id))


@app.route("/plans/<int:plan_id>/executors", methods=["GET", "POST"])
@role_required("admin")
def plan_executors(plan_id):
    plan = Plan.query.get_or_404(plan_id)
    users = User.query.filter_by(is_approved=True).order_by(User.full_name).all()

    if request.method == "POST":
        user_id = request.form.get("user_id")
        topic = request.form.get("topic")
        start_date = request.form.get("start_date")
        end_date = request.form.get("end_date")

        pe = PlanExecutor(
            plan_id=plan.id,
            user_id=user_id,
            topic=topic,
            start_date=dt.datetime.strptime(start_date, "%Y-%m-%d").date(),
            end_date=dt.datetime.strptime(end_date, "%Y-%m-%d").date(),
        )
        db.session.add(pe)
        db.session.commit()

        # создаём активность для исполнителя
        a = Activity(
            title=f"Задача по плану: {topic}",
            description=plan.plan_text,
            owner_id=user_id,
            created_by_id=current_user.id if current_user.id != -1 else None,
            start_date=pe.start_date,
            end_date=pe.end_date,
            status="open",
            plan_id=plan.id,
            type_id=1001
        )
        db.session.add(a)
        db.session.commit()

        flash("Исполнитель добавлен и активность создана", "success")
        return redirect(url_for("plan_executors", plan_id=plan.id))

    return render_template("plan_executors.html", title="Исполнители", plan=plan, users=users)


@app.get("/plans/<int:plan_id>/activities")
@login_required
def plan_activities(plan_id):
    plan = Plan.query.get_or_404(plan_id)
    acts = Activity.query.filter_by(plan_id=plan.id).order_by(Activity.start_date.desc()).all()
    return render_template("plan_activities.html", title="Активности по плану", plan=plan, activities=acts)
@app.post("/plans/<int:plan_id>/executors/ajax")
@role_required("admin")
def plan_executors_add_ajax(plan_id):
    plan = Plan.query.get_or_404(plan_id)
    data = request.get_json()
    try:
        user_id = int(data.get("user_id"))
        topic = data.get("topic", "").strip()
        start_date = dt.datetime.strptime(data.get("start_date"), "%Y-%m-%d").date()
        end_date = dt.datetime.strptime(data.get("end_date"), "%Y-%m-%d").date()
    except Exception as e:
        return jsonify(success=False, error=f"Неверные данные: {e}")

    if not topic or not user_id:
        return jsonify(success=False, error="Не указаны исполнитель или тема")

    pe = PlanExecutor(
        plan_id=plan.id,
        user_id=user_id,
        topic=topic,
        start_date=start_date,
        end_date=end_date
    )
    db.session.add(pe)
    db.session.commit()

    # 🔹 создаём активность
    act = Activity(
        title=f"Задача по плану: {topic}",
        description=plan.plan_text,
        owner_id=user_id,
        created_by_id=current_user.id if current_user.id != -1 else None,
        start_date=start_date,
        end_date=end_date,
        status="open",
        plan_id=plan.id,    type_id=1001
    )
    db.session.add(act)
    db.session.commit()

    return jsonify(
        success=True,
        executor={
            "id": pe.id,
            "user_name": pe.user.full_name,
            "topic": pe.topic,
            "start_date": pe.start_date.strftime("%Y-%m-%d"),
            "end_date": pe.end_date.strftime("%Y-%m-%d"),
            "status": pe.status,
        }
    )
@app.post("/plans/<int:plan_id>/toggle_complete")
@role_required("admin")
def toggle_plan_complete(plan_id):
    plan = Plan.query.get_or_404(plan_id)
    plan.is_completed = not plan.is_completed
    db.session.commit()
    return jsonify(success=True, is_completed=plan.is_completed)
# === 📅 Единый календарь сроков ===
@app.route("/calendar")
@login_required
def deadlines_calendar():
    return render_template("plans_calendar.html", title="Календарь сроков")


@app.route("/plans/calendar")
@login_required
def plans_calendar():
    """Keep the old plans URL working while the calendar becomes a global section."""
    return redirect(url_for("deadlines_calendar"))

# === API для подгрузки сроков задач и планов ===
@app.route("/api/activities/events")
@login_required
def activities_events():
    selected_kinds = set(filter(None, (request.args.get("kinds") or "task,plan,postponed,overdue").split(",")))
    today = dt.date.today()
    privileged_roles = ("admin", "manager", "superadmin", "deputy")
    is_privileged = current_user.role in privileged_roles

    activity_query = (
        Activity.query.options(joinedload(Activity.owner))
        .filter(Activity.archived_at.is_(None))
    )
    if not is_privileged:
        department_id = getattr(current_user, "department_id", None)
        activity_query = activity_query.filter(or_(
            Activity.owner_id == current_user.id,
            Activity.assigned_department_id == department_id if department_id else False,
        ))
    events = []
    for activity in activity_query.all():
        deadline = activity.end_date.date() if activity.end_date else None
        start_date = activity.start_date.date() if activity.start_date else None
        postponed_to = activity.postponed_to.date() if activity.postponed_to else None
        is_postponed = activity.status == "postponed" and postponed_to and postponed_to >= today
        is_overdue = (
            deadline and deadline < today
            and activity.status in ("open", "in_progress", "postponed")
            and not is_postponed
        )
        if is_postponed:
            kind, event_date, color, prefix = "postponed", postponed_to, "#f28c28", "Перенос"
        elif is_overdue:
            kind, event_date, color, prefix = "overdue", deadline, "#dc3545", "Просрочено"
        else:
            kind, event_date, color, prefix = (
                "task", deadline or start_date, "#0dcaf0",
                "Задача" if deadline else "Начало задачи",
            )
        if not event_date or kind not in selected_kinds:
            continue
        events.append({
            "id": "activity-{}-{}".format(activity.id, kind),
            "title": "{}: {}".format(prefix, activity.title),
            "start": event_date.isoformat(),
            "allDay": True,
            "color": color,
            "kind": kind,
            "status": activity.status,
            "owner": activity.owner.full_name if activity.owner else "-",
            "url": url_for("activity_view", activity_id=activity.id),
        })

    if "plan" in selected_kinds:
        plan_items_query = PlanItem.query.options(joinedload(PlanItem.plan), joinedload(PlanItem.executor))
        if not is_privileged:
            plan_items_query = plan_items_query.filter(PlanItem.executor_id == current_user.id)
        for item in plan_items_query.all():
            plan = item.plan
            if not plan:
                continue
            event_date = item.deadline_date or plan.end_date
            if not event_date:
                continue
            status_meta = plan_item_status_meta(item.status)
            events.append({
                "id": "plan-item-{}".format(item.id),
                "title": "План: {}".format(item.task_text),
                "start": event_date.isoformat(),
                "allDay": True,
                "color": "#6f42c1" if item.status == "planned" else "#198754" if item.status == "done" else "#6c757d",
                "kind": "plan",
                "status": status_meta["label"],
                "owner": item.executor.full_name if item.executor else "-",
                "deadline": PLAN_DEADLINES.get(item.deadline_kind, "Срок плана"),
                "url": url_for("plan_detail", plan_id=plan.id),
            })

    return jsonify(events)
# === 🔔 API: напоминания о сроках задач, планов и графиков ===
def plan_item_reminder_deadline(item):
    if item.deadline_date:
        return item.deadline_date
    if item.deadline_kind in ("q1", "q2", "q3", "q4") and item.plan and item.plan.start_date:
        from calendar import monthrange
        quarter_end_month = int(item.deadline_kind[-1]) * 3
        year = item.plan.start_date.year
        return dt.date(year, quarter_end_month, monthrange(year, quarter_end_month)[1])
    return item.plan.end_date if item.plan else None


def deadline_reminder_prefix(deadline, today):
    days_left = (deadline - today).days
    if days_left < 0:
        return "Просрочено на {} дн.".format(abs(days_left)), days_left
    if days_left in (6, 3, 1):
        return "Срок через {} дн.".format(days_left), days_left
    return None, days_left


@app.route("/api/notifications")
@login_required
def notifications():
    today = dt.date.today()
    tasks = Activity.query.filter(
        Activity.owner_id == current_user.id,
        Activity.status.in_(["open", "in_progress", "postponed"]),
        Activity.end_date.isnot(None),
    ).all()
    reminders = []
    for task in tasks:
        deadline = task.end_date.date() if isinstance(task.end_date, dt.datetime) else task.end_date
        prefix, days_left = deadline_reminder_prefix(deadline, today)
        if not prefix:
            continue
        reminders.append({"title": "Задача. {}: {}".format(prefix, task.title), "days_left": days_left})

    plan_items = (
        PlanItem.query.options(joinedload(PlanItem.plan))
        .filter(PlanItem.executor_id == current_user.id, PlanItem.status == "planned")
        .all()
    )
    for item in plan_items:
        deadline = plan_item_reminder_deadline(item)
        if not deadline:
            continue
        prefix, days_left = deadline_reminder_prefix(deadline, today)
        if prefix:
            reminders.append({
                "title": "План. {}: {}".format(prefix, item.task_text),
                "days_left": days_left,
            })

    schedule_items = (
        ScheduleItem.query.options(joinedload(ScheduleItem.schedule), joinedload(ScheduleItem.activity))
        .filter(
            ScheduleItem.responsible_id == current_user.id,
            ScheduleItem.status.in_(["planned", "in_progress"]),
        )
        .all()
    )
    for item in schedule_items:
        deadline = item.activity.end_date.date() if item.activity and item.activity.end_date else item.item_date
        prefix, days_left = deadline_reminder_prefix(deadline, today)
        if prefix:
            schedule_title = item.schedule.title if item.schedule else "График"
            reminders.append({
                "title": "График «{}». {}: {}".format(schedule_title, prefix, item.text),
                "days_left": days_left,
            })

    count = len(reminders)
    if count == 0:
        return jsonify({"count": 0})
    return jsonify({
        "count": count,
        "titles": [item["title"] for item in reminders],
        "message": "У вас {} пункт(ов), требующих внимания по сроку.".format(count),
    })


@app.route("/notifications/send", methods=["GET", "POST"])
@role_required("admin", "superadmin", "deputy")
def send_user_notification():
    users = User.query.filter_by(is_approved=True).order_by(User.full_name).all()
    departments = Department.query.order_by(Department.name).all()

    if request.method == "POST":
        target_type = (request.form.get("target_type") or "").strip()
        target_id = (request.form.get("target_id") or "").strip()
        title = (request.form.get("title") or "").strip()
        message = (request.form.get("message") or "").strip()
        scheduled_raw = (request.form.get("scheduled_at") or "").strip()
        background_color = (request.form.get("background_color") or "#0d6efd").strip().lower()
        recipients = []

        if not title or not message:
            flash("Заполните заголовок и текст уведомления.", "warning")
            return redirect(request.url)
        if not re.fullmatch(r"#[0-9a-f]{6}", background_color):
            flash("Цвет фона должен быть указан в формате #RRGGBB.", "warning")
            return redirect(request.url)
        if scheduled_raw:
            try:
                scheduled_at = dt.datetime.strptime(scheduled_raw, "%Y-%m-%dT%H:%M")
            except ValueError:
                flash("Укажите дату и время отправки.", "warning")
                return redirect(request.url)
            if scheduled_at < dt.datetime.now() - dt.timedelta(minutes=1):
                flash("Время отправки не может быть в прошлом.", "warning")
                return redirect(request.url)
        else:
            scheduled_at = dt.datetime.now()

        if target_type == "all":
            recipients = users
        elif target_type == "user" and target_id.isdigit():
            user = db.session.get(User, int(target_id))
            recipients = [user] if user and user.is_approved else []
        elif target_type == "department" and target_id.isdigit():
            recipients = User.query.filter_by(
                department_id=int(target_id), is_approved=True
            ).order_by(User.full_name).all()

        if not recipients:
            flash("Не удалось найти получателей уведомления.", "warning")
            return redirect(request.url)

        sender_id = current_user.id if current_user.id != -1 else None
        db.session.add_all([
            UserNotification(
                user_id=user.id,
                sender_id=sender_id,
                title=title,
                message=message,
                scheduled_at=scheduled_at,
                background_color=background_color,
            )
            for user in recipients
        ])
        db.session.commit()
        if scheduled_raw:
            flash("Уведомление запланировано на {}: получателей {}.".format(scheduled_at.strftime("%d.%m.%Y %H:%M"), len(recipients)), "success")
        else:
            flash("Уведомление отправлено: получателей {}.".format(len(recipients)), "success")
        return redirect(url_for("send_user_notification"))

    return render_template(
        "notification_send.html",
        title="Отправка уведомления",
        users=users,
        departments=departments,
    )


@app.get("/notifications")
@login_required
def user_notifications_page():
    now = dt.datetime.now()
    notifications_list = (
        UserNotification.query
        .filter(UserNotification.user_id == current_user.id)
        .filter(or_(UserNotification.scheduled_at.is_(None), UserNotification.scheduled_at <= now))
        .order_by(UserNotification.scheduled_at.desc(), UserNotification.id.desc())
        .limit(100)
        .all()
    )
    unread = [item for item in notifications_list if item.read_at is None]
    if unread:
        for item in unread:
            item.read_at = now
        db.session.commit()
    return render_template(
        "notifications.html",
        title="Уведомления",
        notifications_list=notifications_list,
    )


@app.get("/api/user-notifications")
@login_required
def user_notifications_api():
    now = dt.datetime.now()
    notifications_list = (
        UserNotification.query
        .filter(UserNotification.user_id == current_user.id)
        .filter(or_(UserNotification.scheduled_at.is_(None), UserNotification.scheduled_at <= now))
        .order_by(UserNotification.scheduled_at.desc(), UserNotification.id.desc())
        .limit(20)
        .all()
    )
    unread_count = sum(item.read_at is None for item in notifications_list)
    return jsonify({
        "unread_count": unread_count,
        "items": [
            {
                "id": item.id,
                "title": item.title,
                "message": item.message,
                "created_at": item.created_at.strftime("%d.%m.%Y %H:%M"),
                "scheduled_at": (item.scheduled_at or item.created_at).strftime("%d.%m.%Y %H:%M"),
                "background_color": item.background_color or "#0d6efd",
                "is_read": item.read_at is not None,
            }
            for item in notifications_list
        ],
    })


@app.post("/api/user-notifications/<int:notification_id>/read")
@login_required
def mark_user_notification_read(notification_id):
    notification = UserNotification.query.filter_by(
        id=notification_id, user_id=current_user.id
    ).first_or_404()
    if notification.read_at is None:
        notification.read_at = dt.datetime.utcnow()
        db.session.commit()
    return jsonify(success=True)
@app.route("/plans/statistics")
@role_required("admin", "superadmin")
def plans_statistics():
    shared_payload = get_report_share_payload("plans_statistics")
    month_value = str((shared_payload or {}).get("month") or request.args.get("month") or "").strip()
    period = get_plan_period(month_value)
    if not period:
        flash("Неверный месяц.", "warning")
        return redirect(url_for("plans_statistics"))
    period_start, period_end = period
    month_value = period_start.strftime("%Y-%m")
    workshop_id = get_report_workshop_id(shared_payload)

    items_query = (
        PlanItem.query.options(joinedload(PlanItem.linked_activity)).join(Plan).join(User, PlanItem.executor_id == User.id)
        .filter(Plan.start_date <= period_end, Plan.end_date >= period_start)
    )
    if workshop_id:
        items_query = items_query.join(Department, User.department_id == Department.id).filter(Department.workshop_id == workshop_id)
    items = items_query.order_by(PlanItem.executor_id, PlanItem.position).all()
    is_manager = shared_payload is not None or getattr(current_user, "role", None) in ("admin", "superadmin")
    if not is_manager:
        items = [item for item in items if item.executor_id == current_user.id]

    totals = {status: sum(1 for item in items if item.status == status) for status in PLAN_ITEM_STATUSES}
    department_rows, rows = build_plan_department_rows(items)

    total = len(items)
    summary = {
        "total": total,
        "planned": totals["planned"],
        "done": totals["done"],
        "not_done": totals["not_done"],
        "cancelled": totals["cancelled"],
        "completion": round((totals["done"] / total) * 100, 1) if total else 0,
        "planned_hours": round(sum(float(item.planned_hours or 0) for item in items), 2),
        "done_hours": round(sum(float(item.planned_hours or 0) for item in items if item.status == "done"), 2),
        "not_done_hours": round(sum(float(item.planned_hours or 0) for item in items if item.status == "not_done"), 2),
    }
    chart_data = {
        "status_labels": ["Выполнено", "Невыполнено", "Отменено", "Запланировано"],
        "status_values": [summary["done"], summary["not_done"], summary["cancelled"], summary["planned"]],
        "executor_labels": [row["user"].full_name for row in rows],
        "executor_done": [row["done"] for row in rows],
        "executor_not_done": [row["not_done"] for row in rows],
        "executor_planned": [row["planned"] for row in rows],
        "department_labels": [row["name"] for row in department_rows],
        "department_done": [row["done"] for row in department_rows],
        "department_not_done": [row["not_done"] for row in department_rows],
        "department_planned": [row["planned"] for row in department_rows],
        "department_hours": [round(row["planned_hours"], 2) for row in department_rows],
        "department_done_hours": [round(row["done_hours"], 2) for row in department_rows],
    }
    return render_template(
        "plans_statistics.html", title="Аналитика выполнения планов",
        month_value=month_value, period_start=period_start, period_end=period_end,
        summary=summary, department_rows=department_rows, executor_rows=rows, chart_data=chart_data,
    )


@app.route("/plans/monthly-report")
@role_required("admin", "superadmin")
def plans_monthly_report():
    shared_payload = get_report_share_payload("plans_monthly_report")
    month_value = str((shared_payload or {}).get("month") or request.args.get("month") or "").strip()
    period = get_plan_period(month_value)
    if not period:
        flash("Неверный месяц.", "warning")
        return redirect(url_for("plans_monthly_report"))
    period_start, period_end = period
    month_value = period_start.strftime("%Y-%m")
    workshop_id = get_report_workshop_id(shared_payload)
    items_query = (
        PlanItem.query.options(joinedload(PlanItem.linked_activity)).join(Plan).join(User, PlanItem.executor_id == User.id)
        .filter(Plan.start_date <= period_end, Plan.end_date >= period_start)
    )
    if workshop_id:
        items_query = items_query.join(Department, User.department_id == Department.id).filter(Department.workshop_id == workshop_id)
    items = items_query.order_by(User.full_name, PlanItem.position).all()
    plans_count = len({item.plan_id for item in items})
    department_rows, executor_rows = build_plan_department_rows(items)
    counts = {status: sum(1 for item in items if item.status == status) for status in PLAN_ITEM_STATUSES}
    summary = {
        "plans_count": plans_count, "departments_count": len(department_rows),
        "executors_count": len(executor_rows), "items_total": len(items),
        "planned_count": counts["planned"], "done_count": counts["done"],
        "not_done_count": counts["not_done"], "cancelled_count": counts["cancelled"],
        "completion_percent": round((counts["done"] / len(items)) * 100, 1) if items else 0,
        "planned_hours_total": round(sum(float(item.planned_hours or 0) for item in items), 2),
        "done_hours_total": round(sum(float(item.planned_hours or 0) for item in items if item.status == "done"), 2),
        "not_done_hours_total": round(sum(float(item.planned_hours or 0) for item in items if item.status == "not_done"), 2),
    }

    return render_template(
        "plans_monthly_report.html",
        title="Сводный отчет по планам",
        month_value=month_value,
        period_start=period_start,
        period_end=period_end,
        department_rows=department_rows,
        summary=summary,
        get_plan_status_meta=plan_item_status_meta,
        deadline_options=PLAN_DEADLINES,
    )


@app.get("/plans/monthly-report/export")
@role_required("admin", "superadmin")
def plans_monthly_report_export():
    """Export the selected monthly plan report in an Excel-compatible CSV file."""
    month_value = (request.args.get("month") or "").strip()
    period = get_plan_period(month_value)
    if not period:
        flash("Выберите месяц для экспорта отчета.", "warning")
        return redirect(url_for("plans_monthly_report"))
    period_start, period_end = period
    items = (
        PlanItem.query.options(
            joinedload(PlanItem.executor).joinedload(User.department),
            joinedload(PlanItem.linked_activity),
        )
        .join(Plan)
        .filter(Plan.start_date <= period_end, Plan.end_date >= period_start)
        .order_by(User.full_name, PlanItem.position)
        .all()
    )

    payload = io.StringIO(newline="")
    writer = csv.writer(payload, delimiter=";")
    writer.writerow(["СВОДНЫЙ ОТЧЕТ ПО ВЫПОЛНЕНИЮ ПЛАНОВ"])
    writer.writerow(["Период", "{} - {}".format(period_start.strftime("%d.%m.%Y"), period_end.strftime("%d.%m.%Y"))])
    writer.writerow([])
    writer.writerow([
        "Отдел", "Исполнитель", "№", "Мероприятие", "Связанная задача", "Обоснования",
        "Срок выполнения", "Плановые часы", "Результат", "Комментарий",
    ])
    for item in items:
        deadline = (
            "До {}".format(item.deadline_date.strftime("%d.%m.%Y"))
            if item.deadline_kind == "date" and item.deadline_date
            else PLAN_DEADLINES.get(item.deadline_kind, item.deadline_kind)
        )
        status = plan_item_status_meta(item.status)["label"]
        department = item.executor.department.name if item.executor and item.executor.department else "Без отдела"
        linked_task = "#{} {}".format(item.linked_activity.id, item.linked_activity.title) if item.linked_activity else ""
        justification_text = (
            item.linked_activity.outgoing_document
            if item.linked_activity and item.linked_activity.outgoing_document
            else item.justification or "-"
        )
        writer.writerow([
            department,
            item.executor.full_name if item.executor else "",
            item.position,
            item.task_text,
            linked_task,
            justification_text,
            deadline,
            item.planned_hours or 0,
            status,
            item.comment or "",
        ])

    filename = "otchet_po_planam_{}.csv".format(period_start.strftime("%Y-%m"))
    response = app.response_class(
        payload.getvalue().encode("utf-8-sig"),
        mimetype="text/csv; charset=utf-8",
    )
    response.headers["Content-Disposition"] = "attachment; filename={}".format(filename)
    return response


@app.route("/my-plans")
@login_required
def my_plans():
    from datetime import date

    # Получаем все планы, где текущий пользователь — исполнитель
    executors = PlanExecutor.query.filter_by(user_id=current_user.id).all()
    plan_ids = list({ex.plan_id for ex in executors})

    plans = Plan.query.filter(Plan.id.in_(plan_ids)).all()
    my_plans_data = []

    for p in plans:
        # Все активности этого пользователя в рамках плана
        my_activities = Activity.query.filter_by(plan_id=p.id, owner_id=current_user.id).all()

        total = len(my_activities)
        done = sum(1 for a in my_activities if a.status in ("done", "approved"))
        progress = round((done / total) * 100, 1) if total > 0 else 0

        # Дни до конца плана
        days_left = (p.end_date - date.today()).days if p.end_date else None

        my_plans_data.append({
            "id": p.id,
            "start_date": p.start_date,
            "end_date": p.end_date,
            "text": p.plan_text,
            "progress": progress,
            "days_left": days_left,
        })

    return render_template("my_plans.html", title="Мои планы", plans=my_plans_data)
@app.get("/exports/<string:kind>.<string:format_name>")
@login_required
def export_journal(kind, format_name):
    if format_name != "csv" or kind not in ("activities", "letters", "memos", "plans"):
        abort(404)
    today = dt.date.today()
    date_from = request.args.get("date_from", "")
    date_to = request.args.get("date_to", "")
    headers, rows, title = [], [], ""
    if kind == "activities":
        query = Activity.query.options(joinedload(Activity.owner), joinedload(Activity.type))
        if date_from:
            query = query.filter(Activity.start_date >= dt.datetime.strptime(date_from, "%Y-%m-%d"))
        if date_to:
            query = query.filter(Activity.start_date <= dt.datetime.combine(dt.datetime.strptime(date_to, "%Y-%m-%d").date(), dt.datetime.max.time()))
        headers, title = ["№", "Название", "Исполнитель", "Статус", "Срок"], "Журнал задач"
        rows = [[a.id, a.title, a.owner.full_name if a.owner else "-", get_activity_status_meta(a.status)["label"], a.end_date.strftime("%d.%m.%Y") if a.end_date else "-"] for a in query.order_by(Activity.id.desc()).all()]
    elif kind in ("letters", "memos"):
        memo_source = LetterSource.query.filter_by(name="Служебная записка").first()
        query = Letter.query.options(joinedload(Letter.executor), joinedload(Letter.source))
        query = query.filter(Letter.source_id == memo_source.id) if kind == "memos" and memo_source else (query.filter(Letter.source_id != memo_source.id) if kind == "letters" and memo_source else query)
        headers, title = ["Номер", "Дата", "Тема", "Исполнитель"], "Журнал служебных записок" if kind == "memos" else "Журнал писем"
        rows = [[item.reg_number or item.id, item.letter_date.strftime("%d.%m.%Y") if item.letter_date else "-", item.subject, item.executor.full_name if item.executor else "-"] for item in query.order_by(Letter.id.desc()).all()]
    else:
        headers, title = ["№", "Период", "Автор", "Статус"], "Журнал планов"
        rows = [[item.id, "{} - {}".format(item.start_date.strftime("%d.%m.%Y"), item.end_date.strftime("%d.%m.%Y")), item.creator.full_name if item.creator else "-", item.approval_status] for item in Plan.query.options(joinedload(Plan.creator)).order_by(Plan.id.desc()).all()]
    payload = io.StringIO(newline="")
    writer = csv.writer(payload, delimiter=";")
    writer.writerow([title]); writer.writerow(headers); writer.writerows(rows)
    response = app.response_class(payload.getvalue().encode("utf-8-sig"), mimetype="text/csv; charset=utf-8")
    response.headers["Content-Disposition"] = "attachment; filename={}.csv".format(kind)
    return response

@app.route("/search", methods=["GET"])
@login_required
def global_search():
    """Single search across tasks, letters, memos, and attachment text."""
    q = request.args.get("q", "").strip()
    results_acts, results_lets = [], []
    task_results, letter_results, memo_results = [], [], []

    if q:
        like = "%{}%".format(q)
        task_results = Activity.query.filter(or_(
            Activity.title.ilike(like), Activity.description.ilike(like),
            Activity.direction.ilike(like), Activity.outgoing_document.ilike(like),
        )).order_by(Activity.start_date.desc()).limit(50).all()
        memo_source = LetterSource.query.filter_by(name="Служебная записка").first()
        letters_query = Letter.query.filter(or_(
            Letter.subject.ilike(like), Letter.body.ilike(like), Letter.reg_number.ilike(like),
        ))
        if memo_source:
            letter_results = letters_query.filter(Letter.source_id != memo_source.id).order_by(Letter.letter_date.desc()).limit(50).all()
            memo_results = letters_query.filter(Letter.source_id == memo_source.id).order_by(Letter.letter_date.desc()).limit(50).all()
        else:
            letter_results = letters_query.order_by(Letter.letter_date.desc()).limit(50).all()
        # === Поиск по активности ===
        sql_a = text("""
        SELECT ad.id AS doc_id, ad.filename, a.id AS parent_id, a.title AS parent_title,
               'activity' AS source,
               snippet(activity_documents_fts, 0, '<mark>', '</mark>', '...', 20) AS snippet
        FROM activity_documents_fts
        JOIN activity_documents ad ON ad.id = activity_documents_fts.rowid
        LEFT JOIN activities a ON a.id = ad.activity_id
        WHERE activity_documents_fts MATCH :q
        ORDER BY rank;
        """)
        results_acts = db.session.execute(sql_a, {"q": q}).fetchall()

        # === Поиск по письмам ===
        sql_l = text("""
        SELECT ld.id AS doc_id, ld.filename, l.id AS parent_id, l.subject AS parent_title,
               'letter' AS source,
               snippet(letter_documents_fts, 0, '<mark>', '</mark>', '...', 20) AS snippet
        FROM letter_documents_fts
        JOIN letter_documents ld ON ld.id = letter_documents_fts.rowid
        LEFT JOIN letters l ON l.id = ld.letter_id
        WHERE letter_documents_fts MATCH :q
        ORDER BY rank;
        """)
        results_lets = db.session.execute(sql_l, {"q": q}).fetchall()

    return render_template(
        "search_results.html",
        title="Поиск по документам",
        q=q,
        task_results=task_results,
        letter_results=letter_results,
        memo_results=memo_results,
        results_acts=results_acts,
        results_lets=results_lets
    )

def init_fts_indexes():
    """Создание виртуальных таблиц FTS5 и триггеров"""
    with db.engine.begin() as conn:
        # === activity_documents ===
        conn.exec_driver_sql("""
        CREATE VIRTUAL TABLE IF NOT EXISTS activity_documents_fts
        USING fts5(doc_rec, activity_id UNINDEXED, content='activity_documents', content_rowid='id');
        """)
        conn.exec_driver_sql("""
        CREATE TRIGGER IF NOT EXISTS activity_documents_ai AFTER INSERT ON activity_documents
        BEGIN
            INSERT INTO activity_documents_fts(rowid, doc_rec, activity_id)
            VALUES (new.id, new.doc_rec, new.activity_id);
        END;
        """)
        conn.exec_driver_sql("""
        CREATE TRIGGER IF NOT EXISTS activity_documents_ad AFTER DELETE ON activity_documents
        BEGIN
            DELETE FROM activity_documents_fts WHERE rowid = old.id;
        END;
        """)
        conn.exec_driver_sql("""
        CREATE TRIGGER IF NOT EXISTS activity_documents_au AFTER UPDATE ON activity_documents
        BEGIN
            UPDATE activity_documents_fts
            SET doc_rec = new.doc_rec, activity_id = new.activity_id
            WHERE rowid = old.id;
        END;
        """)
        conn.exec_driver_sql("""
        INSERT INTO activity_documents_fts(rowid, doc_rec, activity_id)
        SELECT id, doc_rec, activity_id FROM activity_documents
        WHERE id NOT IN (SELECT rowid FROM activity_documents_fts);
        """)

        # === letter_documents ===
        conn.exec_driver_sql("""
        CREATE VIRTUAL TABLE IF NOT EXISTS letter_documents_fts
        USING fts5(doc_rec, letter_id UNINDEXED, content='letter_documents', content_rowid='id');
        """)
        conn.exec_driver_sql("""
        CREATE TRIGGER IF NOT EXISTS letter_documents_ai AFTER INSERT ON letter_documents
        BEGIN
            INSERT INTO letter_documents_fts(rowid, doc_rec, letter_id)
            VALUES (new.id, new.doc_rec, new.letter_id);
        END;
        """)
        conn.exec_driver_sql("""
        CREATE TRIGGER IF NOT EXISTS letter_documents_ad AFTER DELETE ON letter_documents
        BEGIN
            DELETE FROM letter_documents_fts WHERE rowid = old.id;
        END;
        """)
        conn.exec_driver_sql("""
        CREATE TRIGGER IF NOT EXISTS letter_documents_au AFTER UPDATE ON letter_documents
        BEGIN
            UPDATE letter_documents_fts
            SET doc_rec = new.doc_rec, letter_id = new.letter_id
            WHERE rowid = old.id;
        END;
        """)
        conn.exec_driver_sql("""
        INSERT INTO letter_documents_fts(rowid, doc_rec, letter_id)
        SELECT id, doc_rec, letter_id FROM letter_documents
        WHERE id NOT IN (SELECT rowid FROM letter_documents_fts);
        """)

        print("📚 FTS5 индексы инициализированы.")
from flask import send_file, abort
import os

# === Скачивание документа активности ===
@app.route("/documents/<int:doc_id>/download")
@login_required
def download_activity_document(doc_id):
    doc = ActivityDocument.query.get_or_404(doc_id)
    allowed = False

    if current_user.role in ("admin", "manager"):
        allowed = True
    elif doc.activity.owner_id == current_user.id:
        allowed = True
    else:
        allowed = (
            DocumentAccess.query.filter_by(
                doc_type="activity",
                doc_id=doc.id,
                user_id=current_user.id
            ).count() > 0
        )

    if not allowed:
        # Автор документа
        author = User.query.get(doc.activity.owner_id)
        return render_template(
            "document_request_access.html",
            title="Доступ к документу запрещён",
            doc=doc,
            author=author,
            doc_type="activity"
        ), 403

    file_path = os.path.join("static", doc.filepath)
    if not os.path.exists(file_path):
        abort(404, "Файл отсутствует на сервере")

    # 🧾 Лог загрузки
    db.session.add(DownloadLog(
        user_id=current_user.id,
        doc_type="activity",
        doc_id=doc.id
    ))
    db.session.commit()

    return send_file(file_path, as_attachment=True, download_name=doc.filename)

# === Скачивание документа письма ===
@app.route("/letters/documents/<int:doc_id>/download")
@login_required
def download_letter_document(doc_id):
    doc = LetterDocument.query.get_or_404(doc_id)
    allowed = False

    if current_user.role in ("admin", "manager"):
        allowed = True
    elif doc.letter.author_id == current_user.id:
        allowed = True
    else:
        allowed = (
            DocumentAccess.query.filter_by(
                doc_type="letter",
                doc_id=doc.id,
                user_id=current_user.id
            ).count() > 0
        )

    if not allowed:
        author = User.query.get(doc.letter.author_id)
        return render_template(
            "document_request_access.html",
            title="Доступ к документу запрещён",
            doc=doc,
            author=author,
            doc_type="letter"
        ), 403

    file_path = os.path.join("static", doc.filepath)
    if not os.path.exists(file_path):
        abort(404, "Файл отсутствует на сервере")

    db.session.add(DownloadLog(
        user_id=current_user.id,
        doc_type="letter",
        doc_id=doc.id
    ))
    db.session.commit()

    return send_file(file_path, as_attachment=True, download_name=doc.filename)

@app.get("/document_access")
@login_required
def document_access_manager():
    """Управление доступом к документам (активности + письма)"""
    if current_user.role in ("admin", "manager"):
        activity_docs = ActivityDocument.query.all()
        letter_docs = LetterDocument.query.all()
    else:
        activity_docs = (
            ActivityDocument.query
            .join(Activity)
            .filter(Activity.owner_id == current_user.id)
            .all()
        )
        letter_docs = (
            LetterDocument.query
            .join(Letter)
            .filter(Letter.author_id == current_user.id)
            .all()
        )

    # соберём все доступы сразу
    access_map = {}
    for acc in DocumentAccess.query.all():
        access_map.setdefault((acc.doc_type, acc.doc_id), []).append(acc)

    users = User.query.order_by(User.full_name).all()

    return render_template(
        "document_access_manager.html",
        activity_docs=activity_docs,
        letter_docs=letter_docs,
        users=users,
        access_map=access_map,
        is_admin=current_user.role in ("admin", "manager")
    )

@app.post("/document_access/grant")
@login_required
def grant_doc_access():
    doc_type = request.form["doc_type"]
    doc_id = int(request.form["doc_id"])
    user_id = int(request.form["user_id"])

    # Автор или админ может выдавать
    if not has_document_access(current_user.id, doc_type, doc_id):
        return jsonify({"error": "Нет прав выдавать доступ"}), 403

    exists = DocumentAccess.query.filter_by(
        doc_type=doc_type, doc_id=doc_id, user_id=user_id
    ).first()
    if exists:
        return jsonify({"error": "Уже есть доступ"}), 400

    acc = DocumentAccess(
        doc_type=doc_type,
        doc_id=doc_id,
        user_id=user_id,
        granted_by=current_user.id
    )
    db.session.add(acc)
    db.session.commit()
    return jsonify({"success": True})


@app.post("/document_access/revoke")
@login_required
def revoke_doc_access():
    doc_type = request.form["doc_type"]
    doc_id = int(request.form["doc_id"])
    user_id = int(request.form["user_id"])

    # Автор или админ может отзывать
    if not has_document_access(current_user.id, doc_type, doc_id):
        return jsonify({"error": "Нет прав отзывать доступ"}), 403

    DocumentAccess.query.filter_by(
        doc_type=doc_type, doc_id=doc_id, user_id=user_id
    ).delete()
    db.session.commit()
    return jsonify({"success": True})
@app.route("/chat")
@login_required
def chat_page():
    users = User.query.filter(User.id != current_user.id, User.is_approved == True).all()
    return render_template("chat.html", title="Чат", users=users)


@app.post("/api/chat/send")
@login_required
def chat_send():
    text = request.form.get("text", "").strip()
    receiver_id = request.form.get("receiver_id")
    file = request.files.get("file")

    file_path = None
    if file and file.filename:
        filename = secure_filename_rus(file.filename)
        upload_dir = os.path.join(app.root_path, "static", "chat_uploads")
        os.makedirs(upload_dir, exist_ok=True)
        unique_name = f"{uuid.uuid4().hex}_{filename}"
        file.save(os.path.join(upload_dir, unique_name))
        file_path = f"chat_uploads/{unique_name}"

    msg = ChatMessage(
        sender_id=current_user.id,
        receiver_id=int(receiver_id) if receiver_id and receiver_id != "all" else None,
        text=text,
        file_path=file_path
    )
    db.session.add(msg)
    db.session.commit()
    return jsonify(success=True, message=serialize_chat_message(msg))

@app.get("/api/chat/messages")
@login_required
def chat_messages():
    """Сообщения либо общего чата, либо между пользователями"""
    user_id = request.args.get("user_id")

    q = ChatMessage.query

    if user_id and user_id != "all":
        uid = int(user_id)
        q = q.filter(
            or_(
                and_(ChatMessage.sender_id == current_user.id, ChatMessage.receiver_id == uid),
                and_(ChatMessage.sender_id == uid, ChatMessage.receiver_id == current_user.id)
            )
        )
    else:
        q = q.filter(ChatMessage.receiver_id.is_(None))  # общий чат

    q = q.order_by(ChatMessage.created_at.desc()).limit(50)

    msgs = reversed(q.all())
    data = [
        serialize_chat_message(m)
        for m in msgs
    ]
    return jsonify(data)
@app.get("/api/chat/users")
@login_required
def chat_users():
    """Возвращает список всех активных пользователей для чата"""
    users = User.query.filter(User.id != current_user.id, User.is_approved == True).all()
    data = [
        {"id": u.id, "name": u.full_name, "role": u.role}
        for u in users
    ]
    return jsonify(data)


@app.get("/api/chat/incoming")
@login_required
def chat_incoming():
    direct_after_raw = (request.args.get("direct_after_id") or "0").strip()
    general_after_raw = (request.args.get("general_after_id") or "0").strip()
    try:
        direct_after_id = int(direct_after_raw)
    except ValueError:
        direct_after_id = 0
    try:
        general_after_id = int(general_after_raw)
    except ValueError:
        general_after_id = 0

    direct_messages = (
        ChatMessage.query
        .filter(ChatMessage.receiver_id == current_user.id)
        .filter(ChatMessage.sender_id != current_user.id)
        .filter(ChatMessage.id > direct_after_id)
        .order_by(ChatMessage.id.asc())
        .limit(20)
        .all()
    )

    general_messages = (
        ChatMessage.query
        .filter(ChatMessage.receiver_id.is_(None))
        .filter(ChatMessage.sender_id != current_user.id)
        .filter(ChatMessage.id > general_after_id)
        .order_by(ChatMessage.id.asc())
        .limit(20)
        .all()
    )

    direct_data = [
        {
            "id": m.id,
            "sender": get_chat_sender_name(m),
            "sender_id": m.sender_id,
            "text": m.text or "",
            "time": m.created_at.strftime("%H:%M"),
            "kind": "direct",
        }
        for m in direct_messages
    ]

    general_data = [
        {
            "id": m.id,
            "sender": get_chat_sender_name(m),
            "sender_id": m.sender_id,
            "text": m.text or "",
            "time": m.created_at.strftime("%H:%M"),
            "kind": "general",
        }
        for m in general_messages
    ]

    direct_last_id = direct_data[-1]["id"] if direct_data else direct_after_id
    general_last_id = general_data[-1]["id"] if general_data else general_after_id

    return jsonify({
        "direct_messages": direct_data,
        "general_messages": general_data,
        "direct_last_id": direct_last_id,
        "general_last_id": general_last_id,
    })


@app.get("/api/chat/tray/incoming")
@login_required
def chat_tray_incoming():
    direct_after_raw = (request.args.get("direct_after_id") or "0").strip()
    general_after_raw = (request.args.get("general_after_id") or "0").strip()
    try:
        direct_after_id = int(direct_after_raw)
    except ValueError:
        direct_after_id = 0
    try:
        general_after_id = int(general_after_raw)
    except ValueError:
        general_after_id = 0

    direct_messages = (
        ChatMessage.query
        .filter(ChatMessage.receiver_id == current_user.id)
        .filter(ChatMessage.sender_id != current_user.id)
        .filter(ChatMessage.id > direct_after_id)
        .order_by(ChatMessage.id.asc())
        .limit(20)
        .all()
    )

    general_messages = (
        ChatMessage.query
        .filter(ChatMessage.receiver_id.is_(None))
        .filter(ChatMessage.sender_id != current_user.id)
        .filter(ChatMessage.id > general_after_id)
        .order_by(ChatMessage.id.asc())
        .limit(20)
        .all()
    )

    direct_last_id = direct_messages[-1].id if direct_messages else direct_after_id
    general_last_id = general_messages[-1].id if general_messages else general_after_id

    all_messages = [
        {
            "id": m.id,
            "kind": "direct",
            "sender_id": m.sender_id or 0,
            "sender": get_chat_sender_name(m),
            "text": m.text or "",
            "time": m.created_at.strftime("%H:%M"),
        }
        for m in direct_messages
    ] + [
        {
            "id": m.id,
            "kind": "general",
            "sender_id": m.sender_id or 0,
            "sender": get_chat_sender_name(m),
            "text": m.text or "",
            "time": m.created_at.strftime("%H:%M"),
        }
        for m in general_messages
    ]
    all_messages.sort(key=lambda item: item["id"])

    lines = [
        f"DIRECT_LAST\t{direct_last_id}",
        f"GENERAL_LAST\t{general_last_id}",
    ]
    for item in all_messages:
        lines.append(
            "\t".join([
                "MSG",
                str(item["id"]),
                item["kind"],
                str(item["sender_id"]),
                quote(item["sender"], safe=""),
                quote(item["text"], safe=""),
                quote(item["time"], safe=""),
            ])
        )

    payload = "\n".join(lines)
    return current_app.response_class(payload, mimetype="text/plain; charset=utf-8")


@app.get("/api/chat/tray/users")
@login_required
def chat_tray_users():
    users = (
        User.query
        .filter(User.id != current_user.id, User.is_approved == True)
        .order_by(User.full_name.asc())
        .all()
    )

    lines = []
    for user in users:
        lines.append(
            "\t".join([
                "USER",
                str(user.id),
                quote(user.full_name or "", safe=""),
                quote(user.role or "", safe=""),
            ])
        )

    payload = "\n".join(lines)
    return current_app.response_class(payload, mimetype="text/plain; charset=utf-8")
@app.post("/letters/<int:letter_id>/request-access")
@login_required
def request_letter_access(letter_id):
    le = Letter.query.get_or_404(letter_id)
    author = User.query.get(le.author_id)
    message = request.form.get("message", "").strip()

    # 📎 Формируем ссылку на письмо
    letter_url = url_for("letter_view", letter_id=le.id, _external=True)

    # 📨 Текст уведомления автору
    msg_text = (
        f"📬 Пользователь <b>{current_user.full_name}</b> просит открыть доступ "
        f"к письму №{le.id} — <a href='{letter_url}'>перейти к письму</a>.<br><br>"
        f"Сообщение: {message or '(без текста)'}"
    )

    # 💬 Создаём сообщение автору (внутренний чат)
    chat_msg = ChatMessage(
        sender_id=current_user.id,
        receiver_id=author.id,
        text=msg_text
    )
    db.session.add(chat_msg)
    db.session.commit()

    flash("Запрос автору письма отправлен.", "success")
    return redirect(url_for("letters_page"))
@app.post("/letters/<int:letter_id>/grant-access")
@login_required
def grant_letter_access(letter_id):
    """Автор письма выдаёт доступ другому пользователю"""
    le = Letter.query.get_or_404(letter_id)

    # Только автор или админ может выдавать
    if current_user.id != le.author_id and current_user.role != "admin":
        return jsonify({"error": "Нет прав выдавать доступ"}), 403

    user_id = int(request.form.get("user_id", 0))
    if not user_id:
        return jsonify({"error": "Не указан пользователь"}), 400

    exists = DocumentAccess.query.filter_by(
        doc_type="letter", doc_id=le.id, user_id=user_id
    ).first()
    if exists:
        return jsonify({"error": "Уже есть доступ"}), 400

    acc = DocumentAccess(
        doc_type="letter",
        doc_id=le.id,
        user_id=user_id,
        granted_by=current_user.id
    )
    db.session.add(acc)
    db.session.commit()

    return jsonify({"success": True, "user_id": user_id})
@app.get("/letters/<int:letter_id>/access-list")
@login_required
def letter_access_list(letter_id):
    """Возвращает список пользователей, имеющих доступ к письму"""
    le = Letter.query.get_or_404(letter_id)

    # Только автор, админ или менеджер может смотреть
    if current_user.id != le.author_id and current_user.role not in ("admin", "manager"):
        return jsonify({"error": "Нет прав"}), 403

    accesses = (
        DocumentAccess.query
        .join(User, DocumentAccess.user_id == User.id)
        .filter(DocumentAccess.doc_type == "letter", DocumentAccess.doc_id == le.id)
        .with_entities(DocumentAccess.id, User.full_name.label("user_name"), User.id.label("user_id"))
        .all()
    )

    data = [{"id": a.id, "user_id": a.user_id, "user_name": a.user_name} for a in accesses]
    return jsonify(data)


@app.post("/letters/<int:letter_id>/revoke-access")
@login_required
def revoke_letter_access(letter_id):
    """Отзыв доступа к письму"""
    le = Letter.query.get_or_404(letter_id)

    if current_user.id != le.author_id and current_user.role != "admin":
        return jsonify({"error": "Нет прав отзывать доступ"}), 403

    user_id = int(request.form.get("user_id", 0))
    if not user_id:
        return jsonify({"error": "Не указан пользователь"}), 400

    DocumentAccess.query.filter_by(doc_type="letter", doc_id=le.id, user_id=user_id).delete()
    db.session.commit()

    return jsonify({"success": True})
@app.post("/documents/<doc_type>/<int:doc_id>/request-access")
@login_required
def request_document_access(doc_type, doc_id):
    """Пользователь запрашивает доступ к документу у автора"""
    if doc_type == "letter":
        doc = LetterDocument.query.get_or_404(doc_id)
        author = User.query.get(doc.letter.author_id)
        letter_url = url_for("letter_view", letter_id=doc.letter.id, _external=True)
        msg_text = (
            f"📎 Пользователь <b>{current_user.full_name}</b> просит доступ к документу "
            f"<i>{doc.filename}</i> из письма №{doc.letter.id}.<br>"
            f"<a href='{letter_url}'>Открыть письмо</a><br>"
            f"Сообщение: {request.form.get('message') or '(без текста)'}"
        )

    elif doc_type == "activity":
        doc = ActivityDocument.query.get_or_404(doc_id)
        author = User.query.get(doc.activity.owner_id)
        activity_url = url_for("activity_view", activity_id=doc.activity.id, _external=True)
        msg_text = (
            f"📎 Пользователь <b>{current_user.full_name}</b> просит доступ к документу "
            f"<i>{doc.filename}</i> из активности №{doc.activity.id}.<br>"
            f"<a href='{activity_url}'>Открыть активность</a><br>"
            f"Сообщение: {request.form.get('message') or '(без текста)'}"
        )

    else:
        abort(400, "Неизвестный тип документа")

    # 💬 Отправляем сообщение автору
    chat_msg = ChatMessage(
        sender_id=current_user.id,
        receiver_id=author.id,
        text=msg_text
    )
    db.session.add(chat_msg)
    db.session.commit()

    flash("Запрос автору отправлен.", "success")
    return redirect(url_for("dashboard"))
@app.get("/admin/download-logs")
@role_required("admin", "manager")
def admin_download_logs():
    logs = (
        DownloadLog.query
        .order_by(DownloadLog.timestamp.desc())
        .limit(100)
        .all()
    )
    return render_template("admin_download_logs.html", logs=logs, title="Логи загрузок документов")
@app.route("/files/upload", methods=["GET", "POST"])
@login_required
def upload_general_document():
    if request.method == "POST":
        file = request.files.get("file")
        desc = request.form.get("description", "").strip()

        if not file or not file.filename:
            flash("⚠️ Файл не выбран", "warning")
            return redirect(request.url)

        allowed_ext = {"png", "jpg", "jpeg", "pdf", "txt", "docx", "xlsx", "zip", "rar"}
        ext = file.filename.rsplit(".", 1)[-1].lower()
        if ext not in allowed_ext:
            flash(f"🚫 Недопустимый формат файла: .{ext}", "danger")
            return redirect(request.url)

        upload_dir = os.path.join(app.root_path, "static", "general_uploads")
        os.makedirs(upload_dir, exist_ok=True)

        filename = secure_filename_rus(file.filename)
        random_name = f"{uuid.uuid4().hex}_{filename}"
        save_path = os.path.join(upload_dir, random_name)
        file.save(save_path)

        doc = GeneralDocument(
            filename=filename,
            filepath=f"general_uploads/{random_name}",
            uploaded_by=current_user.id,
            description=desc
        )
        db.session.add(doc)
        db.session.commit()
        flash("📁 Документ успешно загружен.", "success")
        return redirect(url_for("files_list"))

    return render_template("file_upload.html", title="Загрузка документа")
@app.get("/files")
@login_required
def files_list():
    if current_user.role in ("admin", "manager"):
        docs = GeneralDocument.query.order_by(GeneralDocument.uploaded_at.desc()).all()
    else:
        docs = (
            GeneralDocument.query.outerjoin(DocumentAccess,
                (DocumentAccess.doc_type == "file") &
                (DocumentAccess.doc_id == GeneralDocument.id)
            )
            .filter(
                (GeneralDocument.uploaded_by == current_user.id) |
                (DocumentAccess.user_id == current_user.id)
            )
            .distinct()
            .order_by(GeneralDocument.uploaded_at.desc())
            .all()
        )

    # ✅ добавляем активные ссылки
    shared_links = {
        link.doc_id: link
        for link in SharedLink.query.filter_by(doc_type="file").all()
    }

    return render_template("files_list.html", docs=docs, shared_links=shared_links, title="Общие документы")

@app.route("/files/<int:file_id>")
@login_required
def view_file(file_id):
    doc = GeneralDocument.query.get_or_404(file_id)

    # Проверка доступа
    allowed = (
        current_user.role in ("admin", "manager") or
        doc.uploaded_by == current_user.id or
        DocumentAccess.query.filter_by(doc_type="file", doc_id=doc.id, user_id=current_user.id).count() > 0
    )

    if not allowed:
        author = User.query.get(doc.uploaded_by)
        return render_template(
            "document_request_access.html",
            title="Доступ к документу запрещён",
            doc=doc,
            author=author,
            doc_type="file"
        ), 403

    file_path = os.path.join("static", doc.filepath)
    if not os.path.exists(file_path):
        abort(404, "Файл отсутствует на сервере")

    # Лог загрузки
    db.session.add(DownloadLog(
        user_id=current_user.id,
        doc_type="file",
        doc_id=doc.id
    ))
    db.session.commit()

    return send_file(file_path, as_attachment=True, download_name=doc.filename)
@app.post("/files/<int:file_id>/share")
@login_required
def share_file(file_id):
    doc = GeneralDocument.query.get_or_404(file_id)
    if doc.uploaded_by != current_user.id and current_user.role != "admin":
        abort(403)

    # Удаляем старую ссылку, если есть
    SharedLink.query.filter_by(doc_type="file", doc_id=doc.id).delete()

    token = uuid.uuid4().hex
    days = int(request.form.get("days", 0))  # 0 = бессрочно
    expires_at = None
    if days > 0:
        expires_at = dt.datetime.utcnow() + dt.timedelta(days=days)

    link = SharedLink(
        doc_type="file",
        doc_id=doc.id,
        token=token,
        created_by=current_user.id,
        expires_at=expires_at
    )
    db.session.add(link)
    db.session.commit()

    flash("✅ Ссылка создана", "success")
    return redirect(url_for("files_list"))

@app.route("/share/<token>")
@login_required
def open_shared_link(token):
    """Открытие документа по внутренней ссылке"""
    link = SharedLink.query.filter_by(token=token).first_or_404()

    # Проверяем срок действия
    if link.expires_at and link.expires_at < dt.datetime.utcnow():
        abort(403, "Ссылка просрочена")

    # Проверяем тип документа
    if link.doc_type == "file":
        doc = GeneralDocument.query.get_or_404(link.doc_id)
        file_path = os.path.join("static", doc.filepath)
        if not os.path.exists(file_path):
            abort(404, "Файл отсутствует на сервере")

        # Проверка доступа (только авторизованные пользователи)
        allowed = (
            current_user.role in ("admin", "manager") or
            doc.uploaded_by == current_user.id or
            DocumentAccess.query.filter_by(
                doc_type="file", doc_id=doc.id, user_id=current_user.id
            ).count() > 0
        )
        if not allowed:
            author = User.query.get(doc.uploaded_by)
            return render_template(
                "document_request_access.html",
                title="Доступ к документу запрещён",
                doc=doc,
                author=author,
                doc_type="file"
            ), 403

        # Лог загрузки
        db.session.add(DownloadLog(
            user_id=current_user.id,
            doc_type="file",
            doc_id=doc.id
        ))
        db.session.commit()

        return send_file(file_path, as_attachment=True, download_name=doc.filename)

    abort(400, "Неизвестный тип документа")
@app.get("/files/<int:file_id>/access")
@login_required
def file_access(file_id):
    """Управление доступом к файлу (только автор или админ)"""
    doc = GeneralDocument.query.get_or_404(file_id)
    if doc.uploaded_by != current_user.id and current_user.role != "admin":
        abort(403)

    users = User.query.filter(User.id != current_user.id, User.is_approved == True).order_by(User.full_name).all()
    access_list = (
        DocumentAccess.query
        .join(User, DocumentAccess.user_id == User.id)
        .filter(DocumentAccess.doc_type == "file", DocumentAccess.doc_id == doc.id)
        .with_entities(DocumentAccess.user_id, User.full_name)
        .all()
    )

    return render_template("file_access.html", doc=doc, users=users, access_list=access_list, title="Доступ к файлу")


@app.post("/files/<int:file_id>/grant-access")
@login_required
def grant_file_access(file_id):
    doc = GeneralDocument.query.get_or_404(file_id)
    if doc.uploaded_by != current_user.id and current_user.role != "admin":
        abort(403)

    user_id = int(request.form.get("user_id", 0))
    if not user_id:
        return jsonify({"error": "Не выбран пользователь"}), 400

    exists = DocumentAccess.query.filter_by(doc_type="file", doc_id=doc.id, user_id=user_id).first()
    if exists:
        return jsonify({"error": "Уже есть доступ"}), 400

    db.session.add(DocumentAccess(
        doc_type="file",
        doc_id=doc.id,
        user_id=user_id,
        granted_by=current_user.id
    ))
    db.session.commit()
    return jsonify({"success": True})


@app.post("/files/<int:file_id>/revoke-access")
@login_required
def revoke_file_access(file_id):
    doc = GeneralDocument.query.get_or_404(file_id)
    if doc.uploaded_by != current_user.id and current_user.role != "admin":
        abort(403)

    user_id = int(request.form.get("user_id", 0))
    DocumentAccess.query.filter_by(doc_type="file", doc_id=doc.id, user_id=user_id).delete()
    db.session.commit()
    return jsonify({"success": True})
@app.get("/admin/users/<int:user_id>/show-password")
@role_required("admin")
def show_password(user_id):
    if current_user.role != "admin":
        abort(403)
    u = User.query.get_or_404(user_id)
    return render_template("admin_show_password.html",
                           user=u,
                           password=u.get_password())
@app.post("/admin/users/<int:user_id>/toggle-admin")
@role_required("superadmin")
def toggle_admin(user_id):
    """Суперадмин назначает/снимает администратора"""
    u = User.query.get_or_404(user_id)
    if u.role == "admin":
        u.role = "worker"
        flash(f"❎ {u.full_name} больше не администратор", "info")
    else:
        u.role = "admin"
        flash(f"✅ {u.full_name} назначен администратором", "success")
    db.session.commit()
    return redirect(url_for("admin_users"))


@app.post("/admin/users/<int:user_id>/role")
@role_required("admin", "superadmin")
def update_user_role(user_id):
    user = User.query.get_or_404(user_id)
    if not can_manage_user(user):
        abort(403)
    role = (request.form.get("role") or "worker").strip()
    allowed_roles = {"worker", "deputy", "manager"}
    if current_user.role == "superadmin":
        allowed_roles.add("admin")
    if role not in allowed_roles:
        abort(400)
    user.role = role
    db.session.commit()
    flash(f"Роль пользователя {user.full_name} обновлена.", "success")
    return redirect(url_for("admin_users"))
@app.get("/admin/users/<int:user_id>/password")
@role_required("admin", "superadmin")
def admin_show_password(user_id):
    """Просмотр расшифрованного пароля"""
    u = User.query.get_or_404(user_id)
    if not can_manage_user(u):
        abort(403)
    try:
        decrypted = decrypt_str(u.password_enc, SECRET_PASSPHRASE)
    except Exception:
        decrypted = "(ошибка расшифровки)"
    return jsonify({"password": decrypted})
# === Личный кабинет ===
@app.get("/profile")
@login_required
def profile_page():
    """Страница профиля"""
    return render_template("profile.html", title="Личный кабинет", user=current_user)


@app.post("/profile/change-password")
@login_required
def change_password():
    """Смена пароля пользователем"""
    old_pass = request.form.get("old_password", "")
    new_pass = request.form.get("new_password", "")
    confirm = request.form.get("confirm_password", "")

    # Проверяем ввод
    if not old_pass or not new_pass or not confirm:
        flash("Заполните все поля.", "warning")
        return redirect(url_for("profile_page"))

    if new_pass != confirm:
        flash("Новый пароль и подтверждение не совпадают.", "danger")
        return redirect(url_for("profile_page"))

    # Проверка старого пароля
    user = User.query.get(current_user.id)
    if not user.check_password(old_pass):
        flash("Неверный текущий пароль.", "danger")
        return redirect(url_for("profile_page"))

    # Устанавливаем новый
    user.set_password(new_pass)
    db.session.commit()
    flash("✅ Пароль успешно изменён.", "success")
    return redirect(url_for("profile_page"))
# === Доска заданий ===
from datetime import datetime, date

@app.get("/board")
@login_required
def board_page():
    """Просмотр доски заданий с вычисленным статусом"""
    if current_user.role in ("admin", "superadmin"):
        tasks = Task.query.order_by(Task.end_date.desc()).all()
    else:
        tasks = (
            Task.query.join(TaskAssignee)
            .filter(TaskAssignee.user_id == current_user.id)
            .order_by(Task.end_date.desc())
            .all()
        )

    # ⚙️ Добавим вычисленные поля для отображения
    today = date.today()
    for t in tasks:
        t.all_done = all(a.is_done for a in t.assignees)
        t.expired = t.end_date < today
        t.in_progress = not t.all_done and not t.expired

    users = User.query.all()
    return render_template("board.html", tasks=tasks, users=users)


@app.post("/board/add")
@role_required("admin", "superadmin")
def add_task():
    title = request.form.get("title")
    desc = request.form.get("description")
    assigned_to_list = request.form.getlist("assigned_to")  # ✅ несколько пользователей
    start = request.form.get("start_date")
    end = request.form.get("end_date")

    if not title or not assigned_to_list:
        flash("Введите название и выберите хотя бы одного исполнителя.", "warning")
        return redirect(url_for("board_page"))

    try:
        start_date = datetime.strptime(start, "%Y-%m-%d").date()
        end_date = datetime.strptime(end, "%Y-%m-%d").date()
    except Exception:
        flash("Неверный формат даты.", "danger")
        return redirect(url_for("board_page"))

    task = Task(
        title=title,
        description=desc,
        created_by=current_user.id,
        start_date=start_date,
        end_date=end_date,
    )
    db.session.add(task)
    db.session.flush()  # получаем id задания

    # Добавляем всех исполнителей
    for uid in assigned_to_list:
        db.session.add(TaskAssignee(task_id=task.id, user_id=int(uid)))

    db.session.commit()
    flash("✅ Задание создано и назначено пользователям.", "success")
    return redirect(url_for("board_page"))

@app.post("/board/<int:task_id>/toggle")
@login_required
def toggle_task(task_id):
    link = TaskAssignee.query.filter_by(task_id=task_id, user_id=current_user.id).first_or_404()
    link.is_done = not link.is_done
    link.done_at = dt.datetime.utcnow() if link.is_done else None
    db.session.commit()
    return redirect(url_for("board_page"))

if __name__ == "__main__":
    with app.app_context():
        db.create_all()
        init_fts_indexes()  # <=== 👈 вызываем вручную при запуске
        print("✔ Проверка таблиц выполнена:", [t.name for t in db.metadata.sorted_tables])
    start_ocr_worker()
    app.run(
        host="0.0.0.0",
        port=5001,
        debug=True,
        use_reloader=False,
    )

